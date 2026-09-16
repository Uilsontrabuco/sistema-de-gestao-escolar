"""Importação com prévia versionada; documentos institucionais gerados no servidor."""
import base64, copy, csv, hashlib, io, json, re, secrets, time, zipfile
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape
from server import allowed, require, now, fold

SCHOOL='Colégio Adventista de Juazeiro'
ADDRESS='R. Antônio Pedro, 263 - Centro, Juazeiro - BA, 48903-660'
CNPJ='CNPJ: 07.114.699/0047-42'
ALIASES={'code':['codigo','cod','codigo da conta','conta codigo'],'category':['categoria','conta','nome da conta','categoria/conta'],'subaccount':['subconta'],'description':['descricao','historico'],'purpose':['finalidade'],'budget':['orcado','orcamento','valor orcado'],'value':['valor','realizado','valor realizado','valor pago'],'date':['data','data pagamento','data de pagamento'],'document':['documento','id','numero documento','identificador'],'month':['mes','competencia'],'financialPercent':['financeira','inadimplencia financeira','financeira (%)','inadimplencia financeira (%)'],'accountingPercent':['contabil','inadimplencia contabil','contabil (%)','inadimplencia contabil (%)'],'debt':['divida','valor da divida','divida financeira'],'students':['alunos','alunos afetados'],'guardians':['responsaveis','responsaveis afetados']}
FINANCIAL_IMPORT_ALIASES={'className':['turma','classe','turma/serie','turma/série'],'grade':['serie','série','ano','serie/ano','série/ano'],'category':['beneficio','benefício','tipo','categoria','categoria/beneficio','categoria/benefício','bolsa/desconto','tipo de beneficio','tipo de benefício'],'percent':['percentual','percentual de desconto','desconto (%)','percentual (%)','desconto','% desconto'],'quantity':['quantidade','qtd','alunos','quantidade de alunos','qtde'],'student':['aluno','nome do aluno','nome','matricula','matrícula'],'note':['observacao','observação','obs']}
REVENUE_IMPORT_ALIASES={'studentId':['id aluno','identificador unico','identificador único','matricula','matrícula','codigo aluno','código aluno'],'externalId':['id lancamento','id lançamento','id financeiro','documento','referencia','referência'],'className':['turma','classe','turma/serie'],'eventType':['tipo','evento','tipo de lancamento','tipo de lançamento'],'date':['data','data matricula','data matrícula','competencia','competência'],'gross':['valor bruto','bruto','valor'],'discount':['desconto','valor desconto'],'status':['status','situacao','situação'],'settledAmount':['baixa financeira','valor baixado','valor recebido','recebido'],'reversalOf':['estorno de','reversao de','reversão de','lancamento original','lançamento original'],'fromClassName':['turma origem','classe origem']}
FINANCIAL_FIXED={
    ('sem desconto',0):'noDiscount',('bolsa filantropica',100):'philanthropic100',('bolsa filantropica',50):'philanthropic50',
    ('filho de funcionario',100):'staffChild100',('filho de funcionario',80):'staffChild80',('filho de obreiro',100):'workerChild100',('projeto marcando vidas',45):'markingLives45'
}
def normalize_teaching_class_code(code,classes):
    """Interpreta somente gramáticas comprovadas; ausência de seção operacional fica pendente."""
    value=str(code or '').strip().upper();rooms={fold(x['name']):x for x in classes};result={'source_class_code':value,'parsed_segment':None,'parsed_grade':None,'parsed_shift':None,'parsed_section':None,'operational_class_id':None,'status':'PENDENTE'}
    match=re.fullmatch(r'EFUND(\d{2})([MT])([A-Z])',value)
    if match:
        grade=str(int(match.group(1)))+'º';result.update(parsed_segment='Ensino Fundamental',parsed_grade=grade,parsed_shift={'M':'Manhã','T':'Tarde'}[match.group(2)],parsed_section=match.group(3));room=rooms.get(fold(grade+' '+match.group(3)))
        if room:result.update(operational_class_id=room['id'],status='CONCILIADO')
        return result
    match=re.fullmatch(r'(EMERE|EMICN|EMICH)(\d{2})([MT])([A-Z])',value)
    if match:
        grade=str(int(match.group(2)))+'º EM';result.update(parsed_segment='Ensino Médio',parsed_grade=grade,parsed_shift={'M':'Manhã','T':'Tarde'}[match.group(3)],parsed_section=match.group(4));room=rooms.get(fold(grade))
        # A fonte separa seção, mas as 41 turmas operacionais de EM são agregadas: não é vínculo inequívoco.
        if room and match.group(4)=='A':result.update(operational_class_id=room['id'],status='PENDENTE')
        return result
    # EINFA não apareceu no PDF fornecido com gramática comprovável; permanece auditavelmente pendente.
    if value.startswith('EINFA'):result['parsed_segment']='Educação Infantil'
    return result
def extract_teaching_occurrence(professor,professor_id,day,lesson_text,timing_text):
    codes=re.findall(r'\b(?:EINFA|EFUND|EMERE|EMICN|EMICH)[A-Z0-9]+',lesson_text or '')
    timing=re.search(r'(\d{2}:\d{2})\s*-\s*(\d{2,3})',timing_text or '')
    start=timing.group(1) if timing else None;duration=int(timing.group(2)) if timing else None;end=None
    if start and duration is not None:
        h,m=map(int,start.split(':'));total=h*60+m+duration;end=f'{total//60:02d}:{total%60:02d}'
    pairs=re.findall(r'\b((?:EINFA|EFUND|EMERE|EMICN|EMICH)[A-Z0-9]+)\s*-\s*([^|\n]+)',lesson_text or '')
    return {'professor':professor or None,'professor_id':professor_id or None,'day_of_week':day or None,'source_class_codes':codes,'source_class_code':codes[0] if len(codes)==1 else None,'discipline':pairs[0][1].strip() if len(pairs)==1 else None,'start_time':start,'end_time':end,'duration_minutes':duration,'shared_class':len(codes)>1 or '|' in (lesson_text or ''),'source_text':lesson_text or '', 'status':'PENDENTE' if not codes or not timing else 'EXTRACTED'}
def teaching_load_positional_pilot(content,classes,page_number=1):
    """Reconstrói somente uma página semanal por geometria vetorial, sem persistir dados."""
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError('pdfplumber é obrigatório para a prévia posicional da carga horária.') from exc
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        if page_number<1 or page_number>len(pdf.pages):raise ValueError('Página piloto inexistente.')
        page=pdf.pages[page_number-1]
        words=page.extract_words(x_tolerance=1,y_tolerance=2,use_text_flow=False)
        raw_lines=page.lines
    professor_line=' '.join(w['text'] for w in words if 73<=w['top']<=85)
    teacher=re.search(r'Professor:\s*(.+?)\s*\(([^)]+)\)',professor_line)
    professor=teacher.group(1).strip() if teacher else None;professor_id=teacher.group(2).strip() if teacher else None
    turn_line=' '.join(w['text'] for w in words if 89<=w['top']<=101);turn=re.search(r'Turno:\s*(.+)',turn_line)
    weekday_names={'segunda':'Segunda','terca':'Terça','quarta':'Quarta','quinta':'Quinta','sexta':'Sexta'}
    headers=[]
    for word in words:
        # Alguns PDFs legados expõem ç como U+FFFD; somente normalizamos o rótulo já posicionado no cabeçalho.
        name=weekday_names.get(fold(word['text']).replace('\ufffd','c'))
        if name:headers.append((name,word))
    headers.sort(key=lambda item:item[1]['x0'])
    # As linhas verticais longas delimitam fisicamente as células; os nomes dos dias apenas nomeiam cada faixa.
    raw_verticals={round(line['x0'],1) for line in raw_lines if abs(line['x0']-line['x1'])<0.2 and line['top']<=134 and line['bottom']>=198}
    for rect in page.rects:
        if rect['top']<=134 and rect['bottom']>=198:
            raw_verticals.update([round(rect['x0'],1),round(rect['x1'],1)])
    raw_verticals=sorted(raw_verticals)
    verticals=[]
    for value in raw_verticals:
        if not verticals or value-verticals[-1][-1]>2:verticals.append([value])
        else:verticals[-1].append(value)
    verticals=[round(sum(group)/len(group),1) for group in verticals]
    if not headers or len(headers)>5:raise ValueError('Quadro semanal não pôde ser comprovado geometricamente na página piloto.')
    header_centers=[(word['x0']+word['x1'])/2 for _,word in headers]
    candidates=[verticals[index:index+len(headers)+1] for index in range(max(0,len(verticals)-len(headers)))]
    verticals=next((candidate for candidate in candidates if all(candidate[index]<=header_centers[index]<=candidate[index+1] for index in range(len(headers)))),None)
    if not verticals:raise ValueError('Quadro semanal não pôde ser comprovado geometricamente na página piloto.')
    columns=[]
    for index,(day,_) in enumerate(headers):
        left,right=verticals[index],verticals[index+1]
        columns.append({'day':day,'x0':left,'x1':right})
    legend_word=next((word for word in words if fold(word['text']).startswith('legenda')),None)
    legend_top=legend_word['top'] if legend_word else page.height
    # Cada linha de horário tem duas regras horizontais no primeiro dia: início e fim da faixa temporal.
    timing_rules={round(line['top'],1) for line in raw_lines if abs(line['x0']-columns[0]['x0'])<1 and abs(line['x1']-columns[0]['x1'])<1 and 134<line['top']<legend_top}
    for rect in page.rects:
        if abs(rect['x0']-columns[0]['x0'])<1 and abs(rect['x1']-columns[0]['x1'])<1 and rect['top']<legend_top and 134<rect['bottom']<legend_top:timing_rules.add(round(rect['bottom'],1))
    timing_rules=sorted(timing_rules)
    pairs=[(timing_rules[i],timing_rules[i+1]) for i in range(0,len(timing_rules)-1,2) if 7<=timing_rules[i+1]-timing_rules[i]<=13]
    content_starts=[133.5]+[round(end+0.6,1) for _,end in pairs[:-1]]
    rows=[]
    for index,((time_top,time_bottom),content_top) in enumerate(zip(pairs,content_starts),1):
        rows.append({'slot':index,'content_y0':content_top,'content_y1':round(time_top-0.1,1),'time_y0':round(time_top+0.5,1),'time_y1':time_bottom})
    def text_in(box):
        x0,x1,y0,y1=box
        row_words=[word for word in words if x0<=((word['x0']+word['x1'])/2)<=x1 and y0<=word['top']<=y1]
        row_words.sort(key=lambda word:(word['top'],word['x0']))
        return ' '.join(word['text'] for word in row_words)
    def clock(start,duration):
        hour,minute=map(int,start.split(':'));total=hour*60+minute+duration
        return f'{total//60:02d}:{total%60:02d}'
    occurrences=[];time_bands=[]
    for row in rows:
        samples=[]
        for col in columns:
            timing=text_in((col['x0'],col['x1'],row['time_y0'],row['time_y1']))
            match=re.search(r'(\d{2}:\d{2})\s*-\s*(\d{2,3})',timing)
            if match:samples.append((match.group(1),int(match.group(2))))
        start,duration=(samples[0] if samples else (None,None))
        time_bands.append({**row,'start_time':start,'duration_minutes':duration,'end_time':clock(start,duration) if start else None,'status':'CONFIRMADO' if start else 'PENDENTE'})
        for col in columns:
            cell_words=[word for word in words if col['x0']<=((word['x0']+word['x1'])/2)<=col['x1'] and row['content_y0']<=word['top']<=row['content_y1']]
            cell_words.sort(key=lambda word:(word['top'],word['x0']));cell_text=' '.join(word['text'] for word in cell_words)
            pairs=re.findall(r'((?:EINFA|EFUND|EMERE|EMICN|EMICH)[A-Z0-9]+)\s*-\s*([^|\n]+)',cell_text)
            if not pairs:continue
            timing=text_in((col['x0'],col['x1'],row['time_y0'],row['time_y1']));match=re.search(r'(\d{2}:\d{2})\s*-\s*(\d{2,3})',timing)
            start=match.group(1) if match else None;duration=int(match.group(2)) if match else None
            codes=[pair[0] for pair in pairs];shared='|' in cell_text or len(codes)>1;normalizations=[normalize_teaching_class_code(code,classes) for code in codes]
            occurrences.append({'professor':professor,'professor_id':professor_id,'page':page_number,'turn':turn.group(1).strip() if turn else None,'day_of_week':col['day'],'start_time':start,'end_time':clock(start,duration) if start else None,'duration_minutes':duration,'source_class_code':codes[0] if not shared else None,'source_class_codes':codes,'discipline':pairs[0][1].strip() if not shared else None,'disciplines':[pair[1].strip() for pair in pairs],'shared_class':shared,'shared_note':'Aula compartilhada – regra de distribuição pendente' if shared else None,'source_text':cell_text,'cell_x':round(cell_words[0]['x0'],1),'cell_y':round(cell_words[0]['top'],1),'column_x0':col['x0'],'column_x1':col['x1'],'time_band':row['slot'],'status_temporal':'CONFIRMADO' if match else 'PENDENTE','status_normalization':'PENDENTE' if shared else normalizations[0]['status'],'normalization':normalizations[0] if not shared else normalizations})
    # O texto da primeira linha pode estar alguns décimos acima da palavra "Legenda"; ele continua fora do quadro.
    legend_items=[text_in((0,page.width,max(0,legend_top-1),page.height))] if legend_word else []
    return {'page':page_number,'professor':professor,'professor_id':professor_id,'turn':turn.group(1).strip() if turn else None,'columns':columns,'time_bands':time_bands,'occurrences':occurrences,'legend_ignored':bool(legend_word),'legend_text':legend_items,'warnings':[] if occurrences else ['Nenhuma célula ocupada foi comprovada geometricamente.']}
def teaching_load_geometric_preview(content,classes):
    """Processa um PDF inteiro apenas como prévia auditável; não persiste nem altera turmas."""
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError('pdfplumber é obrigatório para a prévia geométrica da carga horária.') from exc
    with pdfplumber.open(io.BytesIO(content)) as pdf:page_count=len(pdf.pages)
    class_names={item['id']:item['name'] for item in classes};records=[];page_reports=[];empty_cells=legend_items=0;warnings=[]
    for page in range(1,page_count+1):
        try:parsed=teaching_load_positional_pilot(content,classes,page)
        except ValueError as exc:
            page_reports.append({'page':page,'status':'PENDENTE','reason':'Página sem quadro semanal geométrico; contém somente material documental/legenda.','occurrences':0});warnings.append({'page':page,'reason':str(exc)});continue
        empty_cells+=len(parsed['columns'])*len(parsed['time_bands'])-len(parsed['occurrences'])
        legend_items+=len(re.findall(r'(?:EINFA|EFUND|EMERE|EMICN|EMICH)[A-Z0-9]+\s*-',' '.join(parsed['legend_text'])))
        page_reports.append({'page':page,'status':'CONFIRMADO','professor':parsed['professor'],'occurrences':len(parsed['occurrences']),'empty_cells_ignored':len(parsed['columns'])*len(parsed['time_bands'])-len(parsed['occurrences']),'legend_ignored':parsed['legend_ignored']})
        for row in parsed['occurrences']:
            normalizations=row['normalization'] if isinstance(row['normalization'],list) else [row['normalization']]
            codes=row['source_class_codes'];families=sorted({code[:5] for code in codes});grades=sorted({item['parsed_grade'] for item in normalizations if item.get('parsed_grade')});sections=sorted({item['parsed_section'] for item in normalizations if item.get('parsed_section')});candidates=sorted({class_names[item['operational_class_id']] for item in normalizations if item.get('operational_class_id') in class_names})
            link_status='CONFIRMADO' if all(item.get('status')=='CONCILIADO' for item in normalizations) and not row['shared_class'] else 'PENDENTE'
            reason='Aula compartilhada – critério de distribuição e vínculo operacional aguardam aprovação.' if row['shared_class'] else ('Código/turma operacional não possui correspondência inequívoca.' if link_status=='PENDENTE' else None)
            records.append({'pagina':row['page'],'professor':row['professor'],'professor_id':row['professor_id'],'dia_semana':row['day_of_week'],'hora_inicio':row['start_time'],'hora_fim':row['end_time'],'duracao_minutos':row['duration_minutes'],'codigo_original_pdf':' | '.join(codes),'codigos_originais_pdf':codes,'familia_codigo':' | '.join(families),'disciplina_componente':' | '.join(row['disciplines']),'ano_serie':' | '.join(grades) if grades else None,'turno':row['turn'],'secao':' | '.join(sections) if sections else None,'turma_operacional_7e7':' | '.join(candidates) if candidates else None,'aula_compartilhada':row['shared_class'],'status_temporal':row['status_temporal'],'status_normalizacao':'CONFIRMADO' if all(item.get('status')=='CONCILIADO' for item in normalizations) else 'PENDENTE','status_vinculo_turma':link_status,'motivo_pendencia':reason,'evidencia_original':f"X={row['cell_x']}; Y={row['cell_y']}; faixa={row['time_band']}; {row['source_text']}",'source_cell_x':row['cell_x'],'source_cell_y':row['cell_y'],'source_class_codes':codes})
    professor_summary={};series_summary={}
    for row in records:
        professor=row['professor'] or 'PENDENTE – professor não identificado';p=professor_summary.setdefault(professor,{'professor':professor,'ocorrencias_semanais':0,'minutos_semanais':0,'duracoes':{40:0,45:0,50:0},'sem_duracao':0,'series':set(),'componentes':set(),'aulas_compartilhadas':0,'pendencias':0})
        key=row['turma_operacional_7e7'] or row['ano_serie'] or 'PENDENTE – série não identificada';s=series_summary.setdefault(key,{'turma_serie':key,'componentes':set(),'professores':set(),'ocorrencias_semanais':0,'minutos_semanais':0,'aulas_compartilhadas':0,'pendencias':0})
        for target in (p,s):target['ocorrencias_semanais']+=1;target['minutos_semanais']+=row['duracao_minutos'] or 0;target['aulas_compartilhadas']+=int(row['aula_compartilhada']);target['pendencias']+=int(row['status_vinculo_turma']=='PENDENTE')
        p['duracoes'][row['duracao_minutos']]=p['duracoes'].get(row['duracao_minutos'],0)+1 if row['duracao_minutos'] else p['duracoes'].get(row['duracao_minutos'],0);p['sem_duracao']+=int(row['duracao_minutos'] is None);p['series'].add(key);p['componentes'].add(row['disciplina_componente']);s['componentes'].add(row['disciplina_componente']);s['professores'].add(professor)
    overlaps=[];seen_cells=set();duplicates=[];grouped={}
    for row in records:
        cell=(row['pagina'],row['source_cell_x'],row['source_cell_y'])
        if cell in seen_cells:duplicates.append(row)
        seen_cells.add(cell)
        if row['hora_inicio'] and row['hora_fim']:grouped.setdefault((row['professor'],row['turno'],row['dia_semana']),[]).append(row)
    for key,rows in grouped.items():
        rows.sort(key=lambda row:row['hora_inicio'])
        for left,right in zip(rows,rows[1:]):
            if left['hora_fim']>right['hora_inicio']:overlaps.append({'professor':key[0],'turno':key[1],'dia':key[2],'left':left['evidencia_original'],'right':right['evidencia_original']})
    professors=[]
    for item in professor_summary.values():
        item['horas_semanais']=round(item['minutos_semanais']/60,2);item['series']=sorted(item['series']);item['componentes']=sorted(item['componentes']);professors.append(item)
    series=[]
    for item in series_summary.values():
        item['horas_semanais']=round(item['minutos_semanais']/60,2);item['componentes']=sorted(item['componentes']);item['professores']=sorted(item['professores']);series.append(item)
    return {'status':'PREVIEW','pages_processed':page_count,'page_reports':page_reports,'records':records,'professors':sorted(professors,key=lambda item:item['professor']),'series':sorted(series,key=lambda item:item['turma_serie']),'summary':{'occurrences':len(records),'professors':len(professor_summary),'empty_cells_ignored':empty_cells,'legend_items_ignored':legend_items,'temporal_confirmed':sum(item['status_temporal']=='CONFIRMADO' for item in records),'pending':sum(item['status_vinculo_turma']=='PENDENTE' for item in records),'shared_lessons':sum(item['aula_compartilhada'] for item in records),'unique_codes':len({code for item in records for code in item['source_class_codes']}),'components':len({item['disciplina_componente'] for item in records})},'inconsistencies':{'overlaps':overlaps,'duplicate_cells':duplicates,'pages_without_weekly_grid':[item['page'] for item in page_reports if item['status']=='PENDENTE'],'missing_teacher':sum(not item['professor'] for item in records),'missing_duration':sum(item['duracao_minutos'] is None for item in records)}}

# Valores "Reajustado" da tabela "Valores de Hora Aula" do Orçamento 2027.
# Esta base é exclusivamente uma referência semanal por ocorrência da grade. O
# documento não comprova equivalência entre aulas de 40/45/50 minutos, DSR,
# hora-atividade ou multiplicador mensal; por isso nenhum desses fatores é aplicado.
TEACHING_HOUR_AULA_REAJUSTED_2027={
    'Educação Infantil':16.20,
    'Fundamental I':16.20,
    'Fundamental II':26.11,
    'Ensino Médio':38.50,
    '3º EM':38.50,
}

def teaching_segment_for_operational_class(class_name):
    """Classifica somente o segmento de uma turma operacional já existente."""
    value=fold(class_name)
    if re.fullmatch(r'g[2-5] [a-z]',value):return 'Educação Infantil'
    grade=re.match(r'(\d+)(?:º|o)?\s+([a-z])$',value)
    if grade:
        number=int(grade.group(1))
        if 1<=number<=5:return 'Fundamental I'
        if 6<=number<=9:return 'Fundamental II'
    if re.fullmatch(r'[123](?:º|o)? em',value):return '3º EM' if value.startswith('3') else 'Ensino Médio'
    return None

def build_teaching_weekly_cost_audit(preview,classes):
    """Monta uma conferência não persistente da grade x valor de hora-aula 2027.

    Aulas compartilhadas e vínculos pendentes são deliberadamente excluídos do
    valor monetário: ainda não existe critério documental de distribuição.
    """
    by_name={fold(item['name']):item for item in classes};by_id={item['id']:item for item in classes}
    class_rows={item['id']:{'class_id':item['id'],'class_name':item['name'],'segment':teaching_segment_for_operational_class(item['name']),'weekly_lessons':0,'weekly_minutes':0,'teachers':set(),'disciplines':set(),'reference_weekly_cost':0.0,'shared_lessons_pending':0,'pending_occurrences':0,'pending_reasons':set(),'source_codes':set()} for item in classes}
    teacher_rows={};segment_rows={segment:{'segment':segment,'hour_aula_reajustado':rate,'weekly_lessons':0,'weekly_minutes':0,'reference_weekly_cost':0.0,'teachers':set(),'classes':set()} for segment,rate in TEACHING_HOUR_AULA_REAJUSTED_2027.items()}
    pending=[];included_occurrences=[];included=0
    def affected_classes(row):
        found=[]
        explicit=by_name.get(fold(row.get('turma_operacional_7e7') or ''))
        if explicit:found.append(explicit)
        for code in row.get('source_class_codes') or []:
            normalized=normalize_teaching_class_code(code,classes);room=by_id.get(normalized.get('operational_class_id'))
            if room and room not in found:found.append(room)
        return found
    for row in preview.get('records',[]):
        teacher=row.get('professor') or 'PENDENTE – professor não identificado';duration=row.get('duracao_minutos') or 0
        teacher_row=teacher_rows.setdefault(teacher,{'professor':teacher,'professor_id':row.get('professor_id'),'weekly_lessons':0,'weekly_minutes':0,'durations':{40:0,45:0,50:0},'classes':set(),'disciplines':set(),'shared_lessons':0,'pending_occurrences':0,'reference_weekly_cost':0.0,'by_class_discipline':{}})
        teacher_row['weekly_lessons']+=1;teacher_row['weekly_minutes']+=duration;teacher_row['durations'][duration]=teacher_row['durations'].get(duration,0)+1 if duration else teacher_row['durations'].get(duration,0)
        teacher_row['disciplines'].add(row.get('disciplina_componente') or 'PENDENTE')
        targets=affected_classes(row)
        for room in targets:teacher_row['classes'].add(room['name'])
        reason=None
        if row.get('status_temporal')!='CONFIRMADO':reason='Horário/duração não confirmados geometricamente.'
        elif row.get('aula_compartilhada'):reason='Aula compartilhada – rateio financeiro aguardando critério documental aprovado.'
        elif row.get('status_vinculo_turma')!='CONFIRMADO' or len(targets)!=1:reason=row.get('motivo_pendencia') or 'Vínculo com turma operacional não inequívoco.'
        elif not teaching_segment_for_operational_class(targets[0]['name']):reason='Segmento da turma operacional não identificado.'
        if reason:
            teacher_row['pending_occurrences']+=1;teacher_row['shared_lessons']+=int(bool(row.get('aula_compartilhada')))
            for room in targets:
                item=class_rows[room['id']];item['pending_occurrences']+=1;item['shared_lessons_pending']+=int(bool(row.get('aula_compartilhada')));item['pending_reasons'].add(reason);item['source_codes'].update(row.get('source_class_codes') or [])
            pending.append({'professor':teacher,'page':row.get('pagina'),'source_codes':row.get('source_class_codes') or [],'discipline':row.get('disciplina_componente'),'duration_minutes':duration,'shared_class':bool(row.get('aula_compartilhada')),'reason':reason,'evidence':row.get('evidencia_original')})
            continue
        room=targets[0];segment=teaching_segment_for_operational_class(room['name']);rate=TEACHING_HOUR_AULA_REAJUSTED_2027[segment]
        # Uma ocorrência da grade equivale a uma unidade da rubrica "hora-aula".
        # Não há conversão entre as durações reais e não há mensalização nesta etapa.
        cost=rate;included+=1
        item=class_rows[room['id']];item['weekly_lessons']+=1;item['weekly_minutes']+=duration;item['teachers'].add(teacher);item['disciplines'].add(row.get('disciplina_componente') or 'PENDENTE');item['reference_weekly_cost']+=cost;item['source_codes'].update(row.get('source_class_codes') or [])
        segment_item=segment_rows[segment];segment_item['weekly_lessons']+=1;segment_item['weekly_minutes']+=duration;segment_item['reference_weekly_cost']+=cost;segment_item['teachers'].add(teacher);segment_item['classes'].add(room['name'])
        teacher_row['reference_weekly_cost']+=cost
        group=teacher_row['by_class_discipline'].setdefault((room['name'],row.get('disciplina_componente') or 'PENDENTE'),{'class_name':room['name'],'segment':segment,'discipline':row.get('disciplina_componente') or 'PENDENTE','weekly_lessons':0,'weekly_minutes':0,'hour_aula_reajustado':rate,'reference_weekly_cost':0.0,'source_codes':set()})
        group['weekly_lessons']+=1;group['weekly_minutes']+=duration;group['reference_weekly_cost']+=cost;group['source_codes'].update(row.get('source_class_codes') or [])
        included_occurrences.append({'professor':teacher,'professor_id':row.get('professor_id'),'class_name':room['name'],'segment':segment,'discipline':row.get('disciplina_componente'),'day_of_week':row.get('dia_semana'),'start_time':row.get('hora_inicio'),'end_time':row.get('hora_fim'),'duration_minutes':duration,'source_codes':row.get('source_class_codes') or [],'shared_class':False,'hour_aula_reajustado':rate,'reference_weekly_cost':cost,'evidence':row.get('evidencia_original')})
    classes_out=[]
    for item in class_rows.values():
        if item['weekly_lessons'] and not item['pending_occurrences']:item['coverage_status']='CONCILIADO'
        elif item['weekly_lessons']:item['coverage_status']='PARCIAL – pendências de rateio/vínculo'
        else:item['coverage_status']='PENDENTE'
        item['teachers']=sorted(item['teachers']);item['disciplines']=sorted(item['disciplines']);item['source_codes']=sorted(item['source_codes']);item['pending_reasons']=sorted(item['pending_reasons']);item['reference_weekly_cost']=round(item['reference_weekly_cost'],2);classes_out.append(item)
    teachers_out=[]
    for item in teacher_rows.values():
        item['classes']=sorted(item['classes']);item['disciplines']=sorted(item['disciplines']);item['reference_weekly_cost']=round(item['reference_weekly_cost'],2);item['clock_hours_weekly']=round(item['weekly_minutes']/60,2)
        groups=[]
        for group in item.pop('by_class_discipline').values():group['source_codes']=sorted(group['source_codes']);group['reference_weekly_cost']=round(group['reference_weekly_cost'],2);groups.append(group)
        item['by_class_discipline']=sorted(groups,key=lambda group:(group['class_name'],group['discipline']));teachers_out.append(item)
    segments_out=[]
    for item in segment_rows.values():item['teachers']=sorted(item['teachers']);item['classes']=sorted(item['classes']);item['reference_weekly_cost']=round(item['reference_weekly_cost'],2);item['clock_hours_weekly']=round(item['weekly_minutes']/60,2);segments_out.append(item)
    total=round(sum(item['reference_weekly_cost'] for item in segments_out),2)
    coverage={'operational_classes':len(classes_out),'with_confirmed_weekly_cost':sum(item['weekly_lessons']>0 for item in classes_out),'fully_conciliated':sum(item['coverage_status']=='CONCILIADO' for item in classes_out),'partial':sum(item['coverage_status'].startswith('PARCIAL') for item in classes_out),'pending':sum(item['coverage_status']=='PENDENTE' for item in classes_out),'pending_classes':[{'class_name':item['class_name'],'reason':'; '.join(item['pending_reasons']) or 'Nenhuma ocorrência com vínculo inequívoco foi encontrada na grade.'} for item in classes_out if item['coverage_status']=='PENDENTE']}
    return {'status':'AUDIT_PREVIEW','source':'carga_horaria.pdf (Base 2026 – projeção provisória para 2027)','rates_source':'Orçamento 2027 · Valores de Hora Aula · coluna Reajustado','rates':copy.deepcopy(TEACHING_HOUR_AULA_REAJUSTED_2027),'calculation_basis':'Ocorrência de aula da grade × valor de hora-aula reajustado do segmento; duração real preservada sem conversão entre 40/45/50 minutos.','monthly_cost_status':'PENDENTE – o Orçamento 2027 não documenta multiplicador mensal, DSR, hora-atividade ou conversão de duração.','professors':sorted(teachers_out,key=lambda item:item['professor']),'classes':sorted(classes_out,key=lambda item:item['class_name']),'segments':segments_out,'included_occurrences':included_occurrences,'pending_occurrences':pending,'summary':{'real_occurrences':len(preview.get('records',[])),'included_unshared_confirmed_occurrences':included,'pending_or_excluded_occurrences':len(pending),'shared_lessons_pending':sum(item['shared_lessons_pending'] for item in classes_out),'reference_weekly_cost':total,'monthly_cost':None,'coverage':coverage}}

TEACHING_SHARED_RATEIO_2027={'criterio':'RATEIO_IGUALITARIO','origem':'DECISAO_ADMINISTRATIVA','finalidade':'CUSTO_POR_TURMA_PE'}

def _shared_rateio_target(code,classes,by_id):
    """Retorna somente uma turma operacional existente; códigos sem turma ficam pendentes."""
    normalized=normalize_teaching_class_code(code,classes)
    if normalized.get('status')=='CONCILIADO':return by_id.get(normalized.get('operational_class_id'))
    return None

def _shared_rate(code):
    segment=None
    if code.startswith('EINFA'):segment='Educação Infantil'
    match=re.fullmatch(r'EFUND(\d{2})[MT][A-Z]',code)
    if match and 1<=int(match.group(1))<=9:
        segment='Fundamental I' if int(match.group(1))<=5 else 'Fundamental II'
    if re.fullmatch(r'(?:EMERE|EMICN|EMICH)0[1-3][MT][A-Z]',code):segment='Ensino Médio'
    if segment:return segment,TEACHING_HOUR_AULA_REAJUSTED_2027[segment]
    return None,None

def build_teaching_weekly_cost_with_shared_rateio_audit(preview,classes):
    """Aplica somente a alocação igualitária aprovada às aulas compartilhadas.

    A carga e o custo do professor entram uma vez. As parcelas por turma, já
    arredondadas em centavos, fecham exatamente o valor da ocorrência. Um
    código sem turma operacional continua em uma parcela pendente, sem criar
    turma fictícia e sem perder o fechamento auditável.
    """
    result=copy.deepcopy(build_teaching_weekly_cost_audit(preview,classes));by_id={item['id']:item for item in classes}
    # Retira apenas as pendências provisórias de compartilhamento. A carga do
    # professor, contabilizada uma vez pela auditoria base, permanece intacta.
    unshared=build_teaching_weekly_cost_audit({'records':[row for row in preview.get('records',[]) if not row.get('aula_compartilhada')]},classes)
    result['classes']=unshared['classes']
    result['pending_occurrences']=unshared['pending_occurrences']
    unshared_teachers={item['professor']:item for item in unshared['professors']}
    for item in result['professors']:
        item['pending_occurrences']=unshared_teachers.get(item['professor'],{}).get('pending_occurrences',0)
    class_rows={item['class_id']:item for item in result['classes']};teacher_rows={item['professor']:item for item in result['professors']};segment_rows={item['segment']:item for item in result['segments']}
    for item in class_rows.values():item['shared_rateio_weekly_cost']=0.0
    for item in teacher_rows.values():item['shared_rateio_weekly_cost']=0.0
    allocations=[];pending_code_allocations=[];shared_total=0.0;class_total=0.0;pending_total=0.0
    shared_pending=0;pending_links=0;unallocated=[];minutes_total=0.0;class_minutes=0.0;pending_minutes=0.0
    def pending_reason(code):
        if code.startswith('EINFA'):return 'EINFA: gramática documental não comprovada.'
        if code.startswith(('EMERE','EMICN','EMICH')):return 'Ensino Médio: seção sem correspondência inequívoca na turma operacional agregada.'
        if _shared_rate(code)==(None,None):return 'Código desconhecido ou inválido; tarifa não identificável.'
        return 'Código sem turma operacional conciliada.'
    for occurrence_index,row in enumerate(preview.get('records',[])):
        if not row.get('aula_compartilhada'):continue
        codes=sorted(set(row.get('source_class_codes') or []));rates={_shared_rate(code) for code in codes}
        teacher=row.get('professor') or 'PENDENTE – professor não identificado'
        duration=row.get('duracao_minutos')
        metadata={'occurrence_index':occurrence_index,'professor':teacher,'professor_id':row.get('professor_id'),'page':row.get('pagina'),'source_codes':codes,'discipline':row.get('disciplina_componente'),'day_of_week':row.get('dia_semana'),'start_time':row.get('hora_inicio'),'end_time':row.get('hora_fim'),'duration_minutes':duration,'shared_class':True,'evidence':row.get('evidencia_original')}
        targets={code:_shared_rateio_target(code,classes,by_id) for code in codes}
        pending_links+=sum(target is None for target in targets.values())
        reason=None
        if row.get('status_temporal')!='CONFIRMADO' or duration not in (40,45,50):reason='Horário/duração não confirmados; ocorrência mantida pendente sem alocação.'
        elif not row.get('professor'):reason='Professor não identificado; ocorrência mantida pendente sem alocação.'
        elif not codes or (None,None) in rates:reason='Código desconhecido ou inválido; valor de hora-aula não identificável para todos os participantes.'
        elif len(rates)!=1:reason='Participantes com segmentos/tarifas diferentes; custo integral não comprovado.'
        target_ids=[target['id'] for target in targets.values() if target]
        if not reason and len(set(target_ids))!=len(target_ids):reason='Códigos diferentes apontam para a mesma turma; participantes não inequívocos.'
        pending_reasons=[reason] if reason else sorted({pending_reason(code) for code,target in targets.items() if not target})
        if pending_reasons:
            shared_pending+=1
            result['pending_occurrences'].append({**metadata,'reason':'; '.join(pending_reasons)})
            teacher_rows[teacher]['pending_occurrences']+=1
            for target_id in set(target_ids):
                item=class_rows[target_id];item['pending_occurrences']+=1;item['shared_lessons_pending']+=1
                item['pending_reasons']=sorted(set(item['pending_reasons'])|set(pending_reasons))
        if reason:
            entry={**metadata,'reason':reason,'status':'PENDENTE','allocations':[]}
            pending_code_allocations.append(entry);unallocated.append(entry)
            continue
        segment,rate=next(iter(rates));cents=round(rate*100);base,remainder=divmod(cents,len(codes))
        teacher_item=teacher_rows.get(teacher)
        if teacher_item:teacher_item['reference_weekly_cost']+=rate;teacher_item['shared_rateio_weekly_cost']+=rate
        segment_item=segment_rows.get(segment)
        if segment_item:
            segment_item['reference_weekly_cost']+=rate;segment_item['weekly_minutes']+=duration;segment_item['weekly_lessons']+=1
        shared_total+=rate;minutes_total+=duration;occurrence=[]
        for index,code in enumerate(codes):
            amount=(base+(1 if index<remainder else 0))/100;target=targets[code]
            minutes=duration/len(codes)
            allocation={'source_code':code,'operational_class_id':target['id'] if target else None,'operational_class_name':target['name'] if target else None,'amount':amount,'fraction':1/len(codes),'allocated_minutes':minutes,'status':'CONCILIADO' if target else 'PENDENTE','reason':None if target else pending_reason(code),**TEACHING_SHARED_RATEIO_2027}
            occurrence.append(allocation)
            if target:
                item=class_rows[target['id']];item['reference_weekly_cost']+=amount;item['shared_rateio_weekly_cost']+=amount;item['teachers']=sorted(set(item['teachers'])|{teacher});item['disciplines']=sorted(set(item['disciplines'])|{row.get('disciplina_componente') or 'PENDENTE'});class_total+=amount
                item['weekly_minutes']+=minutes;item['weekly_lessons']+=1/len(codes);item['source_codes']=sorted(set(item['source_codes'])|{code});class_minutes+=minutes
            else:
                pending_total+=amount
                pending_minutes+=minutes
                pending_code_allocations.append({**metadata,**allocation})
        if round(sum(item['amount'] for item in occurrence),2)!=round(rate,2):raise ValueError('Rateio compartilhado não fecha o custo da ocorrência.')
        if abs(sum(item['allocated_minutes'] for item in occurrence)-duration)>1e-9:raise ValueError('Rateio compartilhado não conserva minutos.')
        if abs(sum(item['fraction'] for item in occurrence)-1)>1e-12:raise ValueError('Rateio compartilhado não fecha 100%.')
        allocations.append({**metadata,'original_cost':rate,'allocations':occurrence,**TEACHING_SHARED_RATEIO_2027})
    for item in result['classes']:item['reference_weekly_cost']=round(item['reference_weekly_cost'],2);item['shared_rateio_weekly_cost']=round(item['shared_rateio_weekly_cost'],2)
    for item in result['professors']:item['reference_weekly_cost']=round(item['reference_weekly_cost'],2);item['shared_rateio_weekly_cost']=round(item['shared_rateio_weekly_cost'],2)
    for item in result['segments']:item['reference_weekly_cost']=round(item['reference_weekly_cost'],2);item['clock_hours_weekly']=round(item['weekly_minutes']/60,2)
    difference=round(shared_total-class_total-pending_total,2);difference=0.0 if abs(difference)<.005 else difference
    result['shared_lesson_rateio']={**TEACHING_SHARED_RATEIO_2027,'occurrences':len(allocations),'allocations':allocations,'pending_code_allocations':pending_code_allocations,'original_cost_total':round(shared_total,2),'allocated_to_operational_classes':round(class_total,2),'allocated_pending_code':round(pending_total,2),'closure_difference':difference,'teacher_cost_counted_once':True}
    result['shared_lesson_rateio'].update(unallocated_occurrences=unallocated,pending_occurrences=shared_pending,pending_links=pending_links,original_minutes_total=minutes_total,allocated_minutes_to_operational_classes=class_minutes,allocated_minutes_pending_code=pending_minutes,minutes_closure_difference=minutes_total-class_minutes-pending_minutes,minutes_tolerance=1e-9,fraction_tolerance=1e-12)
    result['summary']['shared_lessons_allocated']=len(allocations);result['summary']['shared_lessons_pending']=shared_pending;result['summary']['pending_or_excluded_occurrences']=len(result['pending_occurrences']);result['summary']['reference_weekly_cost']=round(result['summary']['reference_weekly_cost']+shared_total,2);result['summary']['rateio_closure_difference']=difference
    for item in result['classes']:
        item['coverage_status']=('PARCIAL – pendências de rateio/vínculo' if item['pending_occurrences'] else 'CONCILIADO') if item['weekly_lessons'] else 'PENDENTE'
    result['summary']['coverage']={'operational_classes':len(result['classes']),'with_confirmed_weekly_cost':sum(item['weekly_lessons']>0 for item in result['classes']),'fully_conciliated':sum(item['coverage_status']=='CONCILIADO' for item in result['classes']),'partial':sum(item['coverage_status'].startswith('PARCIAL') for item in result['classes']),'pending':sum(item['coverage_status']=='PENDENTE' for item in result['classes']),'pending_classes':[{'class_name':item['class_name'],'reason':'; '.join(item['pending_reasons']) or 'Nenhuma ocorrência com vínculo inequívoco foi encontrada na grade.'} for item in result['classes'] if item['coverage_status']=='PENDENTE']}
    return result

def teaching_load_weekly_cost_audit(content,classes):
    """Executa a prévia geométrica e a alocação compartilhada aprovada, sem persistir dados."""
    return build_teaching_weekly_cost_with_shared_rateio_audit(teaching_load_geometric_preview(content,classes),classes)
def teaching_load_preview(name,content,state,target_year=2027):
    if Path(name).suffix.lower()!='.pdf':raise ValueError('Nesta etapa a carga horária aceita somente PDF textual.')
    from pypdf import PdfReader
    digest=hashlib.sha256(content).hexdigest()
    if any(v.get('source_hash')==digest for v in state.get('teachingLoad',{}).get('versions',[])):raise ValueError('Arquivo já processado.')
    pages=[p.extract_text(extraction_mode='layout') or '' for p in PdfReader(io.BytesIO(content)).pages];rooms={fold(x['name']):x for x in state['classes']};records=[];warnings=[];professor='';professor_id='';source_year=2026
    code_re=re.compile(r'\b((?:EINFA|EFUND|EMERE|EMICN|EMICH)[A-Z0-9]+)\s*-\s*([^|\n]+)');dur_re=re.compile(r'(\d{2}:\d{2})\s*-\s*(\d{2,3})')
    for page_no,text in enumerate(pages,1):
        found=re.search(r'Professor:\s*(.+?)\s*\(([^)]+)\)',text);period=re.search(r'Período letivo:\s*(20\d{2})',text)
        if found:professor,professor_id=found.group(1).strip(),found.group(2).strip()
        if period:source_year=int(period.group(1))
        for line in text.splitlines():
            codes=code_re.findall(line)
            if not codes:continue
            shared='|' in line;duration=next((int(x[1]) for x in dur_re.findall(text[text.find(line):text.find(line)+500])),None)
            for code,discipline in codes:
                normalized=normalize_teaching_class_code(code,state['classes']);records.append({'id':secrets.token_hex(10),'professor':professor,'professor_id':professor_id or None,**normalized,'discipline':discipline.strip(),'duration_minutes':duration,'shared_class':shared,'shared_note':'Aula compartilhada – regra de distribuição pendente' if shared else None,'source_text':line.strip(),'reconciliation_status':normalized['status'],'page':page_no})
    if not records:warnings.append('Nenhum registro de aula foi extraído com segurança.')
    codes={r['source_class_code'] for r in records};mapped={r['source_class_code'] for r in records if r['operational_class_id']};return {'version_id':secrets.token_hex(12),'source_year':source_year,'target_year':int(target_year),'source_filename':name,'source_hash':digest,'status':'PREVIEW','notes':'Base 2026 – projeção provisória para 2027','records':records,'class_mappings':[{'source_class_code':c,'operational_class_id':next((r['operational_class_id'] for r in records if r['source_class_code']==c),None),'status':next((r['reconciliation_status'] for r in records if r['source_class_code']==c),'PENDENTE')} for c in sorted(codes)],'warnings':warnings,'professors':len({r['professor'] for r in records if r['professor']}),'codes':len(codes),'conciliated_codes':len(mapped),'pending_codes':len(codes-mapped),'shared_lessons':sum(r['shared_class'] for r in records),'errors':[]}
def decimal(value):
    if isinstance(value,(int,float)):return float(value)
    text=str(value or '').strip().replace('R$','').replace('%','').replace(' ','');negative=text.startswith('(') and text.endswith(')');text=text[1:-1] if negative else text
    if not text:raise ValueError('Valor ausente.')
    if ',' in text:text=text.replace('.','').replace(',','.')
    if not re.fullmatch(r'-?\d+(\.\d+)?',text):raise ValueError('Valor numérico não reconhecido: '+str(value))
    return -float(text) if negative else float(text)
def date_value(value):
    if isinstance(value,datetime):return value.date().isoformat()
    for pattern in ['%Y-%m-%d','%d/%m/%Y','%d-%m-%Y']:
        try:return datetime.strptime(str(value).split('T')[0],pattern).date().isoformat()
        except ValueError:pass
    raise ValueError('Data não reconhecida: '+str(value))
def read_rows(name,content,kind):
    ext=Path(name).suffix.lower()
    if ext=='.csv':
        try:text=content.decode('utf-8-sig')
        except UnicodeDecodeError:text=content.decode('cp1252')
        try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=';,\t')
        except csv.Error:dialect=csv.excel;dialect.delimiter=';'
        return list(csv.reader(io.StringIO(text),dialect)),None
    if ext=='.xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist())>40*1024*1024:raise ValueError('Planilha descompactada excede 40 MB.')
        book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        if kind in ['budget','expenses']:
            if 'DESPESAS 2026' not in book.sheetnames:raise ValueError('A aba DESPESAS 2026 não foi encontrada. As demais abas não serão importadas.')
            sheet=book['DESPESAS 2026']
        else:
            if len(book.sheetnames)!=1:raise ValueError('Relatório de inadimplência deve conter uma única aba para evitar ambiguidade.')
            sheet=book.active
        rows=[[cell.value*100 if isinstance(cell.value,(int,float)) and '%' in cell.number_format else cell.value for cell in row] for row in sheet.iter_rows()];title=sheet.title;book.close();return rows,title
    if ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(content))
        if len(reader.pages)>150:raise ValueError('PDF excede 150 páginas.')
        text='\n'.join(page.extract_text(extraction_mode='layout') or '' for page in reader.pages)
        if not text.strip():raise ValueError('PDF sem texto selecionável. OCR não configurado; forneça CSV/Excel ou um PDF textual.')
        # Deliberadamente conservador: tabelas textuais separadas por colunas; nenhum palpite monetário silencioso.
        lines=[re.split(r'\s{2,}|\t|;',line.strip()) for line in text.splitlines() if line.strip()]
        return lines,None
    if ext=='.xls':raise ValueError('Excel .xls antigo ainda não suportado. Salve como .xlsx; nenhum dado foi importado.')
    raise ValueError('Formato permitido: XLSX, CSV ou PDF textual.')

def financial_import_records(name,content):
    """Lê uma ou as duas modalidades do modelo, ignorando apenas a aba de instruções."""
    if Path(name).suffix.lower()!='.xlsx':
        raw,sheet=read_rows(name,content,'financial_classifications');records,_=mapped(raw,aliases=FINANCIAL_IMPORT_ALIASES);return records,sheet
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist())>40*1024*1024:raise ValueError('Planilha descompactada excede 40 MB.')
    book=load_workbook(io.BytesIO(content),read_only=True,data_only=True);records=[];used=[];errors=[]
    for sheet in book.worksheets:
        if fold(sheet.title)=='instrucoes':continue
        rows=[[cell.value*100 if isinstance(cell.value,(int,float)) and '%' in cell.number_format else cell.value for cell in row] for row in sheet.iter_rows()]
        if not any(any(value not in [None,''] for value in row) for row in rows):continue
        try:items,_=mapped(rows,aliases=FINANCIAL_IMPORT_ALIASES);records.extend(items);used.append(sheet.title)
        except ValueError as error:errors.append(sheet.title+': '+str(error))
    book.close()
    if not records:raise ValueError('Nenhuma aba nominal ou consolidada reconhecida. '+(' '.join(errors) if errors else ''))
    return records,', '.join(used)

def financial_import_workbook(homologation=False):
    """Modelo seguro sem fórmulas/macros; a variante de homologação contém somente nomes fictícios."""
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    book=Workbook();nominal=book.active;nominal.title='Nominal';consolidated=book.create_sheet('Consolidada por turma');instructions=book.create_sheet('Instruções')
    nominal.append(['Aluno','Turma','Série/Ano','Categoria/Benefício','Percentual','Observação'])
    consolidated.append(['Turma','Série/Ano','Categoria/Benefício','Percentual','Quantidade','Observação'])
    if homologation:
        nominal_rows=[['ALUNO FICTÍCIO 001','G2 A','Grupo 2','Sem desconto',0,'TESTE — não confirmar'],['ALUNO FICTÍCIO 002','G2 A','Grupo 2','Desconto variável',16,'TESTE — faixa variável'],['ALUNO FICTÍCIO 003','G3 A','Grupo 3','Projeto Marcando Vidas',45,'TESTE'],['ALUNO FICTÍCIO 004','G3 A','Grupo 3','Bolsa filantrópica',50,'TESTE'],['ALUNO FICTÍCIO 005','1º A','1º ano','Filho de funcionário',100,'TESTE'],['ALUNO FICTÍCIO 006','TURMA FICTÍCIA INEXISTENTE','—','Sem desconto',0,'TESTE — turma inexistente'],['ALUNO FICTÍCIO 007','','—','Sem desconto',0,'TESTE — turma ausente'],['ALUNO FICTÍCIO 008','G2 A','Grupo 2','Categoria ambígua',35,'TESTE — não presumir']]
        for row in nominal_rows:nominal.append(row)
        for row in [['G2 A','Grupo 2','Bolsa filantrópica',100,999,'TESTE — divergência proposital']]:consolidated.append(row)
    instructions.append(['MODELO 7&7 — BOLSAS, BENEFÍCIOS E DESCONTOS']);instructions.append(['Preencha a aba Nominal (uma linha por aluno), a aba Consolidada por turma (uma linha por combinação turma/desconto), ou ambas.']);instructions.append(['Turma é obrigatória. Registros sem turma ou ambíguos ficam pendentes; nunca são distribuídos entre turmas.']);instructions.append(['Categorias reconhecidas']);
    for label in ['Sem desconto — 0%','Bolsa filantrópica — 100%','Bolsa filantrópica — 50%','Filho de funcionário — 100%','Filho de funcionário — 80%','Filho de obreiro — 100%','Projeto Marcando Vidas — 45%','Outros descontos variáveis — informe um percentual entre 0% e 100%']:instructions.append([label])
    for sheet in [nominal,consolidated,instructions]:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='075EAF');cell.alignment=Alignment(wrap_text=True)
        for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=min(55,max(16,max(len(str(cell.value or '')) for cell in column)+2))
    stream=io.BytesIO();book.save(stream);book.close();return stream.getvalue()
def mapped(rows,mapping=None,aliases=ALIASES):
    for idx,row in enumerate(rows[:80]):
        headers=[fold(x) for x in row]
        matches={field:next((i for i,h in enumerate(headers) if h in [fold(alias) for alias in names]),None) for field,names in aliases.items()}
        if mapping:matches.update({k:int(v) for k,v in mapping.items() if v is not None and v!=''})
        if sum(v is not None for v in matches.values())>=2:
            data=[{k:(r[i] if i<len(r) else None) for k,i in matches.items() if i is not None} for r in rows[idx+1:] if any(x not in [None,''] for x in r)]
            return data,headers
    raise ValueError('Cabeçalho não reconhecido. Use colunas Código, Categoria, Orçado ou Data, Descrição, Valor; inadimplência: Mês, Financeira (%), Contábil (%), Dívida, Alunos, Responsáveis. PDFs devem possuir colunas textuais legíveis.')

def _financial_category(category,percent):
    label=fold(category).replace('bolsista ','bolsa ').replace('colaborador','funcionario').replace('funcionário','funcionario')
    aliases={'filantropica':'bolsa filantropica','filantropico':'bolsa filantropica','filantropica integral':'bolsa filantropica','filantropica parcial':'bolsa filantropica','filho funcionario':'filho de funcionario','filho obreiro':'filho de obreiro','marcando vidas':'projeto marcando vidas','sem beneficio':'sem desconto','sem benefício':'sem desconto','integral':'sem desconto'}
    label=aliases.get(label,label)
    if (label,percent) in FINANCIAL_FIXED:return FINANCIAL_FIXED[(label,percent)]
    if label in ['desconto variavel','outros descontos variaveis','outro desconto','desconto'] or (not label and 0<percent<100):
        return 'variable'
    return None

def _financial_percent(category,value):
    if value not in [None,'']:return decimal(value)
    label=fold(category).replace('bolsista ','bolsa ').replace('colaborador','funcionario').replace('funcionário','funcionario')
    label={'filantropica':'bolsa filantropica','filantropico':'bolsa filantropica','filantropica integral':'bolsa filantropica','filantropica parcial':'bolsa filantropica','filho funcionario':'filho de funcionario','filho obreiro':'filho de obreiro','marcando vidas':'projeto marcando vidas','sem beneficio':'sem desconto','sem benefício':'sem desconto','integral':'sem desconto'}.get(label,label)
    candidates={percent for (name,percent),_ in FINANCIAL_FIXED.items() if name==label}
    if len(candidates)==1:return float(next(iter(candidates)))
    raise ValueError('Percentual ausente ou ambíguo para a categoria informada.')

def _financial_preview_summary(state,year,rows,invalid):
    """Projeta a importação nas 41 turmas sem alterar o estado persistido."""
    from server import FINANCIAL_CATEGORIES
    rooms={str(room['id']):room for room in state.get('classes',[])};academic=next(item for item in state.get('academicYears',[]) if int(item['year'])==int(year));pending_by_class={};unassigned=0
    imported={}
    for item in rows:
        if item.get('status')=='ready':
            marker=(str(item['classId']),item['classificationKey'],float(item['percent']))
            imported[marker]=imported.get(marker,0)+int(item['quantity'])
        elif item.get('classId'):pending_by_class[str(item['classId'])]=pending_by_class.get(str(item['classId']),0)+1
        else:unassigned+=1
    summary=[]
    for room in state.get('classes',[]):
        class_id=str(room['id']);current=copy.deepcopy(academic.get('classifications',{}).get(class_id,{}));fixed={key:int((current.get('fixed') or {}).get(key,0)) for key,_,_ in FINANCIAL_CATEGORIES};variables={float(row.get('percent',0)):int(row.get('quantity',0)) for row in current.get('variables',[])}
        for (target,key,percent),quantity in imported.items():
            if target!=class_id:continue
            if key=='variable':variables[percent]=quantity
            else:fixed[key]=quantity
        classified=sum(fixed.values())+sum(variables.values());enrolled=int(room.get('students',0));divergence=max(0,classified-enrolled)
        summary.append({'classId':class_id,'className':room['name'],'enrolled':enrolled,'fixed':fixed,'variables':[{'percent':percent,'quantity':quantity} for percent,quantity in sorted(variables.items()) if quantity],'classified':classified,'unclassified':max(0,enrolled-classified),'pending':pending_by_class.get(class_id,0),'duplicates':sum(1 for item in rows if str(item.get('classId'))==class_id and item.get('status')=='duplicate'),'unrecognized':pending_by_class.get(class_id,0),'divergence':divergence,'status':'blocked' if divergence else 'complete' if classified==enrolled else 'pending'})
    return summary,unassigned+len(invalid),not any(item['divergence'] for item in summary)

def financial_classification_preview(name,content,state,year):
    """Prévia conservadora de bolsas/descontos; nunca altera o estado recebido."""
    records,sheet=financial_import_records(name,content)
    target_year=next((item for item in state.get('academicYears',[]) if int(item.get('year',0))==int(year)),None)
    if not target_year:raise ValueError('Ano letivo financeiro não encontrado.')
    rooms={fold(item['name']):item for item in state.get('classes',[])}
    result={'name':name,'kind':'financial_classifications','year':int(year),'sheet':sheet,'rows':[],'invalid':[],'duplicates':0,'automatic':0,'pending':0,'fileHash':hashlib.sha256(content).hexdigest()}
    if any(item.get('fileHash')==result['fileHash'] and item.get('kind')==result['kind'] for item in state.get('imports',[])):raise ValueError('Este arquivo já foi confirmado para descontos e benefícios.')
    seen=set()
    for line,raw_item in enumerate(records,2):
        item={'id':secrets.token_hex(12),'line':line,'className':str(raw_item.get('className') or '').strip(),'grade':str(raw_item.get('grade') or '').strip(),'category':str(raw_item.get('category') or '').strip(),'student':str(raw_item.get('student') or '').strip(),'source':{key:str(value or '').strip() for key,value in raw_item.items()}}
        try:
            item['percent']=_financial_percent(item['category'],raw_item.get('percent'));quantity_raw=raw_item.get('quantity');item['quantity']=1 if item['student'] and quantity_raw in [None,''] else decimal(quantity_raw)
            if not 0<=item['percent']<=100:raise ValueError('Percentual fora de 0% a 100%.')
            if item['quantity']<0 or not item['quantity'].is_integer():raise ValueError('Quantidade deve ser inteiro não negativo.')
            item['quantity']=int(item['quantity']);room=rooms.get(fold(item['className']));item['classId']=room['id'] if room else None;item['classificationKey']=_financial_category(item['category'],item['percent'])
            if not room:
                item.update(status='pending_review',reason='Turma inexistente ou não identificada exatamente.')
            elif not item['classificationKey']:
                item.update(status='pending_review',reason='Tipo de benefício/desconto e percentual não correspondem a uma faixa configurada.')
            else:
                signature=(item['classId'],fold(item['student']),item['classificationKey'],item['percent']) if item['student'] else (item['classId'],item['classificationKey'],item['percent'])
                if signature in seen:
                    item.update(status='duplicate',reason='Duplicidade dentro do arquivo.');result['duplicates']+=1
                else:
                    seen.add(signature)
                    item.update(status='ready',reason='Aluno, turma e faixa reconhecidos.' if item['student'] else 'Turma e faixa reconhecidas exatamente.');result['automatic']+=1
            if item['status']=='pending_review':result['pending']+=1
            result['rows'].append(item)
        except (ValueError,TypeError) as error:
            result['invalid'].append({'line':line,'error':str(error),'source':item['source']})
    if not result['rows'] and not result['invalid']:raise ValueError('Nenhuma linha encontrada no arquivo.')
    result['classSummary'],result['unrecognized'],result['canConfirm']=_financial_preview_summary(state,year,result['rows'],result['invalid'])
    return result

def apply_financial_classification_preview(state,preview,decisions,user):
    """Aplica somente decisões explícitas e válidas, mantendo o arquivo/auditoria."""
    from server import FINANCIAL_CATEGORIES
    if not user.get('isAdmin'):raise PermissionError('Somente administradores podem aprovar importações financeiras.')
    updated=copy.deepcopy(state);rows={item['id']:item for item in preview['rows']};rooms={str(item['id']):item for item in updated['classes']};year=next((item for item in updated['academicYears'] if item['year']==preview['year']),None)
    if not year:raise ValueError('Ano letivo não encontrado ao confirmar.')
    selections=[]
    for decision in decisions:
        original=rows.get(str(decision.get('id','')))
        if not original:raise ValueError('Item da prévia não encontrado.')
        action=decision.get('action')
        if action not in ['accept','edit','ignore']:raise ValueError('Decisão de prévia inválida.')
        reason=str(decision.get('reason') or '').strip()
        if action in ['edit','ignore'] and not reason:raise ValueError('Editar ou ignorar exige justificativa.')
        if action=='ignore':continue
        if original.get('status')!='ready' and action!='edit':raise ValueError('Item pendente exige edição com justificativa ou deve ser ignorado.')
        class_id=str(decision.get('classId') or original.get('classId') or '');category=str(decision.get('category') or original.get('category') or '');percent=decimal(decision.get('percent',original.get('percent')));quantity=decimal(decision.get('quantity',original.get('quantity')))
        if class_id not in rooms:raise ValueError('Turma selecionada não existe.')
        if not 0<=percent<=100 or quantity<0 or not quantity.is_integer():raise ValueError('Percentual ou quantidade inválidos.')
        key=_financial_category(category,percent)
        if not key:raise ValueError('Faixa de benefício/desconto não reconhecida. Corrija a categoria e o percentual.')
        selections.append({'classId':class_id,'category':category,'key':key,'percent':percent,'quantity':int(quantity),'action':action,'reason':reason,'sourceId':original['id']})
    grouped={}
    for item in selections:
        marker=(item['classId'],item['key'],item['percent'])
        if marker in grouped:grouped[marker]['quantity']+=item['quantity']
        else:grouped[marker]=item
    by_class={}
    for item in grouped.values():by_class.setdefault(item['classId'],[]).append(item)
    for class_id,items in by_class.items():
        current=copy.deepcopy(year['classifications'].get(class_id,{}));fixed=current.get('fixed',{});variables=current.get('variables',[])
        for item in items:
            if item['key']=='variable':
                variables=[x for x in variables if float(x.get('percent',0))!=item['percent']]
                if item['quantity']:variables.append({'id':'import-'+item['sourceId'],'percent':item['percent'],'quantity':item['quantity']})
            else:fixed[item['key']]=item['quantity']
        total=sum(int(fixed.get(key,0)) for key,_,_ in FINANCIAL_CATEGORIES)+sum(int(row.get('quantity',0)) for row in variables)
        if total>int(rooms[class_id]['students']):raise ValueError('Classificação importada excede os matriculados da turma '+rooms[class_id]['name']+'.')
        year['classifications'][class_id]={'fixed':fixed,'variables':variables}
    stamp=now();updated['imports'].append({'id':secrets.token_hex(12),'name':preview['name'],'kind':'financial_classifications','year':preview['year'],'fileHash':preview['fileHash'],'date':stamp,'count':len(selections),'automatic':0,'pending':sum(1 for item in preview['rows'] if item['status']!='ready'),'duplicates':preview['duplicates'],'rejected':len(preview['invalid']),'version':preview.get('version'),'approvedBy':user['name'],'approvedById':user['id'],'decisions':[{key:value for key,value in item.items() if key!='sourceId'} for item in selections]})
    return updated

def revenue_year(state,year):
    layer=state.setdefault('revenuePlanning',{'years':{}});years=layer.setdefault('years',{})
    key=str(int(year))
    return years.setdefault(key,{'year':int(year),'calendar':{'enrollmentMonth':1,'tuitionMonths':list(range(2,13))},'campaigns':[],'imports':[],'records':[]})

def revenue_preview(name,content,state,year):
    """Prévia imutável de receitas por aluno identificado; não aceita nome como chave."""
    raw,sheet=read_rows(name,content,'revenue')
    records,_=mapped(raw,aliases=REVENUE_IMPORT_ALIASES)
    year=int(year);config=revenue_year(copy.deepcopy(state),year);rooms={fold(row['name']):row for row in state.get('classes',[])}
    known={str(row.get('externalId')) for row in config.get('records',[])};known_business={(str(row.get('studentId')),str(row.get('eventType')),str(row.get('date'))) for row in config.get('records',[]) if row.get('eventType') in ['enrollment','tuition']};result={'name':name,'kind':'revenue_consolidated','year':year,'sheet':sheet,'rows':[],'invalid':[],'duplicates':0,'pending':0,'ready':0,'fileHash':hashlib.sha256(content).hexdigest()}
    if any(row.get('fileHash')==result['fileHash'] for row in config.get('imports',[])):raise ValueError('Este arquivo consolidado já foi confirmado para este ano.')
    file_ids=set();file_business=set()
    for line,raw_item in enumerate(records,2):
        item={'id':secrets.token_hex(12),'line':line,'source':{key:str(value or '').strip() for key,value in raw_item.items()}}
        try:
            item.update(studentId=str(raw_item.get('studentId') or '').strip(),externalId=str(raw_item.get('externalId') or '').strip(),className=str(raw_item.get('className') or '').strip(),eventType=fold(raw_item.get('eventType')).replace(' ','_'),status=fold(raw_item.get('status')).replace(' ','_'),reversalOf=str(raw_item.get('reversalOf') or '').strip(),fromClassName=str(raw_item.get('fromClassName') or '').strip())
            aliases={'matricula':'enrollment','matrícula':'enrollment','mensalidade':'tuition','transferencia':'transfer','transferência':'transfer','cancelamento':'cancellation','estorno':'refund','realizada':'realized','realizado':'realized','a_receber':'receivable','projetada':'projected','projetado':'projected'}
            item['eventType']=aliases.get(item['eventType'],item['eventType']);item['status']=aliases.get(item['status'],item['status']);item['date']=date_value(raw_item.get('date'))
            if not item['studentId']:raise ValueError('Identificador único do aluno é obrigatório; nome não é aceito como chave.')
            if not item['externalId']:raise ValueError('Identificador único do lançamento financeiro é obrigatório.')
            if item['eventType'] not in ['enrollment','tuition','transfer','cancellation','refund']:raise ValueError('Tipo deve ser matrícula, mensalidade, transferência, cancelamento ou estorno.')
            if item['status'] not in ['realized','receivable','projected']:raise ValueError('Status deve ser realizada, a receber ou projetada.')
            room=rooms.get(fold(item['className']));item['classId']=room['id'] if room else None
            month=int(item['date'][5:7])
            if item['eventType']=='enrollment' and month!=1:raise ValueError('Matrícula pertence exclusivamente a janeiro; não é 12ª mensalidade.')
            if item['eventType']=='tuition' and month not in range(2,13):raise ValueError('Mensalidade deve pertencer ao período de fevereiro a dezembro.')
            item['gross']=decimal(raw_item.get('gross',0));provided_discount=decimal(raw_item.get('discount',0)) if raw_item.get('discount') not in [None,''] else 0.0
            item['settledAmount']=decimal(raw_item.get('settledAmount')) if raw_item.get('settledAmount') not in [None,''] else None
            if item['gross']<0 or provided_discount<0 or provided_discount>item['gross']:raise ValueError('Valor bruto/desconto inválido.')
            campaign=None
            if item['eventType']=='enrollment':
                when=datetime.strptime(item['date'],'%Y-%m-%d').date()
                matches=[c for c in config.get('campaigns',[]) if c.get('appliesEnrollment') and datetime.strptime(c['startDate'],'%Y-%m-%d').date()<=when<=datetime.strptime(c['endDate'],'%Y-%m-%d').date() and item['studentId'] not in c.get('exceptions',[]) and (not c.get('classExceptions') or item['classId'] not in c.get('classExceptions',[]))]
                if len(matches)>1:raise ValueError('Mais de uma campanha aplicável; exige conferência.')
                campaign=matches[0] if matches else None
            calculated_discount=round(item['gross']*float(campaign['discountPercent'])/100,2) if campaign else provided_discount
            item.update(campaignId=campaign.get('id') if campaign else None,campaignName=campaign.get('name') if campaign else None,discount=calculated_discount,net=round(item['gross']-calculated_discount,2),direction=-1 if item['eventType'] in ['cancellation','refund'] else 0 if item['eventType']=='transfer' else 1)
            if not room and item['eventType']!='refund':item.update(statusReview='pending_review',reason='Turma inexistente ou sem correspondência exata.')
            elif item['externalId'] in known or item['externalId'] in file_ids or item['eventType'] in ['enrollment','tuition'] and (item['studentId'],item['eventType'],item['date']) in known_business|file_business:item.update(statusReview='duplicate',reason='Lançamento financeiro ou evento do aluno já existe; não será somado novamente.');result['duplicates']+=1
            elif item['eventType'] in ['cancellation','refund'] and item['reversalOf'] not in known and item['reversalOf'] not in file_ids:item.update(statusReview='pending_review',reason='Cancelamento/estorno exige identificador do lançamento original.')
            elif campaign and provided_discount not in [0,calculated_discount]:item.update(statusReview='pending_review',reason='Desconto informado diverge da campanha aplicável.')
            else:item.update(statusReview='ready',reason='Identificadores, competência e valores validados.');result['ready']+=1
            item['financialDivergence']=None if item['settledAmount'] is None else round(item['net']-item['settledAmount'],2)
            if item['statusReview']=='pending_review':result['pending']+=1
            file_ids.add(item['externalId']);
            if item['eventType'] in ['enrollment','tuition']:file_business.add((item['studentId'],item['eventType'],item['date']))
            result['rows'].append(item)
        except (ValueError,TypeError) as error:result['invalid'].append({'line':line,'error':str(error),'source':item['source']})
    if not result['rows'] and not result['invalid']:raise ValueError('Nenhuma linha encontrada no arquivo consolidado.')
    return result

def apply_revenue_preview(state,preview,decisions,user):
    if not user.get('isAdmin'):raise PermissionError('Somente administradores podem aprovar receitas consolidadas.')
    updated=copy.deepcopy(state);config=revenue_year(updated,preview['year']);rows={row['id']:row for row in preview['rows']};approved=[]
    for decision in decisions:
        original=rows.get(str(decision.get('id','')))
        if not original:raise ValueError('Linha da prévia não encontrada.')
        action=decision.get('action');reason=str(decision.get('reason') or '').strip()
        if action not in ['accept','edit','ignore'] or action in ['edit','ignore'] and not reason:raise ValueError('Editar ou ignorar exige justificativa.')
        if action=='ignore':continue
        if original['statusReview']!='ready':raise ValueError('Linha pendente ou duplicada não pode ser aplicada sem nova prévia corrigida.')
        approved.append({**original,'decision':action,'decisionReason':reason})
    used={str(row.get('externalId')) for row in config['records']}
    if any(row['externalId'] in used for row in approved):raise ValueError('Uma receita já foi confirmada por este identificador financeiro.')
    stamp=now();version=len(config['imports'])+1
    for row in approved:
        record={key:value for key,value in row.items() if key not in ['source','statusReview','reason']};record.update(id=secrets.token_hex(12),importVersion=version,importedAt=stamp,importedBy=user['name'],importedById=user['id'])
        config['records'].append(record)
    config['imports'].append({'id':secrets.token_hex(12),'version':version,'name':preview['name'],'fileHash':preview['fileHash'],'date':stamp,'approvedBy':user['name'],'approvedById':user['id'],'count':len(approved),'pending':preview['pending'],'duplicates':preview['duplicates'],'invalid':len(preview['invalid']),'decisions':[{'id':x.get('id'),'externalId':x.get('externalId'),'action':x.get('action'),'reason':x.get('reason','')} for x in decisions]})
    return updated
def fingerprint(row):
    # Documento + data + valor, quando fornecido; sem documento, descrição/finalidade/subconta.
    key=[row.get('date'),round(row['value'],2),fold(row.get('document'))]
    if not row.get('document'):key += [fold(row.get(k)) for k in ['description','purpose','subaccount','code']]
    return hashlib.sha256(json.dumps(key,ensure_ascii=False).encode()).hexdigest()
def classify(row,state):
    accounts=state['budget'];code=fold(row.get('code'));category=fold(row.get('category'))
    if code:
        matches=[x for x in accounts if fold(x.get('code'))==code]
        if len(matches)==1:return matches[0]['id'],'high','Código exato'
    if category:
        matches=[x for x in accounts if fold(x.get('category'))==category]
        if len(matches)==1:return matches[0]['id'],'high','Conta exata'
    signature='|'.join(fold(row.get(k)) for k in ['subaccount','description','purpose'])
    matches=[x for x in state.get('classificationRules',[]) if x['signature']==signature and any(a['id']==x['accountId'] for a in accounts)]
    if signature.strip('|') and len({x['accountId'] for x in matches})==1 and matches:return matches[0]['accountId'],'high','Associação confirmada anteriormente'
    text=' '.join(fold(row.get(k)) for k in ['subaccount','description','purpose'])
    suggestions=[x for x in accounts if fold(x['category']) and fold(x['category']) in text]
    return (suggestions[0]['id'] if len(suggestions)==1 else None),'review','Requer confirmação humana'
def build_preview(name,content,kind,state,source_confirmed=False):
    if kind not in ['budget','expenses','financial','accounting','classes']:raise ValueError('Tipo de importação inválido.')
    from source_import import preview as source_preview
    specific=source_preview(name,content,kind,state)
    if specific is not None:return specific
    if kind=='classes':raise ValueError('Use o relatório Turmas CAJ 2027 em PDF textual.')
    raw,sheet=read_rows(name,content,kind)
    if kind in ['budget','expenses'] and sheet is None and not source_confirmed:raise ValueError('Confirme que CSV/PDF corresponde exclusivamente a DESPESAS 2026.')
    records,headers=mapped(raw);result={'name':name,'kind':kind,'sheet':sheet,'rows':[],'invalid':[],'duplicates':0,'automatic':0,'pending':0,'fileHash':hashlib.sha256(content).hexdigest()}
    if any(x.get('fileHash')==result['fileHash'] and x.get('kind')==kind for x in state['imports']):raise ValueError('Este arquivo já foi confirmado neste tipo de importação.')
    fingerprints={x.get('fingerprint') for a in state['budget'] for x in a.get('expenses',[])}|{x.get('fingerprint') for x in state['reviewQueue']}
    for line,r in enumerate(records,2):
        try:
            item={k:str(v or '').strip() for k,v in r.items()};item['line']=line;item['id']=secrets.token_hex(12)
            if kind=='budget':
                item['budget']=decimal(r.get('budget'));item['category']=item.get('category','');item['code']=item.get('code','')
                if item['budget']<0 or not item['category']:raise ValueError('Categoria e orçamento não negativo são obrigatórios.')
                if any(x.get('code')==item['code'] and item['code'] or fold(x.get('category'))==fold(item['category']) for x in result['rows']):raise ValueError('Conta repetida no arquivo.')
                item['status']='ready'
            elif kind=='expenses':
                item['value']=decimal(r.get('value'));item['date']=date_value(r.get('date'))
                if item['date'][:4]!='2026':raise ValueError('Despesa fora do exercício de 2026.')
                if item['value']<=0 or not (item.get('description') or item.get('purpose')):raise ValueError('Valor positivo e descrição/finalidade obrigatórios.')
                item['fingerprint']=fingerprint(item)
                if item['fingerprint'] in fingerprints:result['duplicates']+=1;continue
                fingerprints.add(item['fingerprint']);item['accountId'],confidence,item['reason']=classify(item,state);item['status']='ready' if confidence=='high' else 'review';result['automatic' if confidence=='high' else 'pending']+=1
            else:
                month=str(r.get('month','')).strip()
                if re.fullmatch(r'\d{2}/\d{4}',month):month=month[3:]+'-'+month[:2]
                datetime.strptime(month,'%Y-%m');item['month']=month;field='financialPercent' if kind=='financial' else 'accountingPercent';item[field]=decimal(r.get(field))
                if not 0<=item[field]<=100:raise ValueError('Percentual fora de 0 a 100.')
                for k in ['debt','students','guardians']:
                    if r.get(k) not in [None,'']:
                        item[k]=decimal(r[k])
                        if item[k]<0 or k!='debt' and not item[k].is_integer():raise ValueError('Valor/contagem inválida.')
                item['status']='ready'
            result['rows'].append(item)
        except (ValueError,TypeError) as error:result['invalid'].append({'line':line,'error':str(error)})
    if not result['rows'] and not result['duplicates']:raise ValueError('Nenhum registro válido encontrado. '+('; '.join(x['error'] for x in result['invalid'][:3])))
    return result

def _budget_class_rows(pages):
    """Lê linhas inequívocas da Ficha Financeira sem inferir qualquer lacuna."""
    labels=['revenue','otherIncome','totalRevenue','gratuities','commercialDiscount','netRevenue','personalPayroll','supportPayrollAllocation','generalExpenses','loans','totalExpenses','result']
    money=r'\(?-?[\d.]+,\d{2}\)?|-'
    pattern=re.compile(r'^\s*((?:Grupo\s+[2-5]|[1-9]º\s+Ano|[1-3]º\s+E\.\s*M[ée]dio)\s+[A-C])\s+(\d+)\s+(.*)$',re.I)
    rows=[]
    for page_no,page in enumerate(pages,1):
        for line_no,line in enumerate(page.splitlines(),1):
            found=pattern.match(line)
            if not found: continue
            tokens=re.findall(money,found.group(3))
            if len(tokens)!=12: continue
            pe=re.search(r'\s(\d+)\s*$',found.group(3))
            if not pe: continue
            # Na tabela original, "-" representa célula sem valor informado.
            raw={key:(0 if value=='-' else decimal(value)) for key,value in zip(labels,tokens)}
            students=int(found.group(2));pending=students==0 and any(raw.values())
            rows.append({'id':secrets.token_hex(8),'className':found.group(1),'students':students,'breakEvenStudents':int(pe.group(1)),**raw,
                'calculatedNetRevenue':round(raw['totalRevenue']-raw['gratuities']-raw['commercialDiscount'],2),
                'calculatedResult':round(raw['netRevenue']-raw['totalExpenses'],2),
                'netRevenueDifference':round(raw['totalRevenue']-raw['gratuities']-raw['commercialDiscount']-raw['netRevenue'],2),
                'resultDifference':round(raw['netRevenue']-raw['totalExpenses']-raw['result'],2),
                'origin':f'página {page_no}, linha {line_no}','status':'pending_review' if pending else 'recognized',
                'pendingReason':'Turma sem alunos com valores financeiros; exige conciliação administrativa.' if pending else None,'sourceValues':tokens})
    return rows

def budget_version_preview(name,content):
    """Prévia conservadora de orçamento anual: nunca grava e sempre preserva a origem."""
    ext=Path(name).suffix.lower(); text=''; pages=[]
    if ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(content))
        if len(reader.pages)>150:raise ValueError('PDF excede 150 páginas.')
        pages=[page.extract_text(extraction_mode='layout') or '' for page in reader.pages];text='\n'.join(pages)
    elif ext=='.xlsx':
        book=__import__('openpyxl').load_workbook(io.BytesIO(content),read_only=True,data_only=True);text='\n'.join(' '.join(str(c.value or '') for c in row) for sheet in book for row in sheet.iter_rows());book.close()
    else:raise ValueError('Orçamento anual aceita PDF textual ou XLSX.')
    def origin(found):return f'página {next((i+1 for i,p in enumerate(pages) if found and found.group(0) in p),"—")}' if found else 'não identificado'
    def amount(pattern,label,kind='Indicador'):
        found=re.search(pattern,text,re.I|re.M);return {'id':secrets.token_hex(8),'category':label,'subcategory':'','description':label,'originalDescription':label,'detected':decimal(found.group(1)) if found else None,'monthly':None,'annual':decimal(found.group(1)) if found else None,'type':kind,'costCenter':None,'className':None,'period':'anual','origin':origin(found),'confidence':'high' if found else 'none','status':'recognized' if found else 'pending_review'}
    year=re.search(r'20\d{2}',Path(name).stem) or re.search(r'20\d{2}',text[:1500]);items=[{'id':secrets.token_hex(8),'category':'Ano letivo','subcategory':'','description':'Ano detectado','originalDescription':'Ano detectado','detected':int(year.group()) if year else None,'monthly':None,'annual':None,'type':'Indicador','costCenter':None,'className':None,'period':'anual','origin':'nome do arquivo' if year and year.group() in Path(name).stem else 'primeira página','confidence':'high' if year else 'none','status':'recognized' if year else 'pending_review'},amount(r'[Mm]eta\s+de\s+[Aa]lunos\s+(\d{3,4})','Meta de alunos'),amount(r'(?:[Pp]revis.o\s+)?[Ii]nadimpl.ncia\s+(\d+[,.]\d+)\s*%','Inadimplência'),amount(r'(?m)^TOTAL\s+DESPESAS(?!\s+MANUT).*?\s([\d.]+,[0-9]{2})\s*$','Despesas totais','Despesa'),amount(r'(?m)^RESULTADO\s+DO\s+EXERC.CIO.*?\s([\d.]+,[0-9]{2})\s*$','Resultado','Indicador')]
    groups=[('folha','Folha de pagamento'),('encarg','Encargos'),('prof','Professores'),('aux','Auxiliares'),('estagi','Estagiários'),('energia','Energia'),('agua','Água'),('internet','Internet'),('software','Softwares'),('publicidade','Publicidade'),('manut','Manutenção'),('depreci','Depreciação'),('pcld','PCLD'),('fundo','Fundo de Educação'),('rateio sede','Rateio Sede'),('contribui','Contribuições'),('material','Materiais')]
    money_pattern=r'\(?-?[\d.]+,\d{2}\)?'
    for page_no,page in enumerate(pages,1):
        for line_no,line in enumerate(page.splitlines(),1):
            values=re.findall(money_pattern,line)
            code=re.match(r'\s*(\d{5,})\s+(.+)',line)
            if not code or len(values)<1:continue
            label=code.group(2).strip();folded=fold(label);match=next((x for x in groups if x[0] in folded),None)
            annual=decimal(values[-1]);monthly=decimal(values[-2]) if len(values)>1 else None
            items.append({'id':secrets.token_hex(8),'category':match[1] if match else 'Pendente de classificação','subcategory':'','description':label,'originalDescription':label,'detected':annual,'monthly':monthly,'annual':annual,'type':'Despesa','costCenter':None,'className':None,'period':'anual','origin':f'página {page_no}, linha {line_no}, conta {code.group(1)}','confidence':'high' if match else 'low','status':'recognized' if match else 'pending_review'})
    class_rows=_budget_class_rows(pages) if pages else []
    official_totals={}
    total_line=next((line for page in pages for line in page.splitlines() if re.match(r'^\s*TOTAL\s+\d+',line)),None)
    if total_line:
        values=re.findall(money_pattern,total_line)
        pe=re.search(r'\s(\d+)\s*$',total_line)
        students=re.search(r'^\s*TOTAL\s+(\d+)',total_line)
        if len(values)>=11 and pe and students:
            totals=[0 if value=='-' else decimal(value) for value in values[:12]]
            # A coluna Empréstimos vazia pode não ser emitida pelo extrator PDF.
            offset=1 if len(values)==11 else 0
            official_totals={'students':int(students.group(1)),'grossRevenueMonthly':totals[2],'netRevenueMonthly':totals[5],'totalExpensesMonthly':totals[10-offset],'resultMonthly':totals[11-offset],'breakEvenStudents':int(pe.group(1))}
    annual_line=next((line for page in pages for line in page.splitlines() if re.match(r'^\s*TOTAL\s+DO\s+ANO',line)),None)
    if annual_line:
        annual_values=re.findall(money_pattern,annual_line)
        if len(annual_values)>=11:
            annual=[0 if value=='-' else decimal(value) for value in annual_values]
            offset=1 if len(annual_values)==11 else 0
            official_totals.update({'grossRevenueAnnual':annual[2],'netRevenueAnnual':annual[5],'totalExpensesAnnual':annual[10-offset],'resultAnnual':annual[11-offset]})
    return {'name':name,'fileHash':hashlib.sha256(content).hexdigest(),'items':items,'classRows':class_rows,'officialTotals':official_totals,'pages':len(pages) if pages else None,'status':'analysis','warnings':['Prévia não aplica dados. Itens sem categoria segura permanecem pendentes de conferência.','Valores mensal e anual permanecem separados quando o arquivo os informa.','Linhas de turma sem estrutura inequívoca não são associadas automaticamente às turmas operacionais.']}
def apply_preview(state,preview,user):
    if preview.get('blocked'):raise ValueError('Importação bloqueada: resolva as divergências apresentadas na prévia.')
    s=copy.deepcopy(state);kind=preview['kind'];stamp=now()
    for row in preview['rows']:
        if preview.get('sourceAdapter'):
            record={k:v for k,v in row.items() if k not in ['line','status']};record.update(source=preview['name'],fileHash=preview['fileHash'],updatedAt=stamp)
            if kind=='budget':record['expenses']=[];s['budget'].append(record)
            elif kind=='classes':s['classes'].append(record)
            else:
                d=s['delinquency'];d[kind+'Source']=record;d[kind+'Percent']=row['percent'];d[kind+'Debt']=row['debt'];d['updatedAt']=stamp
                if kind=='financial':d['debt']=row['debt']
                monthly=d.setdefault('monthly',[]);current=next((x for x in monthly if x['month']==row['month']),None)
                if current is None:current={'month':row['month']};monthly.append(current)
                current.update({kind+'Percent':row['percent'],kind+'Debt':row['debt'],kind+'Reference':row['reference']})
            continue
        if kind=='budget':
            match=next((a for a in s['budget'] if row.get('code') and str(a.get('code'))==row['code'] or fold(a['category'])==fold(row['category'])),None)
            if match:match.update(budget=row['budget'],updatedAt=stamp,year=2026,code=row.get('code') or match.get('code',''))
            else:s['budget'].append({'id':row['id'],'category':row['category'],'code':row.get('code',''),'budget':row['budget'],'actual':0,'year':2026,'updatedAt':stamp,'expenses':[]})
        elif kind=='expenses':
            if row['status']=='review':s['reviewQueue'].append(dict(row,source=preview['name'],createdAt=stamp));continue
            account=next((a for a in s['budget'] if a['id']==row['accountId']),None)
            if not account:raise ValueError('Conta alterada. Gere uma nova prévia.')
            account.setdefault('expenses',[]).append(dict(row,user=user['name'],userId=user['id'],source=preview['name']));account['actual']=round(account['actual']+row['value'],2);account['updatedAt']=stamp
        else:
            d=s['delinquency'];monthly=d.setdefault('monthly',[]);current=next((x for x in monthly if x['month']==row['month']),None)
            if not current:current={'month':row['month']};monthly.append(current)
            for key in ['financialPercent','accountingPercent','debt','students','guardians']:
                if key in row:current[key]=row[key]
            current['updatedAt']=stamp
            latest=max(monthly,key=lambda x:x['month']);d.update({k:latest.get(k) for k in ['financialPercent','accountingPercent']});d.update({k:v for k,v in latest.items() if k in ['debt','students','guardians','month']});d['updatedAt']=stamp
    s['imports'].append({'id':secrets.token_hex(12),'name':preview['name'],'kind':kind,'type':'inadimplência' if kind in ['financial','accounting'] else 'orçamento','date':stamp,'fileHash':preview['fileHash'],'count':len(preview['rows']),'automatic':preview['automatic'],'pending':preview['pending'],'duplicates':preview['duplicates'],'rejected':len(preview['invalid']),'user':user['name']})
    return s

def financial_report_rows(state):
    """Linhas de exportação da projeção por turma; nunca infere classificações."""
    year=next((x for x in state.get('academicYears',[]) if x.get('year')==state.get('activeAcademicYear')),None)
    if not year:raise ValueError('Ano financeiro ativo não encontrado.')
    tuition={x['id']:x.get('tuition') for x in year.get('parameters',[])}
    fixed_labels=[(key,label,percent) for key,label,percent in [
        ('noDiscount','Sem desconto',0),('philanthropic100','Bolsa filantrópica 100%',100),('philanthropic50','Bolsa filantrópica 50%',50),
        ('staffChild100','Filho de funcionário 100%',100),('staffChild80','Filho de funcionário 80%',80),
        ('workerChild100','Filho de obreiro 100%',100),('markingLives45','Projeto Marcando Vidas 45%',45)]]
    def segment(room):
        if room.get('stage')=='Educação Infantil':return 'early'
        if room.get('stage')=='Fundamental Anos Iniciais':return 'fundamental1'
        if room.get('stage')=='Fundamental Anos Finais':return 'fundamental2'
        return 'secondary3' if str(room.get('name','')).startswith('3º') else 'secondary12'
    columns=['Ano','Turma','Segmento','Matriculados','Mensalidade (R$)','Sem desconto','Filantrópica 100%','Filantrópica 50%','Funcionário 100%','Funcionário 80%','Obreiro 100%','Marcando Vidas 45%','Variáveis','Não classificados financeiramente','Receita bruta conhecida (R$)','Descontos (R$)','Receita líquida (R$)','Ticket médio (R$)']
    rows=[]
    for room in state['classes']:
        classification=year.get('classifications',{}).get(room['id'],{})
        fixed=classification.get('fixed',{})
        variables=classification.get('variables',[])
        counts=[fixed.get(key,0) for key,_,_ in fixed_labels]
        variable_count=sum(x.get('quantity',0) for x in variables)
        classified=sum(counts)+variable_count
        total=room.get('students',0)
        value=tuition.get(segment(room))
        gross=classified*value if value is not None else None
        discount=None if value is None else sum(count*value*percent/100 for count,(_,_,percent) in zip(counts,fixed_labels))+sum(x.get('quantity',0)*value*x.get('percent',0)/100 for x in variables)
        net=None if gross is None else gross-discount
        rows.append([year['year'],room['name'],segment(room),total,value,*counts,variable_count,max(0,total-classified),gross,discount,net,round(net/classified,2) if classified and net is not None else None])
    return columns,rows

def report_rows(state,kind):
    c=state['classes'];hist=state['enrollments']['history'];columns=[];rows=[]
    if kind=='financial':
        return financial_report_rows(state)
    if kind in ['classes','enrollments']:
        columns=['Turma','Capacidade','Rematrículas','Alunos novos','Total','Vagas','Excedentes','Ocupação']
        for room in c:
            count=lambda typ:room.get('opening',{}).get(typ,0)+sum(x['quantity'] for x in hist if str(x['classId'])==str(room['id']) and x['type']==typ)
            new,re=count('new'),count('re');total=new+re;cap=room['capacity'];rows.append([room['name'],cap,re,new,total,max(0,cap-total),max(0,total-cap),f'{total/cap*100:.1f}%' if cap else 'Sem capacidade'])
    elif kind=='budget':columns=['Conta','Código','Orçado (R$)','Realizado (R$)','Disponível (R$)','Executado','Atualização'];rows=[[x['category'],x.get('code',''),x['budget'],x['actual'],round(x['budget']-x['actual'],2),f"{x['actual']/x['budget']*100:.1f}%" if x['budget'] else 'Não definido',x.get('updatedAt','')] for x in state['budget']]
    elif kind=='benefits':columns=['Benefício','Tipo','Quantidade','Situação'];rows=[[x['name'],x.get('type',''),x['quantity'],'Ativo' if x.get('active') else 'Inativo'] for x in state['benefits']]
    elif kind=='requests':columns=['Solicitação','Solicitante','Destinatário','Status','Prazo','Motivo da reprovação'];rows=[[x['title'],x.get('creatorName',x.get('requestedBy','')),x.get('recipientName',x.get('owner','')),x['status'],x.get('due',''),x.get('rejectionReason','')] for x in state['requests']]
    elif kind=='delinquency':columns=['Mês','Financeira (%)','Contábil (%)','Dívida (R$)','Alunos','Responsáveis'];rows=[[x.get(k) for k in ['month','financialPercent','accountingPercent','debt','students','guardians']] for x in state['delinquency'].get('monthly',[])]
    elif kind=='guardians':columns=['Nome','Vínculo','WhatsApp','E-mail','Aluno','Turma'];rows=[[x.get(k,'') for k in ['name','relationship','phone','email','student','className']] for x in state['guardians']]
    elif kind=='audit':columns=['Usuário','Ação','Módulo','Data/hora','Informação','Anterior','Novo'];rows=[[x.get('user'),x.get('action'),x.get('module'),x.get('date'),x.get('field'),str(x.get('before','')),str(x.get('after',''))] for x in state['audit']]
    elif kind=='summary':
        total=state['enrollments']['new']+state['enrollments']['re'];d=state['delinquency'];columns=['Indicador','Valor'];rows=[['Matriculados',total],['Alunos novos',state['enrollments']['new']],['Rematrículas',state['enrollments']['re']],['Meta oficial',1065],['Percentual alcançado',f'{total/1065*100:.2f}%'],['Faltam',max(0,1065-total)],['Orçado (R$)',sum(x['budget'] for x in state['budget'])],['Realizado (R$)',sum(x['actual'] for x in state['budget'])],['Inadimplência financeira (%)',d.get('financialPercent')],['Inadimplência contábil (%)',d.get('accountingPercent')],['Dívida (R$)',d.get('debt',0)]]
    else:raise ValueError('Relatório inválido.')
    return columns,rows
def pdf_report(title,columns,rows,logo=None):
    if not logo or not Path(logo).exists():logo=Path(__file__).parent/'assets'/'educacao-adventista.jpg'
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
    stream=io.BytesIO();size=landscape(A4) if len(columns)>5 else A4;w,h=size
    doc=SimpleDocTemplate(stream,pagesize=size,leftMargin=32,rightMargin=32,topMargin=96,bottomMargin=48)
    styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='CellCAJ',fontName='Helvetica',fontSize=8,leading=11,textColor=colors.HexColor('#152238'),wordWrap='CJK'));styles.add(ParagraphStyle(name='HeadCAJ',parent=styles['CellCAJ'],textColor=colors.white,fontName='Helvetica-Bold'))
    def page(canvas,doc):
        canvas.saveState();canvas.setFillColor(colors.HexColor('#063b76'));canvas.rect(0,h-76,w,76,fill=1,stroke=0);canvas.setFillColor(colors.HexColor('#f6c62f'));canvas.rect(0,h-80,w,4,fill=1,stroke=0);canvas.setFillColor(colors.white);canvas.setFont('Helvetica-Bold',15);canvas.drawString(32,h-27,SCHOOL);canvas.setFont('Helvetica',9);canvas.drawString(32,h-43,ADDRESS);canvas.drawString(32,h-58,CNPJ)
        if logo and Path(logo).exists():
            canvas.drawImage(str(logo),w-84,h-66,width=48,height=48,preserveAspectRatio=True,mask='auto');canvas.saveState();canvas.setFillAlpha(.065);canvas.drawImage(str(logo),w/2-150,h/2-150,width=300,height=300,preserveAspectRatio=True,mask='auto');canvas.restoreState()
        canvas.setStrokeColor(colors.HexColor('#d5e2f0'));canvas.line(32,36,w-32,36);canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#53677f'));canvas.drawString(32,23,SCHOOL+' | Gestão Escolar CAJ');canvas.drawRightString(w-32,23,'Página '+str(doc.page));canvas.restoreState()
    def cell(x,header=False):return Paragraph(escape(str('Não informado' if x is None else x)).replace('\n','<br/>'),styles['HeadCAJ' if header else 'CellCAJ'])
    story=[Paragraph(escape(title),styles['Heading2']),Paragraph('Emitido em '+datetime.now().strftime('%d/%m/%Y %H:%M'),styles['CellCAJ']),Spacer(1,10)]
    if not logo or not Path(logo).exists():story.extend([Paragraph('Logo e marca-d’água institucional pendentes de fornecimento.',styles['CellCAJ']),Spacer(1,8)])
    widths=[(w-64)/len(columns)]*len(columns)
    if len(columns)>5:widths[0]*=1.55;scale=(w-64)/sum(widths);widths=[x*scale for x in widths]
    data=[[cell(x,True) for x in columns]]+[[cell(x) for x in row] for row in rows]
    if not rows:data.append([cell('Sem registros disponíveis.')]+[cell('') for _ in columns[1:]])
    table=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT',splitInRow=1);table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#075eaf')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('LINEBELOW',(0,0),(-1,-1),.3,colors.HexColor('#d5e2f0')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f4f8fd')])]))
    story.append(table);doc.build(story,onFirstPage=page,onLaterPages=page);return stream.getvalue()
def xlsx_report(columns,rows):
    # Documento XLSX de dados, sem fórmulas externas, macros ou interpretação de strings como fórmulas.
    stream=io.BytesIO();xml=['<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" state="frozen"/></sheetView></sheetViews><cols>']
    xml.extend(f'<col min="{i+1}" max="{i+1}" width="26" customWidth="1"/>' for i in range(len(columns)));xml.append('</cols><sheetData>')
    for i,row in enumerate([columns]+rows,1):
        xml.append(f'<row r="{i}">')
        for value in row:
            if isinstance(value,(int,float)):xml.append(f'<c><v>{value}</v></c>')
            else:xml.append('<c t="inlineStr"><is><t xml:space="preserve">'+escape(str(value if value is not None else ''))+'</t></is></c>')
        xml.append('</row>')
    xml.append('</sheetData></worksheet>')
    with zipfile.ZipFile(stream,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml','<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
        z.writestr('_rels/.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        z.writestr('xl/workbook.xml','<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Relatório CAJ" sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr('xl/_rels/workbook.xml.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>');z.writestr('xl/worksheets/sheet1.xml',''.join(xml))
    return stream.getvalue()
def api_service(handler,user,path,p):
    store=handler.store
    if path=='/api/teaching-load/preview':
        if not user.get('isAdmin'):raise PermissionError('Somente administradores podem importar carga horária.')
        state,_=store.state();return handler.respond(200,teaching_load_preview(p['name'],base64.b64decode(p['content'],validate=True),state,p.get('targetYear',2027)))
    if path=='/api/budget-version/preview':
        require(user,'financial','edit')
        return handler.respond(200,budget_version_preview(p['name'],base64.b64decode(p['content'],validate=True)))
    if path=='/api/financial-classifications/template':
        require(user,'financial','edit');return handler.binary(financial_import_workbook(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','modelo-descontos-7e7.xlsx')
    if path=='/api/export':
        require(user,'reports');kind=p['kind']
        if kind=='audit':require(user,'users')
        if kind=='financial':require(user,'financial')
        state=store.snapshot(user)['state'];columns,rows=report_rows(state,kind);query=fold(p.get('query',''));rows=[r for r in rows if query in fold(' '.join(str(x) for x in r))]
        if p.get('format')=='pdf':return handler.binary(pdf_report(p.get('title','Relatório CAJ'),columns,rows,Path(store.path).parent/'logo.png'),'application/pdf','relatorio-caj.pdf')
        if p.get('format')=='xlsx':return handler.binary(xlsx_report(columns,rows),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','relatorio-caj.xlsx')
        raise ValueError('Formato de exportação inválido.')
    if path=='/api/financial-classifications/preview':
        require(user,'financial','edit');state,version=store.state();preview=financial_classification_preview(p['name'],base64.b64decode(p['content'],validate=True),state,p.get('year'))
        token=secrets.token_urlsafe(28);preview['version']=version
        with store.db() as db:db.execute('DELETE FROM previews WHERE expires<?',(time.time(),));db.execute('INSERT INTO previews VALUES(?,?,?,?,?)',(token,user['id'],version,json.dumps(preview),time.time()+1800))
        return handler.respond(200,dict(preview,token=token))
    if path=='/api/financial-classifications/confirm':
        if not user.get('isAdmin'):raise PermissionError('Somente administradores podem aprovar importações financeiras.')
        with store.lock,store.db() as db:
            row=db.execute('SELECT * FROM previews WHERE token=? AND user_id=? AND expires>?',(p.get('token'),user['id'],time.time())).fetchone()
            if not row:raise ValueError('Prévia expirada ou já utilizada. Gere novamente.')
            preview=json.loads(row['payload'])
            if preview.get('kind')!='financial_classifications':raise ValueError('Prévia não pertence à importação de descontos e benefícios.')
            state,version=store.state(db)
            if row['version']!=version:raise RuntimeError('Dados alterados após a prévia. Gere uma nova prévia antes de confirmar.')
            if p.get('confirmation')!='APROVAR IMPORTAÇÃO FINANCEIRA':raise ValueError('Confirmação administrativa obrigatória.')
            updated=apply_financial_classification_preview(state,preview,p.get('decisions',[]),user);changes={key:value for key,value in updated.items() if value!=state[key]}
            db.commit();store.patch(user,changes,version,'financial','import');db.execute('DELETE FROM previews WHERE token=?',(p['token'],))
            return handler.respond(200,store.snapshot(user))
    if path=='/api/revenue-2027/preview':
        require(user,'financial','edit');state,version=store.state();preview=revenue_preview(p['name'],base64.b64decode(p['content'],validate=True),state,p.get('year',2027));token=secrets.token_urlsafe(28);preview['version']=version
        with store.db() as db:db.execute('DELETE FROM previews WHERE expires<?',(time.time(),));db.execute('INSERT INTO previews VALUES(?,?,?,?,?)',(token,user['id'],version,json.dumps(preview),time.time()+1800))
        return handler.respond(200,dict(preview,token=token))
    if path=='/api/revenue-2027/confirm':
        if not user.get('isAdmin'):raise PermissionError('Somente administradores podem aprovar receitas consolidadas.')
        with store.lock,store.db() as db:
            row=db.execute('SELECT * FROM previews WHERE token=? AND user_id=? AND expires>?',(p.get('token'),user['id'],time.time())).fetchone()
            if not row:raise ValueError('Prévia expirada ou já utilizada. Gere novamente.')
            preview=json.loads(row['payload']);state,version=store.state(db)
            if version!=row['version']:raise RuntimeError('Os dados mudaram desde a prévia. Gere uma nova prévia.')
            if p.get('confirmation')!='APROVAR RECEITAS CONSOLIDADAS':raise ValueError('Confirmação administrativa obrigatória.')
            updated=apply_revenue_preview(state,preview,p.get('decisions',[]),user)
            from server import validate
            validate(updated,state);store.audit(db,user,'financial','approve','revenuePlanning',state.get('revenuePlanning'),updated.get('revenuePlanning'));db.execute('UPDATE state SET payload=?,version=? WHERE id=1',(json.dumps(updated),version+1));db.execute('DELETE FROM previews WHERE token=?',(p['token'],))
        return handler.respond(200,store.snapshot(user))
    if path=='/api/import/preview':
        kind=p['kind'];require(user,'delinquency' if kind in ['financial','accounting'] else 'classes' if kind=='classes' else 'budget','edit');state,version=store.state();preview=build_preview(p['name'],base64.b64decode(p['content'],validate=True),kind,state,p.get('sourceConfirmed',False));token=secrets.token_urlsafe(28)
        with store.db() as db:db.execute('DELETE FROM previews WHERE expires<?',(time.time(),));db.execute('INSERT INTO previews VALUES(?,?,?,?,?)',(token,user['id'],version,json.dumps(preview),time.time()+1800))
        return handler.respond(200,dict(preview,token=token,version=version))
    with store.lock,store.db() as db:
        row=db.execute('SELECT * FROM previews WHERE token=? AND user_id=? AND expires>?',(p['token'],user['id'],time.time())).fetchone()
        if not row:raise ValueError('Prévia expirada ou já utilizada. Gere novamente.')
        preview=json.loads(row['payload']);state,version=store.state(db)
        if row['version']!=version:raise RuntimeError('Dados alterados após a prévia. Gere uma nova prévia antes de confirmar.')
        if p.get('confirmation')!='CONFIRMAR IMPORTAÇÃO':raise ValueError('Confirmação obrigatória.')
        updated=apply_preview(state,preview,user);changes={k:v for k,v in updated.items() if v!=state[k]};module='delinquency' if preview['kind'] in ['financial','accounting'] else 'classes' if preview['kind']=='classes' else 'budget'
        # patch utiliza outra conexão: nenhuma escrita nesta transação até a validação confirmar.
        db.commit();store.patch(user,changes,version,module,'import');db.execute('DELETE FROM previews WHERE token=?',(p['token'],))
    return handler.respond(200,store.snapshot(user))
