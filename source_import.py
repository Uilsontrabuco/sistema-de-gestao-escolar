"""Adaptadores conservadores para os relatórios CAJ; somente prévias, sem escrita no banco."""
import hashlib, io, re, unicodedata
from decimal import Decimal, ROUND_HALF_UP

def folded(value):
    return ''.join(c for c in unicodedata.normalize('NFD',str(value or '')) if not unicodedata.combining(c)).lower().strip()

def money(value):
    if value is None:return Decimal(0)
    if isinstance(value,(int,float)):return Decimal(str(value))
    return Decimal(str(value).replace('.','').replace(',','.').replace('(','-').replace(')',''))

def close(a,b,tolerance='0.01'):return abs(money(a)-money(b))<=Decimal(tolerance)

def preview(name,content,kind,state):
    digest=hashlib.sha256(content).hexdigest()
    out=dict(name=name,kind=kind,fileHash=digest,rows=[],invalid=[],duplicates=0,automatic=0,pending=0,warnings=[],reconciliation={},sourceAdapter=True)
    if any(x.get('fileHash')==digest and x.get('kind')==kind for x in state.get('imports',[])):raise ValueError('Este arquivo já foi importado.')
    def check(ok,message):
        if not ok:out['invalid'].append({'line':'conferência','error':message})
    if kind=='budget' and name.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        candidates=[s for s in book if folded(s.cell(1,1).value)=='despesas anuais - 2026']
        if len(candidates)!=1:book.close();return None
        sheet=candidates[0];out['sheet']=sheet.title
        headers=[folded(sheet.cell(2,i).value) for i in range(1,8)]
        if headers!=['finalidade','cod','nome sub conta','orcado','despesas','disponivel','%']:book.close();raise ValueError('Cabeçalho do demonstrativo anual foi alterado. Revise o mapeamento.')
        totals=None
        for line,values in enumerate(sheet.iter_rows(min_row=3,values_only=True),3):
            a,code,sub,budget,actual,available,percent=values[:7]
            if folded(sub)=='total':totals=[budget,actual,available,percent];continue
            if not any(v is not None for v in values):continue
            if totals:check(False,'Há linhas após o total.');continue
            check(bool(a),f'Linha {line}: finalidade ausente.')
            b,r=money(budget),money(actual)
            check(b>=0 and r>=0,f'Linha {line}: valor negativo não suportado.')
            if available is not None:check(abs(b-r-money(available))<=Decimal('.01'),f'Linha {line}: disponível divergente.')
            if percent is not None and b:check(abs(Decimal(str(percent))-r/b)<=Decimal('.0001'),f'Linha {line}: percentual divergente.')
            raw=dict(zip(['FINALIDADE','COD','NOME SUB CONTA','ORÇADO','DESPESAS/REALIZADO','DISPONÍVEL','percentual'],values[:7]))
            out['rows'].append(dict(id='budget-'+digest[:12]+'-'+str(line),line=line,category=str(a or '').strip(),purpose=str(a or '').strip(),code='' if code is None else str(code),subaccount=str(sub or ''),budget=float(b),actual=float(r),year=2026,sourceValues=raw,budgetInformed=budget is not None,actualInformed=actual is not None,status='ready'))
        book.close()
        sums=[sum((Decimal(str(x[k])) for x in out['rows']),Decimal(0)) for k in ['budget','actual']];sums.append(sums[0]-sums[1])
        check(totals is not None,'Total original não encontrado.')
        if totals:
            for i,label in enumerate(['orçado','realizado','disponível']):check(abs(sums[i]-money(totals[i]))<=Decimal('.01'),'Total '+label+' divergente.')
            check(not sums[0] or abs(Decimal(str(totals[3]))-sums[1]/sums[0])<=Decimal('.0001'),'Percentual total divergente.')
        out['reconciliation']=dict(budget=float(sums[0]),actual=float(sums[1]),available=float(sums[2]),percent=float(sums[1]/sums[0]*100) if sums[0] else None,original=totals)
        out['warnings']=['A aba '+out['sheet']+' foi reconhecida pelo título DESPESAS ANUAIS - 2026 e pelo cabeçalho completo.','Células vazias foram preservadas em sourceValues; zero é usado somente no cálculo agregado.','COD repetido não agrupa finalidades diferentes.']
        check(not state.get('budget'),'Já existem contas. Este demonstrativo é um saldo acumulado; concilie os saldos existentes antes de importar para evitar duplicação.')
    elif kind in ['classes','financial','accounting'] and name.lower().endswith('.pdf'):
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(content))
        if len(reader.pages)>150:raise ValueError('PDF excede 150 páginas.')
        text='\n'.join(p.extract_text(extraction_mode='layout') or '' for p in reader.pages)
        if kind=='classes':
            if 'matriculas 2027 caj' not in folded(text):raise ValueError('Título Matrículas 2027 CAJ não encontrado.')
            for line,m in enumerate(re.finditer(r'^(.+?)\s{2,}(\d+)\s+(\d+)\s+(\d+)\s+(-?\d+)\s*$',text,re.M),1):
                label=m[1].strip();cap,total,new,vac=map(int,m.groups()[1:])
                check(total>=new and cap-total==vac,'Turma '+label+': total ou vagas divergentes.')
                out['rows'].append(dict(id='class-2027-'+hashlib.sha256(label.encode()).hexdigest()[:16],line=line,name=label,capacity=cap,opening={'new':new,'re':total-new},students=total,unclassified=0,year=2027,sourceValues={'total':total,'new':new,'vacancies':vac},status='ready'))
            printed=re.search(r'^\s*Total\s+(\d+)\s+(\d+)\s+(-?\d+)\s*$',text,re.M)
            sums=[sum(x[k] for x in out['rows']) for k in ['capacity','students']];sums.append(sums[0]-sums[1])
            check(bool(printed) and sums==list(map(int,printed.groups())),'Totais de capacidade, matriculados ou vagas divergentes.')
            check(bool(out['rows']),'Nenhuma turma reconhecida.')
            out['reconciliation']=dict(capacity=sums[0],total=sums[1],vacancies=sums[2],new=sum(x['opening']['new'] for x in out['rows']),re=sum(x['opening']['re'] for x in out['rows']),excess=sum(max(0,x['students']-x['capacity']) for x in out['rows']))
            out['warnings']=['Rematrículas derivadas de Total Alunos Matri menos Alunos novos; o PDF não traz coluna explícita de rematrículas.','Vagas restantes preservam valores negativos. Excedentes são destacados separadamente.']
            check(not state.get('classes') and not state.get('enrollments',{}).get('history') and not any(state.get('enrollments',{}).get('unallocated',{}).values()),'Existem turmas ou matrículas. Concilie a base atual antes de importar este saldo inicial.')
        else:
            title='inadimplencia financeira' if kind=='financial' else 'inadimplencia contabil'
            if title not in folded(text):return None
            check('2026 - ano letivo' in folded(text),'Período de 2026 não identificado.')
            entity=[l for l in text.splitlines() if '13421|501' in l]
            total=[l for l in text.splitlines() if re.match(r'^\s*Total\s',l)]
            pattern=r'\(?-?\d[\d.]*,\d{2}\)?'
            check(len(entity)==1 and len(total)==1,'Entidade ou total ambíguo.')
            if len(entity)!=1 or len(total)!=1:raise ValueError('Não foi possível identificar uma entidade e um total únicos.')
            values=[money(v) for v in re.findall(pattern,entity[0])];grand=[money(v) for v in re.findall(pattern,total[0])]
            expected=8 if kind=='financial' else 9
            if len(values)!=expected:raise ValueError('Número de colunas financeiras inesperado.')
            check(values==grand,'Linha da entidade diverge do total do relatório.')
            debt,pct=values[-2:];base=values[0]
            if kind=='financial':
                calculated=sum(values[:5])-values[5];reference=re.search(r'Parcelas\s+Até:\s*(\d{2}/\d{2}/\d{4})',text)
            else:
                calculated=values[5]-values[6];reference=re.search(r'Até:\s*(\d{2}/\d{4})',text)
                check(abs(sum(values[:5])-values[5])<=Decimal('.01'),'Total contábil diverge dos componentes.')
            check(abs(calculated-debt)<=Decimal('.01'),'Dívida diverge dos componentes do relatório.')
            check(base>0 and (debt/base*100).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)==pct,'Percentual diverge da base original.')
            if not reference:raise ValueError('Data/período de referência não encontrado.')
            ref=reference[1];month=ref[-4:]+'-'+ref[-7:-5]
            fields=['total','interest','financialDiscount','conditionalDiscount','acceptances','paid','debt','percent'] if kind=='financial' else ['expected','conditionalDiscount','financialDiscount','adjustments','charges','total','paid','debt','percent']
            item=dict(line=1,id=digest[:24],month=month,reference=ref,period='2026 - Ano Letivo',debt=float(debt),percent=float(pct),base=float(base),baseLabel='Total' if kind=='financial' else 'Previsto',sourceValues=dict(zip(fields,map(float,values))),status='ready')
            item[kind+'Percent']=float(pct);out['rows']=[item];out['reconciliation']=dict(item)
            out['warnings']=['Percentual preservado do relatório; denominador: '+item['baseLabel']+'.','O relatório não informa quantidade de alunos ou responsáveis inadimplentes.']
            check(not state.get('delinquency',{}).get(kind+'Source'),'Já há relatório deste indicador. Concilie a referência existente antes de substituir.')
    else:return None
    out['automatic']=len(out['rows']);out['blocked']=bool(out['invalid'])
    return out
