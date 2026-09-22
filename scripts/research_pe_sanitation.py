"""Pesquisa dirigida nas fontes locais, sem modificar originais."""
import hashlib
import json
import re
import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from pe_real import AUDIT, normalized


def main():
    import openpyxl
    from pypdf import PdfReader
    records = json.loads((AUDIT/'registros-locais.json').read_text(encoding='utf-8'))
    targets = [r for r in records if r['sourceRow'] in (275,276,482,697)]
    names = {r['studentKey'] for r in targets} | {'PAULA ARAUJO DIAS','JAILANE RAMISE LIMOEIRO DA CRUZ','ROMILTON SILVA DA COSTA','VERONEIDE MARTINS SILVA'}
    files = [p for p in Path('D:/').iterdir() if p.is_file() and
             re.search(r'descontos.*2027|descontos_2027|funcion.rio|estagi|admiss|folha - abril',normalized(p.name),re.I)
             and p.suffix.lower() in ('.xlsx','.pdf')]
    result=[]
    for p in files:
        entry=dict(path=str(p),sha256=hashlib.sha256(p.read_bytes()).hexdigest(),matches=[],sheets=[])
        if p.suffix.lower()=='.xlsx':
            wb=openpyxl.load_workbook(p,read_only=True,data_only=True)
            for sheet in wb:
                entry['sheets'].append(dict(name=sheet.title,rows=sheet.max_row,columns=sheet.max_column))
                for i,row in enumerate(sheet.values,1):
                    text=normalized(' | '.join(str(v or '') for v in row))
                    hits=[name for name in names if name in text]
                    if hits:entry['matches'].append(dict(sheet=sheet.title,row=i,targets=hits,values=list(row)))
            wb.close()
        else:
            for i,page in enumerate(PdfReader(p).pages,1):
                text=page.extract_text() or ''
                hits=[name for name in names if name in normalized(text)]
                if hits:entry['matches'].append(dict(page=i,targets=hits,text=text))
        result.append(entry)
    prior=[]
    patterns=['PAULA ARAUJO DIAS','JAILANE','ROMILTON','VERONEIDE','EFUND04TD','EFUND06TD','EFUND07TC','EMERE01MB','EMERE01TC','3182019','3195130','4126005','4126007','3149026']
    for p in (ROOT/'output').rglob('*'):
        if p.suffix not in ('.txt','.json','.md') or 'testes' in p.name or AUDIT in p.parents:continue
        text=normalized(p.read_text(encoding='utf-8',errors='replace'))
        hits=[term for term in patterns if term in text]
        if hits:prior.append(dict(path=str(p.relative_to(ROOT)),sha256=hashlib.sha256(p.read_bytes()).hexdigest(),terms=hits))
    (AUDIT/'pesquisa-saneamento.json').write_text(json.dumps(dict(sources=result,previousResults=prior),ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print(json.dumps(dict(sources=len(result),previousResults=len(prior),matches=[dict(path=x['path'],hits=len(x['matches'])) for x in result if x['matches']]),ensure_ascii=True))


if __name__=='__main__':main()
