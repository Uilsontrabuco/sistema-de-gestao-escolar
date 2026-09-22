"""Ponte documental exclusiva dos G2. Não importa nem grava o estado do 7&7."""
import hashlib
import json
import re
import sys
from decimal import Decimal, ROUND_CEILING, ROUND_HALF_UP
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/reconciliacao-g2-pe-documental-15'
PRIOR = ROOT / 'output/reconciliacao-pendencias-pe-2027'
PDF = Path('D:/Orçamento COLEGIO ADVENTISTA DE JUAZEIRO 2027 (1).pdf')
BOOK = Path('D:/Orcamento Escolar - Juazeiro.xlsm')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def money(value):
    return ('R$ ' + f'{Decimal(str(value)):,.2f}').replace(',', 'X').replace('.', ',').replace('X', '.')


def rounded(value):
    return Decimal(value).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)


def ceiling(cost, ticket):
    return int((Decimal(cost) / Decimal(ticket)).to_integral_value(rounding=ROUND_CEILING))


def extract_sources():
    # Somente leitura; a planilha é o modelo histórico 2024, não orçamento 2027.
    from pypdf import PdfReader
    from openpyxl import load_workbook
    import pypdfium2
    OUT.mkdir(parents=True, exist_ok=True)
    sources = {str(p): sha(p) for p in (PDF, BOOK)}
    reader = PdfReader(PDF)
    pages = {str(n): reader.pages[n-1].extract_text() for n in (2, 4, 10, 12)}
    book = load_workbook(BOOK, read_only=True, data_only=False)
    cached = load_workbook(BOOK, read_only=True, data_only=True)
    cells = []
    for row in book['Resumo'].iter_rows(min_row=7, max_row=8, max_col=15):
        for cell in row:
            if cell.value is not None:
                cells.append(dict(sheet='Resumo', cell=cell.coordinate, formula=cell.value,
                                  format=cell.number_format, cached=cached['Resumo'][cell.coordinate].value))
    years = {name: [dict(cell=c.coordinate, value=c.value) for row in cached[name]
                   for c in row if c.value == 2024] for name in ('Menu', 'CAPA')}
    book.close(); cached.close()
    doc = pypdfium2.PdfDocument(str(PDF))
    for n in (4, 10, 12):
        page = doc[n-1]
        bitmap = page.render(scale=2)
        bitmap.to_pil().save(OUT / f'fonte-pagina-{n}.png')
        bitmap.close(); page.close()
    doc.close()
    if sources != {str(p): sha(p) for p in (PDF, BOOK)}:
        raise AssertionError('Fonte alterada')
    evidence = dict(sourceHashes=sources, pdfPages=pages, historicalWorkbookCells=cells,
                    historicalYearCells=years, historicalWorkbookIs2027=False)
    (OUT / 'fontes.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')


def parse_pdf_g2(text, name):
    line = next(line for line in text.splitlines() if line.startswith('Grupo ' + name[1:] + ' '))
    match = re.match(r'Grupo 2 [AB]\s+(\d+)\s+', line)
    fields = line[match.end():].split() if match else []
    if not match or len(fields) != 13:
        raise AssertionError(f'Estrutura da linha documental mudou: {fields}')
    nums = [Decimal(v.replace('.', '').replace(',', '.')) if v != '-' else Decimal(0)
            for v in fields]
    keys = ['tuitionRevenue', 'otherRevenue', 'totalRevenue', 'freeTuition', 'commercialDiscount',
            'netRevenuePage4', 'payroll', 'support', 'general', 'loans', 'totalCost', 'printedResult', 'printedPE']
    return dict(students=int(match[1]), **dict(zip(keys, nums)))


def build():
    evidence = json.loads((OUT / 'fontes.json').read_text(encoding='utf-8'))
    prior = json.loads((PRIOR / 'reconciliacao.json').read_text(encoding='utf-8'))
    results = []
    for name in ('G2 A', 'G2 B'):
        pdf = parse_pdf_g2(evidence['pdfPages']['4'], name)
        current = next(r for r in prior['fullAudit']['classes'] if r['class'] == name)
        target = next(r for r in prior['targetClasses'] if r['className'] == name)
        parts = {p['component']: Decimal(p['monthlyCents'])/100 for p in target['components']}
        pdf_ticket = pdf['netRevenuePage4'] / pdf['students']
        actual_ticket = Decimal(str(current['netTicket']))
        pdf_cost = pdf['totalCost']
        current_cost = Decimal(str(current['totalCostMonthly']))
        additions = parts['auxiliares-segmento'] + parts['coordenacao-segmento'] + parts['pessoal-institucional']
        pcld = Decimal(str(current['pcldNeutralizedMonthly']))
        residuals = parts['residual-g2'] + parts['apoio-direto-g2']
        assert parts['docentes'] + parts['residual-g2'] == pdf['payroll']
        assert parts['estagiarias'] + parts['apoio-direto-g2'] == pdf['support']
        assert pdf['payroll'] + pdf['support'] + pdf['general'] == pdf_cost
        assert pdf_cost - pcld + additions == current_cost
        assert sum(parts.values()) == current_cost
        assert current_cost - pdf_cost == Decimal('5178.48')
        gross = Decimal(str(current['grossTuition']))
        discounted = gross * Decimal('.97')
        stages = [
            ('PDF: receita líquida p. 4 por aluno, incluindo outras receitas', pdf_cost, pdf_ticket),
            ('Mesmos custos; somente mensalidade após 3%, sem outras receitas', pdf_cost, discounted),
            ('Mesmos custos; ticket 7&7 após 3% e 4,5%', pdf_cost, actual_ticket),
            ('Neutralização gerencial da PCLD; base agregada original preservada', pdf_cost-pcld, actual_ticket),
            ('Acrescentar auxiliares do segmento e coordenação: pendência de cobertura', pdf_cost-pcld+parts['auxiliares-segmento']+parts['coordenacao-segmento'], actual_ticket),
            ('Acrescentar institucional: pendência nominal e de cobertura', current_cost, actual_ticket),
        ]
        stages = [dict(step=label, cost=cost, ticket=ticket, quotient=cost/ticket, ceiling=ceiling(cost, ticket))
                  for label, cost, ticket in stages]
        results.append(dict(name=name, pdf=pdf, current=current, parts=parts,
            pdfTicket=pdf_ticket, currentTicket=actual_ticket, stages=stages,
            documentaryPE=int(pdf['printedPE']), documentaryReproducedPE=ceiling(pdf_cost, pdf_ticket),
            newDefinitivePE=None, provisionalPE=ceiling(current_cost, actual_ticket),
            costDifference=current_cost-pdf_cost, pendingExistingResiduals=residuals,
            pendingAddedPersonnel=additions, pendingPositionsTotal=residuals+additions,
            rateioOverlapProven=False,
            institutionalWeight=dict(numerator=18, denominator=1103),
            auxiliaryWeight=dict(numerator=18, denominator=196), coordinationDivisor=9))
    return dict(scope=['G2 A', 'G2 B'], classes=results,
                policy='E_SEPARADA_COMO_PENDENCIA_NAO_SUBSTITUI_PE_DOCUMENTAL',
                historicalFormulaConfirmed=True, original2027WorkbookFormulaAvailable=False,
                sources=evidence['sourceHashes'])


def report(data):
    lines = ['# G2: ponte entre o PE documental de 15 e o cenário provisório de 22', '',
        '**Referência documental preservada: G2 A = 15; G2 B = 15. Novo PE definitivo não comprovado.** '
        'O cenário de 22 permanece exclusivamente como simulação com pendências. Nenhum valor foi excluído para produzir 15, '
        'nenhum lançamento foi feito e nenhuma das outras 39 turmas foi modificada.', '',
        '## Fonte e reconstrução da fórmula', '',
        'Fonte primária: **D:/Orçamento COLEGIO ADVENTISTA DE JUAZEIRO 2027 (1).pdf**, '
        'p. 4, “Orçamento Resumido / Ficha Financeira”, linhas Grupo 2 A/B. Parâmetros na p. 2; '
        'mensalidades e inadimplência por G2 na p. 10. As páginas relevantes foram inspecionadas visualmente.', '',
        'O modelo histórico **D:/Orcamento Escolar - Juazeiro.xlsm**, `Resumo!O7` e `O8`, contém '
        '`=IFERROR(B7/H7*M7,0)` e `=IFERROR(B8/H8*M8,0)`, formato numérico `0`. '
        'Ou seja: quantidade de alunos × despesas ÷ receita líquida. `H7/H8` somam outras receitas e deduzem '
        'gratuidade e desconto comercial; não deduzem inadimplência. **É o arquivo histórico de 2024, não a planilha geradora de 2027.** '
        'A fórmula histórica fornece evidência do método e reproduz o PDF 2027, mas a identidade da fórmula interna 2027 não pode ser certificada sem sua planilha.', '',
        'Para **cada G2**, com os valores impressos de 2027:', '',
        '- Receita de mensalidades: R$ 14.890,97 para 16 alunos; desconto comercial de R$ 446,73 (3%).',
        '- Outras receitas: R$ 341,11; gratuidade/convênio: zero impresso.',
        '- Receita líquida p. 4: 14.890,97 + 341,11 − 446,73 = **R$ 14.785,35**.',
        '- Base média p. 4: 14.785,35 ÷ 16 = **R$ 924,084375 por aluno**, incluindo outras receitas.',
        '- Custo: 6.271,93 + 1.520,00 + 5.751,16 = **R$ 13.543,09**.',
        '- PE reproduzido: **16 × 13.543,09 ÷ 14.785,35 = 14,6556855 → 15 alunos**.', '',
        'O formato `0` da planilha histórica mostra o número sem casas decimais; não é uma fórmula de teto. '
        'Neste G2, tanto a exibição inteira histórica quanto o teto matemático resultam em 15. '
        'O 7&7 usa explicitamente teto, isto é, o menor inteiro que cobre o custo.', '',
        'A p. 10 mostra R$ 14.444,24 após desconto comercial, inadimplência de R$ 649,99 (4,5%), '
        'receita efetivamente recebida de R$ 13.794,25 e “Mens. Líquida” de R$ 902,77 **antes da inadimplência**. '
        'Não confundir essa mensalidade com os R$ 924,084375 usados na reconstrução da p. 4. '
        'O ticket atual é `930,69 × 0,97 × 0,955 = 862,1446815 → R$ 862,14`.', '',
        '**Origem dos R$ 341,11 também rastreada:** p. 12, conta 3182019 (Receita Direitos de Propriedade) '
        'R$ 2.805,00 + conta 3195130 (Outras Subvenções) R$ 24.116,03 = R$ 26.921,03/mês. '
        'A p. 4 distribui R$ 341,11 a cada G2; `26.921,03 × 14.890,97 ÷ 1.175.229,48 = 341,10806 → 341,11` '
        'reproduz o rateio pela receita bruta. A fórmula histórica `Resumo!D7:D8` usa esse mesmo direcionador. '
        'Portanto, trata-se de receita institucional alocada, não aumento da mensalidade. '
        'A origem orçamentária está comprovada; as condições de repasse/uso no cenário gerencial não foram localizadas.', '',
        'Precisão: a mensalidade reajustada parte de `853,84 × 1,09 = 930,6856`; '
        'isso explica por que 16 × o valor exibido 930,69 difere R$ 0,07 da receita impressa. '
        'O resultado impresso na p. 4 é R$ 1.242,25, enquanto a subtração dos totais exibidos dá R$ 1.242,26. '
        'A diferença de um centavo é compatível com casas decimais internas; não foi criada despesa para ajustá-la.', '',
        '## Classificação das diferenças', '',
        '**A:** comprovada e deve permanecer no seu escopo; **B:** duplicidade comprovada a retirar; '
        '**C:** pertencimento comprovado a outro centro/turma; **D:** institucional com composição e critério comprovados; '
        '**E:** sem comprovação suficiente, separada como pendência. Comprovação histórica de um funcionário não prova '
        'que seu custo deva ser adicionado aos blocos já atribuídos no orçamento 2027.', '',
        'Não foi comprovada parcela quantificável B ou C nesta reconciliação. Não é correto transformar suspeita de sobreposição em exclusão. '
        'O cálculo por capacidade está comprovado no código, mas isso, isoladamente, não qualifica os R$ 3.301,29 como D.', '']
    for r in data['classes']:
        name=r['name']; p=r['parts']; pdf=r['pdf']; c=r['current']
        teacher=p['docentes']; intern=p['estagiarias']; residual=p['residual-g2']; support=p['apoio-direto-g2']
        intern_name='Joyce dos Santos Pereira' if name=='G2 A' else 'Larissa Lorrana Miranda de Jesus'
        page=26 if name=='G2 A' else 29
        lines += [f'## {name} — os 13 itens solicitados', '',
            '| Item | PDF original | 7&7 provisório atual |', '|---|---|---|',
            '| 1. Mensalidade | R$ 930,69, p. 2/10 | R$ 930,69 |',
            '| 2. Desconto estrutural | 3%; R$ 446,73 na base de 16 alunos | 3%; R$ 27,9207 por aluno, antes do arredondamento final |',
            '| 3. Inadimplência | 4,5%; R$ 649,99 na p. 10; não deduzida na receita da p. 4 | 4,5% após o desconto; R$ 40,6246185 por aluno |',
            '| 4. Ticket/base do PE | R$ 924,084375 com outras receitas; mensalidade líquida p. 10 R$ 902,77 antes de inadimplência | R$ 862,14, somente mensalidade após os dois percentuais |',
            f'| 5. Custo docente | Não discriminado; dentro da folha/encargos de R$ 6.271,93 | {money(teacher)}, grade homologada × 4,5 |',
            f'| 6. Auxiliares/estagiárias | Apoio agregado de R$ 1.520,00; sem nomes na p. 4 | {money(intern)} de {intern_name}; auxiliares do segmento R$ 1.458,24 em rateio separado |',
            f'| 7. Encargos | Incluídos no bloco folha/encargos, sem parcela segregada | Sem valor autônomo comprovado; complemento de {money(residual)} não pode ser chamado integralmente de encargo |',
            f'| 8. Outros diretos | Não discriminados em linha autônoma | Auxiliar exclusiva: R$ 0,00 lançado; complemento de apoio {money(support)} pendente de natureza |',
            '| 9. Rateios | Apoio R$ 1.520,00 e gerais R$ 5.751,16; folha agregada não equivale a lotação direta | Auxiliares R$ 1.458,24 + coordenação R$ 1.106,28 + gerais líquidos R$ 5.063,83; complementos apresentados à parte |',
            '| 10. Institucionais | Sem linha adicional segregada; podem já estar dentro dos blocos | R$ 3.301,29 adicionais na linha do G2, classe E |',
            '| 11. Custo total | R$ 13.543,09 | R$ 18.721,57 |',
            '| 12. Fórmula | 16 × 13.543,09 ÷ 14.785,35 | teto(18.721,57 ÷ 862,14) |',
            '| 13. PE | 14,6556855 → **15 impresso** | 21,7152319 → **22 provisório** |', '',
            'Os itens acima explicam a composição e contêm subtotais; não devem ser somados entre si. A ponte abaixo usa parcelas sem sobreposição.', '',
            f'### {name}: comparação por componentes', '',
            '| COMPONENTE | PDF ORIGINAL | 7&7 ATUAL | DIFERENÇA | ORIGEM DA DIFERENÇA | COMPROVADO? |',
            '|---|---:|---:|---:|---|---|',
            '| Folha + encargos, total | 6.271,93 | 6.271,93 | 0,00 | Mesma linha orçamentária, decomposta abaixo | A agregado; E composição residual |',
            f'| ↳ Docentes capturados | Incluídos, não segregados | {money(teacher)} | Não mensurável isoladamente | Grade homologada; não somar novamente à folha | A grade; comparação nominal 2027 incompleta |',
            f'| ↳ Complemento folha/encargos | Dentro de 6.271,93 | {money(residual)} | 0,00 no agregado | 6.271,93 − {money(teacher)} | E nominal/natureza |',
            '| Apoio, total | 1.520,00 | 1.520,00 | 0,00 | Mesma linha de rateio de apoio | A agregado; E composição residual |',
            f'| ↳ Estágio completo | Incluído, não segregado | {money(intern)} | Não mensurável isoladamente | Folha agosto, p. {page}: {intern_name} | A base histórica/lotação aprovada |',
            f'| ↳ Complemento apoio | Dentro de 1.520,00 | {money(support)} | 0,00 no agregado | 1.520,00 − 838,35 | E nominal/natureza |',
            '| Despesas gerais brutas | 5.751,16 | 5.751,16 | 0,00 | Linha original do G2 preservada | A agregado documental |',
            '| Neutralização PCLD gerencial | Sem estorno segregado | −687,33 | −687,33 | Conta 4124001; critério atual 18/1.103 | A decisão gerencial; incidência exata no bloco G2 original ainda E |',
            '| Auxiliares de segmento adicionais à linha | Não segregados | 1.458,24 | +1.458,24 na apresentação atual | Controle agosto 15.878,57 × 18/196, maior resto | E cobertura/sobreposição com folha/apoio originais |',
            '| Coordenação adicional à linha | Não segregada | 1.106,28 | +1.106,28 na apresentação atual | Controle agosto 9.956,51 ÷ 9 turmas EI, maior resto | E cobertura/sobreposição com folha/apoio originais |',
            '| Institucional adicional à linha | Não segregado | 3.301,29 | +3.301,29 na apresentação atual | Saldo 202.296,04 × 18/1.103, maior resto | E composição nominal e correspondência orçamentária |',
            '| **TOTAL CUSTOS** | **13.543,09** | **18.721,57** | **+5.178,48** | **−687,33 + 1.458,24 + 1.106,28 + 3.301,29** | **Novo total definitivo não comprovado** |',
            '| Mensalidade bruta/aluno | 930,69 | 930,69 | 0,00 | Mesma fonte | A |',
            '| Desconto estrutural | 3% | 3% | 0 p.p. | Mesma fonte | A |',
            '| Outras receitas/aluno na base PE | 21,319375 | 0 | −21,319375 | 341,11 ÷ 16; 7&7 considera só mensalidades | A diferença de método; permanência futura da receita E |',
            '| Inadimplência no ticket do PE | Não deduzida p. 4 | 4,5% | −40,6246185/aluno | Parâmetro p. 2 e valor p. 10 aplicado uma vez | A regra; conciliada com PCLD gerencial |',
            '| **BASE POR ALUNO** | **924,084375** | **862,14** | **−61,944375** | **Outras receitas, inadimplência e precisão dos valores exibidos** | **Bases distintas documentadas** |', '',
            '“Não segregado” não significa custo zero no PDF. Os acréscimos acima descrevem a montagem da linha atual; '
            'não comprovam nova despesa econômica. As linhas com ↳ decompõem o subtotal imediatamente anterior.', '',
            f'### {name}: ponte de resultados', '',
            '| Etapa | Custo | Base por aluno | Quociente | PE pelo teto |', '|---|---:|---:|---:|---:|']
        for s in r['stages']:
            lines.append(f"| {s['step']} | {money(s['cost'])} | {s['ticket']:.6f} | {s['quotient']:.6f} | {s['ceiling']} |")
        lines += ['', 'A sequência isola efeitos matemáticos, não aprova lançamentos. A etapa de 15 após PCLD '
            'ainda usa os blocos agregados do orçamento e não é certificação nominal nova. '
            'O quociente intermediário sem outras receitas é muito próximo de 15; as casas decimais '
            'importam e o teto não deve ser trocado por arredondamento para obter um resultado desejado.', '',
            f'**Pendências segregadas em {name}:** {money(r["pendingExistingResiduals"])} dentro dos blocos originais '
            f'+ {money(r["pendingAddedPersonnel"])} de parcelas pessoais acrescidas à linha = '
            f'{money(r["pendingPositionsTotal"])} de posições que exigem vínculo/composição/cobertura. '
            'Esse total não é aumento líquido nem montante autorizado a excluir; mistura residual original e novos rateios. '
            'Não foi usado para emitir novo PE definitivo.', '']
    lines += ['## Origem dos rateios e passagem agosto → 2027', '',
        '- **Folha R$ 6.271,93 / apoio R$ 1.520,00:** valores de 2027 já impressos antes desta reconciliação. '
        'A planilha histórica `Resumo!I7:J8` mostra rateios de folha por receita, incluindo folha administrativa e encargos; '
        'não prova pessoas exclusivas do G2. A vinculação nominal de 2027 permanece ausente.',
        '- **Estágios:** Joyce (folha p. 26) e Larissa (p. 29), cada uma: bolsa 10361 R$ 759,00 + seguro 95006 '
        'R$ 19,35 + consultoria 95014 R$ 60,00 = R$ 838,35. A passagem de bolsa R$ 750,00 para custo completo '
        'R$ 838,35 reduziu o complemento de apoio de R$ 770,00 para R$ 681,65; impacto no total R$ 1.520,00: **zero**.',
        '- **Docentes:** G2 A: Jose Rodrigo R$ 210,60 + Kay R$ 72,90 + Patiara R$ 2.114,10 = R$ 2.397,60; '
        'G2 B: Danielly R$ 72,90 + Jose Rodrigo R$ 187,47 + Monise R$ 2.114,10 = R$ 2.374,47. '
        'A diferença R$ 23,13 é compensada pelo complemento: folha total idêntica, R$ 6.271,93 em ambos. '
        'Nenhum fator de encargos adicional foi aplicado sobre esses complementos.',
        '- **Auxiliares:** controle agosto R$ 15.878,57, formado por Gabriela, Joane, Katia, Maria Eliane, Roberlania '
        '(R$ 2.638,43 cada) e Tamara (R$ 2.686,42). Código vigente usa 18/196 vagas da EI. '
        'R$ 1.458,24 por G2 tem rastreio histórico e algorítmico; falta provar abatimento equivalente no bloco original de pessoal do G2.',
        '- **Coordenação:** Edvirgens R$ 5.222,69 + Valeria R$ 4.733,82 = R$ 9.956,51 históricos; '
        'divisão pelas nove turmas da EI resulta em R$ 1.106,28 por G2 após distribuição de centavos. '
        'Falta a mesma ponte para orçamento 2027. A tabela nominal/células/páginas permanece em '
        '`../reconciliacao-pendencias-pe-2027/NOMINAL-AGOSTO.md`.',
        '- **Institucional:** R$ 387.591,93 de pessoal − R$ 185.295,89 já atribuídos = R$ 202.296,04. '
        'O algoritmo distribui esse saldo pelo maior resto, não por arredondamento independente de cada turma: '
        '202.296,04 × 18/1.103 = 3.301,2953037; atribuição registrada R$ 3.301,29. '
        'Rastrear a fórmula não comprova as pessoas nem que esse custo estivesse ausente dos R$ 7.791,93 de folha/apoio do G2.',
        '- **PCLD:** conta orçamentária R$ 42.117,80; modelo gerencial distribui a neutralização por capacidade '
        '18/1.103 → R$ 687,33 por G2. Como os gerais G2 permaneceram na distribuição original por receita, '
        'falta memória que assegure a incidência exata dessa conta na linha original G2. Não foi alterado o estorno vigente nesta execução.', '',
        'O orçamento 2027 foi emitido em **30/07/2026**, antes da folha de agosto. '
        'Assim, a simples existência de um custo na folha posterior não comprova que foi acrescido ou já estava previsto. '
        'A soma global conciliada não resolve a alocação por turma. Há risco de sobreposição, mas nenhum valor foi classificado B sem prova.', '',
        '## Resposta objetiva', '',
        '| Turma | PE documental | PE comprovado após reconciliação | PE provisório com pendências | Diferença em R$ de custo/mês |',
        '|---|---:|---|---:|---:|',
        '| G2 A | 15 | 15 reproduzido no cenário do PDF; novo PE definitivo não comprovado | 22 | +5.178,48 (18.721,57 − 13.543,09) |',
        '| G2 B | 15 | 15 reproduzido no cenário do PDF; novo PE definitivo não comprovado | 22 | +5.178,48 (18.721,57 − 13.543,09) |', '',
        'Para distinguir diferença de custo e falta de receita: com 15 alunos, o cenário atual arrecada '
        '15 × R$ 862,14 = R$ 12.932,10 e apresenta insuficiência de **R$ 5.789,47** por G2. '
        'Isso não é o aumento de custos de R$ 5.178,48 nem uma prova de que o PE definitivo seja 22.', '',
        '## Documento/informação que ainda falta', '',
        '**Planilha geradora do orçamento 2027 que originou este PDF, com fórmulas e memória nominal de rateios por conta, '
        'funcionário e centro de custo**, particularmente as células equivalentes a `Resumo!B7:O8` e suas dependências. '
        'Ela deve demonstrar a composição dos R$ 6.271,93 de folha e R$ 1.520,00 de apoio, a inclusão/compensação dos '
        'auxiliares e coordenação de agosto, a parcela institucional e a PCLD já contidas nos G2. '
        'Uso: classificar definitivamente cada pendência como custo que permanece, duplicidade ou outro centro; '
        'confirmar a fórmula interna do PE 2027 e a precisão dos números impressos.', '',
        '**Condições de repasse/uso das receitas 3182019 e 3195130 no cenário gerencial.** '
        'A composição dos R$ 341,11 já foi identificada; falta decidir sua disponibilidade e permanência '
        'para cobertura do custo escolar, além da alocação no orçamento. '
        'Sem essa informação, a comparação de 15 com 22 mistura bases de receita diferentes.', '',
        '## Verificação e preservação', '',
        'Ver `TESTES.md` e `integridade.json`. Fontes e arquivos anteriores foram verificados por hash. '
        'Esta análise é aditiva e contém somente os dois G2; não modifica código de cálculo existente, banco, '
        'cadastro, capacidades ou relatórios anteriores. Nenhum desconto foi importado; sem deploy, push, '
        'acesso ao Supabase remoto ou alteração de produção.', '']
    return '\n'.join(lines)


def main():
    if '--extract' in sys.argv:
        extract_sources()
    data = build()
    (OUT / 'reconciliacao.json').write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    (OUT / 'RELATORIO.md').write_text(report(data), encoding='utf-8')
    print(json.dumps([dict(turma=r['name'],documental=r['documentaryPE'],novoDefinitivo=r['newDefinitivePE'],
                          provisorio=r['provisionalPE'],diferenca=str(r['costDifference'])) for r in data['classes']]))


if __name__ == '__main__':
    main()
