"""Importação exclusivamente local/temporária, com documentos originais intactos."""
import sys, json, hashlib
from collections import Counter, defaultdict
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from pe_real import import_rows, calculate, AUDIT, normalized


def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def write(name, value):
    (AUDIT/name).write_text(json.dumps(value,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
def br(cents):
    return 'PENDENTE' if cents is None else f'{cents/100:,.2f}'.replace(',','X').replace('.',',').replace('X','.')
def val(x): return 'PENDENTE' if x is None else str(x)


def main():
    import openpyxl
    from pypdf import PdfReader
    AUDIT.mkdir(parents=True,exist_ok=True)
    primary=Path('D:/Descontos_2027_CAJ_Progressao_Series_Atualizada.xlsx')
    comparison=Path('D:/Descontos_2027_CAJ_Todos_Descontos 2027 OFOCIAL.xlsx')
    a=openpyxl.load_workbook(primary,data_only=True)
    b=openpyxl.load_workbook(comparison,data_only=True)
    raw=list(a['Todos os Descontos'].values)[2:]
    other=list(b['Todos os Descontos'].values)[2:]
    records=import_rows(raw,primary.name)
    write('registros-locais.json',records)
    result=calculate(records)
    manifest=[]
    sources=[primary,comparison,Path('D:/Relação CAJ - Agosto 2026.pdf'),Path('D:/Planilha CAJ.xlsx'),
        Path('D:/Relação Estagiárias OF.xlsx'),Path('D:/Relação Estagiárias.xlsx'),
        Path('D:/Controle Estagiarios IEL.xlsx'),Path('D:/Status de contrato - Estagiárias.xlsx'),
        Path('D:/IEL - Estagiarios CAJ (5).pdf'),Path('D:/LISTAS DE MONITORAS DE CLASSE E ESTAGIÁRIAS.pdf'),
        Path('D:/Orçamento COLEGIO ADVENTISTA DE JUAZEIRO 2027 (1).pdf'),ROOT/'pe_layers_2027_source.json',
        ROOT/'output/reconciliacao-pendencias-pe-2027/reconciliacao.json']
    for p in sources: manifest.append(dict(path=str(p),sha256=sha(p)))
    # Conferir a integridade dos documentos já extraídos e homologados.
    previous=json.loads((ROOT/'output/reconciliacao-pendencias-pe-2027/fontes/manifesto.json').read_text(encoding='utf-8'))
    prior={str(Path(x['path'])):x['sha256'] for x in previous}
    for item in manifest:
        if item['path'] in prior and item['sha256']!=prior[item['path']]:
            raise ValueError('Fonte histórica mudou; não reutilizar extração: '+item['path'])
    supplemental=[]
    for p in sources[6:10]:
        if p.suffix.lower()=='.pdf':
            pages=[x.extract_text() for x in PdfReader(p).pages]
            supplemental.append(dict(source=str(p),pages=pages,conclusion='Lista histórica sem custo de Paula e sem ponte para orçamento 2027.'))
        else:
            wb=openpyxl.load_workbook(p,data_only=True)
            supplemental.append(dict(source=str(p),sheets={s.title:list(s.values) for s in wb.worksheets},conclusion='Controle de situação/contrato sem valores mensais.'))
    write('fontes-adicionais.json',supplemental)
    changed=Counter(tuple(i for i in range(7) if x[i]!=y[i]) for x,y in zip(raw,other))
    diff=dict(primary=str(primary),comparison=str(comparison),primaryRecords=len(records),
        comparisonRecords=sum(bool(r[3]) for r in other),blankRows=sum(not any(v is not None for v in r) for r in raw),
        changedColumns={str(k):v for k,v in changed.items()},
        conclusion='1016 linhas mantêm aluno/benefício e mudam série/turma. A base atualizada retira 30 concluintes; nenhuma segunda progressão aplicada.',
        removedRows=[i+3 for i,(x,y) in enumerate(zip(raw,other)) if not x[3] and y[3]],
        summaryUsage='Resumo não recalculado: suas categorias divergem da base nominal. Usada somente Todos os Descontos, F/G, com validação.')
    write('comparacao-fontes.json',diff)
    # Não se força associação contábil por semelhança de percentual.
    benefits=defaultdict(lambda:dict(count=0,discountCents=0,rows=[]))
    for s in records:
        if s['benefitValid'] and s['discountCents'] is not None:
            key=s['benefit']
            benefits[key]['count']+=1; benefits[key]['discountCents']+=s['discountCents']; benefits[key]['rows'].append(s['sourceRow'])
    prior_cost=json.loads((ROOT/'output/reconciliacao-pendencias-pe-2027/reconciliacao.json').read_text(encoding='utf-8'))
    bridge=dict(benefits=dict(benefits),accounts=prior_cost['discounts'],commercialAccount=prior_cost['commercialDiscountRevenueAccount'],
        matchedCents=None,matchedMeaning='Nenhum vínculo por evento com conta contábil certificado; não é afirmação de sobreposição zero.',
        treatment='Benefícios válidos reduzem receita uma vez. 4126005/4126007 continuam fora do custo comprovado, dentro da exposição pendente histórica. Não há neutralização adicional nem rateio novo.',
        status='PARCIALMENTE COMPROVADO / POSSÍVEL DUPLICIDADE',
        grossDifferenceAgainstFinancialAccountsCents=sum(v['discountCents'] for v in benefits.values())-17790108,
        differenceMeaning='Comparação de universos distintos, NÃO valor conciliado nem saldo a distribuir.')
    bridge['scholarshipAccounts']=[
        dict(code='3149111',description='Bolsas Func/Dep.100%',monthlyCents=5771994),
        dict(code='3149114',description='Bolsas Func/Dep.Parcial',monthlyCents=621262),
        dict(code='3149121',description='Bolsas 100%-Lei 12101',monthlyCents=18427997),
        dict(code='3149122',description='Bolsas 50%-Lei 12101',monthlyCents=1413950)]
    bridge['scholarshipSource']='Orçamento 2027 (1), página 12; contrapartidas de receita 3141111/3141114/3141121/3141122 não somadas como despesas.'
    # CEBAS tem vínculo de natureza, sem afirmar identidade de eventos/valores.
    bridge['scholarshipComparisons']=[]
    for rate,account,budget in [('1','3149121',18427997),('0.5','3149122',1413950)]:
        subset=[r for r in records if r['benefitValid'] and r['benefit']=='CEBAS - Ebolsa' and r['rate']==rate]
        amount=sum(r['discountCents'] for r in subset)
        bridge['scholarshipComparisons'].append(dict(benefit='CEBAS '+str(int(float(rate)*100))+'%',
            candidateAccount=account,individualCents=amount,budgetCents=budget,differenceCents=amount-budget,
            status='PARCIALMENTE COMPROVADO',basis='Natureza legal compatível; identidade individual e competência não conciliadas.'))
    write('ponte-descontos.json',bridge)
    result.update(source=diff,manifest=manifest,discountBridge=bridge,
        exceptions=dict(unmapped=dict(Counter(s['sourceClass'] for s in records if s['classId'] is None)),
            invalidBenefits=[dict(row=s['sourceRow'],sourceClass=s['sourceClass'],issues=s['issues']) for s in records if not s['benefitValid']],
            duplicateStudents=sum('ALUNO DUPLICADO/HOMÔNIMO SEM MATRÍCULA' in s['issues'] for s in records)))
    result['summary']['allValidatedWorkbookDiscountCents']=sum(v['discountCents'] for v in benefits.values())
    write('resultado.json',result)
    write('fontes.json',manifest)
    render_report(result)
    print(json.dumps(result['summary'],ensure_ascii=True))


def render_report(data):
    s=data['summary']
    text=['# PE REAL 2027 — RESULTADO DA RECONCILIAÇÃO','',
        'Importação na auditoria local concluída. Não houve gravação no cadastro escolar, produção ou Supabase. Indicador principal: PE real em alunos, calculado com a distribuição individual validada. **Nenhum PE foi declarado definitivo:** custos completos e parte da distribuição ainda carecem de documentação.','',
        f"Base atualizada: {s['records']} registros; {s['validRecords']} vinculados e válidos; {s['rejectedRecords']} em exceção (incluindo {s['unmappedRecords']} sem turma homologada). Nenhuma segunda progressão. A OFOCIAL conserva a série anterior nos 1.016 registros comuns; não foi misturada à atualizada.",
        'A aba Resumo contém percentuais divergentes dos registros nominais e não foi usada como fonte de cálculo. A comparação e as linhas de origem estão em comparacao-fontes.json e registros-locais.json. A base não contém matrícula única nem declaração de completude; ausência de duplicata por nome não certifica identidade absoluta.','',
        f"Custo mensal considerado no cálculo: **R$ {br(s['consideredCostCents'])}**, integralmente composto pela base comprovada parcial. Exposição de custo pendente: **R$ {br(s['pendingCostCents'])}**, além de Paula sem valor. Não se acrescentou R$ 177.901,08 ao custo nem se aplicou o desconto estrutural de 3%.",
        f"Descontos dos registros vinculados e válidos: R$ {br(s['discountCents'])}; receita líquida: R$ {br(s['netRevenueCents'])}; inadimplência: R$ {br(s['delinquencyCents'])}. Descontos de todos os registros com benefício validado, inclusive turmas não vinculadas: R$ {br(s['allValidatedWorkbookDiscountCents'])}. Valores conflitantes permanecem desconhecidos, portanto esses totais não são o total definitivo da escola.",
        f"PE comprovado: **{s['confirmed']}**; parcialmente comprovado: **{s['partial']}**; em auditoria: **{s['audit']}**; inviabilidade comprovada: **{s['infeasible']}**. PCLD documental: R$ 42.117,80; incluída no custo real: R$ 0,00. Inadimplência de 4,5% aplicada uma vez ao total pós-desconto da turma.",'',
        '## Tabela única — 41 turmas / 1.103 vagas','',
        'Valores em R$/mês. APE e PE usam apenas custo comprovado parcial; não comprovam cobertura integral. N = registros válidos / encontrados. Campos pendentes não são zero. Receita e ticket de turmas com conflito representam somente o subconjunto validado; seu PE principal permanece pendente.','',
        '| Turma | Cap. | Mensalidade | N | Bruta | Descontos | Pós-desconto | Inadimplência | Ticket líquido | Docentes | Aux./estágio | Coord./rateio comprovado | Outros diretos/encargos/operação | Custo parcial | APE parcial | PE real parcial | PE % | Margem física | Status |',
        '|---|---:|---:|---|---:|---:|---:|---:|---:|---:|---:|---:|---|---:|---:|---:|---:|---:|---|']
    for r in data['rows']:
        values=[r['name'],r['capacity'],br(r['tuitionCents']),f"{r['acceptedStudentCount']}/{r['sourceStudentCount']}"]
        values += [br(r[k]) for k in ('grossCents','discountCents','postDiscountCents','delinquencyCents','ticketCents','teacherCents','assistantsInternsCents','coordinationCents')]
        values += ['PENDENTE',br(r['verifiedCostCents']),f"{r['ape']:.4f}",val(r['pe']),f"{r['percentCapacity']:.2f}" if r['pe'] is not None else 'PENDENTE',val(r['physicalMargin']),r['status']]
        text.append('| '+' | '.join(map(str,values))+' |')
    text += ['', '## G2 como controle','',
        'G2 A e G2 B não existem na distribuição atualizada. Os 16 alunos de cada antiga G2 estão identificados como G3 A/B. Usar o arquivo anterior para lhes atribuir desconto G2 reverteria a progressão autorizada. **PE REAL G2 A/B: PENDENTE por ausência de mix.** Não é 0, 15, 22 ou 7. APE sobre custo parcial: vide tabela. A referência documental 15 e a simulação histórica 22 permanecem apenas na memória. Custos parciais: G2 A R$ 5.800,47 e G2 B R$ 5.777,34. Joyce/Larissa e critérios homologados preservados.','',
        '## Método e comprovação financeira','',
        'Cada desconto = mensalidade oficial × percentual original, arredondado em centavos por aluno (half-up). Soma das receitas individuais, depois × 0,955 e arredondamento por turma. Ticket = receita líquida / alunos válidos. PE = teto(custo comprovado × N / receita líquida), sem arredondar o ticket antes do teto. APE = custo / (mensalidade oficial × 0,955), sem teto, com referência explicitamente sem benefício. O mix afeta o ticket; duplicar proporcionalmente uma distribuição não multiplica o PE. Capacidade afeta somente os indicadores físicos.','',
        'Percentuais F/G iguais descrevem o mesmo benefício e são reconhecidos uma vez. Quatro conflitos são excluídos da receita confirmada: bolsa 45 com condicional 100 (dois), bolsa 50 com condicional 45 (um), e bolsa 50 | 45 com condicional 95 (um). Não há regra documental para sua acumulação. Gratuidades integrais consistentes geram receita zero. Categorias e percentuais originais, inclusive frações não arredondadas para percentuais inteiros, são preservados.','',
        'Comprovação de custo exibida = custo nominal comprovado / envelope conhecido (comprovado + pendente histórico). É cobertura da exposição conhecida, não do custo final: Paula e eventuais valores ausentes impedem percentual definitivo. Comprovação de receita = mensalidades brutas dos registros validados / mensalidades brutas dos registros encontrados na turma; mede cobertura financeira da planilha, não cobertura de matrícula. Confiança do PE = PENDENTE: não se inventou probabilidade ou média subjetiva entre coberturas.','',
        '| Turma | Cobertura custo conhecido % | Cobertura financeira registros % | Confiança PE % |','|---|---:|---:|---|']
    for r in data['rows']:
        text.append(f"| {r['name']} | {r['costCoveragePercent']:.2f} | {r['revenueCoveragePercent'] if r['revenueCoveragePercent'] is not None else 'PENDENTE'} | PENDENTE |")
    text += ['', '## Ponte de benefícios e contas','',
        '| Benefício | Valor individual validado | Conta relacionada | Orçamento | Diferença | Tratamento | Status |',
        '|---|---:|---|---:|---|---|---|']
    for label,b in data['discountBridge']['benefits'].items():
        text.append(f"| {label} | {br(b['discountCents'])} | Não identificada por evento | Não atribuível sem duplicar conta | PENDENTE | Reduz receita uma vez | PARCIALMENTE COMPROVADO |")
    text += ['| Contas financeiras agregadas | Não conciliado nominalmente | 4126005 / 4126007 | 48.388,20 / 129.512,88 | PENDENTE | Fora do custo comprovado | POSSÍVEL DUPLICIDADE |',
        '| Desconto comercial | Não separado pela planilha | 3149026 | 34.646,32 | PENDENTE | Sem nova redução estrutural de 3% | SEM CORRESPONDÊNCIA |','',
        f"Diferença aritmética entre todos os benefícios validados e R$ 177.901,08: R$ {br(data['discountBridge']['grossDifferenceAgainstFinancialAccountsCents'])}. Universos distintos: não é conciliação nem custo residual. Valor nominal conciliado com conta específica: **não determinado**. Bolsas legais/filantrópicas, dissídio, institucional e Marcando Vidas permanecem discriminados sem atribuição automática a uma conta.",'',
        '## Custos e documentos cruzados','',
        'Orçamento, folha de agosto, cadastro nominal e relações OF de estagiárias tiveram hashes conferidos contra as extrações anteriores. A memória nominal anterior contém salário, DSR, FGTS, provisões e rubricas excluídas, por pessoa/página; permanece em ../reconciliacao-pendencias-pe-2027/NOMINAL-AGOSTO.md e reconciliacao.json. O motor conserva 50 docentes, 4,5 e lotações/compartilhamentos homologados. Lohana é auxiliar de coordenação; Marcelo não integra docentes. Não se acrescentaram encargos de agosto por cima das provisões ou da projeção docente.','',
        'Fontes adicionais: Controle Estagiarios IEL e Status de contrato não contêm valores; IEL (5) contém contratos históricos 2022–2025; lista de monitoras/estagiárias traz pessoas/lotação histórica incompatíveis com a relação OF atual e não contém valores. Não autorizam custo de Paula, troca nominal de Romilton por Jailane, novo custo de Veroneide ou nova distribuição institucional. As evidências e divergências foram preservadas em fontes-adicionais.json, sem sobrepor homologações.','',
        'Cada parcela comprovada e seu critério estão em resultado.json, rows[].costs; os rateios de auxiliares de segmento e coordenação mantêm seus pesos homologados congelados. Nenhum saldo foi redistribuído. O envelope institucional R$ 202.296,04 segue sem ponte nominal para 2027. As novas planilhas de descontos reduzem a incerteza de receita, mas não comprovam a composição de custos.','',
        '## Bolsas legais e gratuidades — orçamento página 12','',
        'Reduções 3149111 (func./dependentes integral): R$ 57.719,94; 3149114 (parcial): R$ 6.212,62; 3149121 (bolsas legais integral): R$ 184.279,97; 3149122 (50%): R$ 14.139,50. Total: R$ 262.352,03. As contas de serviço educacional correspondentes são apresentação bruta; não se somam essas duas faces como despesas. Dissídio não foi equiparado a dependentes sem vínculo documental. Bolsa Institucional e Marcando Vidas não foram reclassificadas como bolsa legal por inferência.','',
        '| Benefício | Planilha validada | Conta de natureza compatível | Orçamento | Diferença aritmética | Status |','|---|---:|---|---:|---:|---|']
    for c in data['discountBridge']['scholarshipComparisons']:
        text.append(f"| {c['benefit']} | {br(c['individualCents'])} | {c['candidateAccount']} | {br(c['budgetCents'])} | {br(c['differenceCents'])} | {c['status']} |")
    text += ['', 'As diferenças acima comparam populações/competências ainda não identificadas por evento. Não autorizam ajuste de custo ou declaração de conciliação nominal.','',
        '## Exceções rastreáveis','',
        '| Linha da aba Todos os Descontos | Turma escrita | Motivo |','|---:|---|---|']
    for e in data['exceptions']['invalidBenefits']:
        text.append(f"| {e['row']} | {e['sourceClass']} | {', '.join(e['issues'])} |")
    text += ['', 'Códigos não vinculados: '+', '.join(f'{k}: {v} registros' for k,v in data['exceptions']['unmapped'].items())+'. Nenhuma letra foi fundida automaticamente.','',
        '## Pendências que realmente ainda impedem o fechamento','',
        '| Documento/informação | Valor envolvido | Turmas afetadas | Por que é necessário |','|---|---|---|---|',
        '| Memória de custos 2027 por funcionário/rubrica/centro, encargos e critérios, reconciliada com agosto e lotações homologadas | R$ 623.208,36 de exposição conhecida; inclui R$ 177.901,08 tratados na linha de descontos abaixo (não somar) | 41 | Confirmar custos institucionais/operacionais e residuais, sem criar rateio de saldo |',
        '| Custo/inclusão de Paula e evolução dos postos Jailane/Romilton e Veroneide | Não determinado | G3 B e centros dos postos | Concluir posição adicional e mudanças sem presumir salário |',
        '| Regra dos quatro benefícios conflitantes e identificação contábil por evento de descontos/bolsas/comercial | R$ 177.901,08 das contas financeiras + R$ 34.646,32 comercial; valores individuais conflitantes não determinados | Turmas das linhas registradas em exceções e reconciliação geral | Distinguir benefícios, neutralizar apenas sobreposição comprovada e fechar receita individual |',
        '| Distribuição de descontos G2 A/B, G5 C e 8º C; vínculo dos códigos não homologados e confirmação de completude da base | Receita não determinada | Turmas sem base e EFUND04TD, EFUND06TD, EFUND07TC, EMERE01MB/TC | Evitar atribuição de mix de outra turma ou projeção sobre amostra incompleta |',
        '| Ponte financeira das turmas sem correspondência entre orçamento e cadastro | Parcela contida nos custos pendentes, ainda não individualizada | Turmas D, 8º C, 7º C e 3º EM B do PDF | Identificar destino do custo sem copiar ou inventar lotação |','',
        '## Verificação','', 'Contagem e logs da execução final: TESTES.md.']
    (AUDIT/'RELATORIO_PE_REAL_2027.md').write_text('\n'.join(text)+'\n',encoding='utf-8')


if __name__=='__main__': main()
