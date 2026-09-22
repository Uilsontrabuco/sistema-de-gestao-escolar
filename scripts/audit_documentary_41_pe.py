"""Auditoria documental aditiva: nenhum estado/PE existente é escrito."""
import hashlib
import json
import re
import sys
from collections import Counter
from decimal import Decimal, ROUND_HALF_UP, ROUND_CEILING
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'output/auditoria-documental-41-turmas-2027'
PRIOR = ROOT/'output/reconciliacao-pendencias-pe-2027/reconciliacao.json'
PDF = Path('D:/Orçamento COLEGIO ADVENTISTA DE JUAZEIRO 2027 (1).pdf')
BOOK = Path('D:/Orcamento Escolar - Juazeiro.xlsm')
LABEL = r'(Grupo \d [A-C]|\dº Ano [A-D]|\dº E\. Médio [A-C])'
FIN_KEYS = ['grossRevenue','otherRevenue','totalRevenue','freeTuition','commercialDiscount',
            'netRevenue','payroll','support','generalExpenses','loans','totalCost','result','pe']
STUDENT_KEYS = ['tuition','grossRevenue','freePartial','freeFull','commercialDiscount',
                'collectivePartial','collectiveFull','delinquency','netRevenue','receivedRevenue','netTuition']


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def dec(v):
    return Decimal(str(v))


def amount(v):
    return Decimal(0) if v=='-' else Decimal(v.replace('.','').replace(',','.').replace('(','-').replace(')',''))


def money(v):
    return 'ND' if v is None else f'{dec(v):,.2f}'.replace(',', 'X').replace('.', ',').replace('X','.')


def round_int(v):
    return int(v.quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def extract():
    from pypdf import PdfReader
    import pdfplumber
    from openpyxl import load_workbook
    OUT.mkdir(parents=True, exist_ok=True)
    hashes = {str(p):sha(p) for p in (PDF,BOOK)}
    reader = PdfReader(PDF)
    pages = {str(n):reader.pages[n-1].extract_text() for n in (2,4,5,8,9,10,11,12,13,14,15)}
    # Evidência de posição: somente cabeçalho e zeros de subtotais na coluna capacidade.
    capacity = []
    with pdfplumber.open(PDF) as document:
        for n in (8,9):
            words = document.pages[n-1].extract_words()
            column = [w for w in words if 218<w['x0']<272]
            for w in column:
                line = ' '.join(x['text'] for x in words if abs(x['top']-w['top'])<2)
                capacity.append(dict(page=n,text=w['text'],top=w['top'],line=line))
                if w['text'].isdigit() and not ('Subtotal' in line or 'TOTAL' in line):
                    raise AssertionError('Capacidade individual encontrada; revisar leitura')
    book = load_workbook(BOOK,read_only=True,data_only=False)
    formulas = [dict(cell=c.coordinate,formula=c.value,numberFormat=c.number_format)
                for row in book['Resumo'].iter_rows(min_row=7,max_row=54,min_col=15,max_col=15)
                for c in row]
    book.close()
    if hashes != {str(p):sha(p) for p in (PDF,BOOK)}:
        raise AssertionError('Fonte alterada')
    (OUT/'fontes.json').write_text(json.dumps(dict(sourceHashes=hashes,pages=pages,
        capacityColumnEvidence=capacity,historicalPEFormulas=formulas,
        historicalWorkbookYear=2024,original2027WorkbookAvailable=False),ensure_ascii=False,indent=2),encoding='utf-8')


def parse_financial(text):
    result = {}
    for line in text.splitlines():
        match = re.match('^'+LABEL+r'\s+(\d+)\s+(.+)$',line)
        if not match: continue
        name,n,tail = match.groups()
        values = tail.split()
        if len(values)!=13: raise AssertionError((name,values))
        if name in result: raise AssertionError('Linha duplicada')
        result[name] = dict(sourceName=name,students=int(n),capacity=None,
            capacityStatus='COLUNA_CAPACIDADE_EM_BRANCO',sourcePage=4,
            **dict(zip(FIN_KEYS,map(amount,values))))
    if len(result)!=48: raise AssertionError('Exige 48 linhas originais inclusive inativas')
    return result


def parse_students(pages):
    result = {}
    for n in (10,11):
        for line in pages[str(n)].splitlines():
            match = re.match('^'+LABEL+r'\s+(.+)$',line)
            if not match: continue
            name,tail = match.groups()
            values = tail.split()
            if len(values)!=11: raise AssertionError((name,values))
            result[name] = dict(sourcePage=n,**dict(zip(STUDENT_KEYS,map(amount,values))))
    if len(result)!=48: raise AssertionError('Resumo de alunos incompleto')
    return result


def source_name(name):
    if name.startswith('G'): return 'Grupo '+name[1:]
    if name.endswith(' EM'): return name.split()[0]+' E. Médio A'
    ordinal,section = name.split()
    return ordinal+' Ano '+section


def revenue_bridge(pdf,current):
    n=pdf['students']; tuition=dec(current['grossTuition']); ticket=dec(current['netTicket'])
    gross=n*tuition; commercial=gross*Decimal('.03'); delinquency=(gross-commercial)*Decimal('.045')
    parts = dict(
        printedRowPrecision=pdf['grossRevenue']+pdf['otherRevenue']-pdf['freeTuition']-pdf['commercialDiscount']-pdf['netRevenue'],
        tuitionPrecision=gross-pdf['grossRevenue'],
        removeInstitutionalRevenue=-pdf['otherRevenue'],
        removeDocumentaryFreeTuition=pdf['freeTuition'],
        replaceCommercialDiscount=pdf['commercialDiscount']-commercial,
        applyStructuralDelinquency=-delinquency,
        roundCurrentTicket=n*ticket-(gross-commercial-delinquency),
        forecastToCapacity=(current['capacity']-n)*ticket)
    target=ticket*current['capacity']
    if pdf['netRevenue']+sum(parts.values())!=target: raise AssertionError('Ponte receita não fecha')
    return dict(start=pdf['netRevenue'],parts=parts,end=target,currentAtDocumentaryStudents=n*ticket,
                documentAverage=pdf['netRevenue']/n if n else None,currentTicket=ticket)


def cost_bridge(pdf,current,parts):
    personnel=sum(v for k,v in parts.items() if k!='gerais-liquidos-pcld')
    pcld=dec(current['pcldNeutralizedMonthly'])
    general=parts['gerais-liquidos-pcld']+pcld
    delta=dict(printedRowPrecision=pdf['payroll']+pdf['support']+pdf['generalExpenses']+pdf['loans']-pdf['totalCost'],
               personnelReallocation=personnel-pdf['payroll']-pdf['support'],
               generalGrossReallocation=general-pdf['generalExpenses'],
               neutralizePCLD=-pcld,removeLoans=-pdf['loans'])
    if pdf['totalCost']+sum(delta.values())!=sum(parts.values()): raise AssertionError('Ponte custo não fecha')
    return dict(start=pdf['totalCost'],parts=delta,end=sum(parts.values()),
                currentPersonnel=personnel,currentGeneralGross=general,
                institutionalIncludedInPersonnel=parts['pessoal-institucional'],
                nominal2027BridgeVerified=False)


def classify(row):
    # Categorias mutuamente exclusivas. Coincidência de PE não certifica receita/custo.
    if row.get('provenDuplicate'): return 'D'
    if row['document'] is None or row['document']['students']==0: return 'E'
    if row['unprovenReasons']: return 'E'
    if row['costDifference']==0 and row['revenueDifference']==0 and row['peDifference']==0: return 'A'
    if row.get('onlyRoundingVerified'): return 'B'
    return 'C'


def build():
    source=json.loads((OUT/'fontes.json').read_text(encoding='utf-8'))
    prior=json.loads(PRIOR.read_text(encoding='utf-8'))
    current=prior['fullAudit']; pdf=parse_financial(source['pages']['4']); pupils=parse_students(source['pages'])
    for name,p in pdf.items():
        p['studentFinancials']=pupils[name]
        ratio=p['students']*p['totalCost']/p['netRevenue'] if p['netRevenue'] else None
        p['ratio']=ratio
        p['reconstructedDisplayedPE']=round_int(ratio) if ratio is not None else 0
        p['minimumCoveringInteger']=int(ratio.to_integral_value(rounding=ROUND_CEILING)) if ratio is not None else None
        if p['reconstructedDisplayedPE']!=p['pe']: raise AssertionError('PE documental não reproduzido: '+name)
        p['formula']='alunos * custo / receita líquida; exibição sem casas; zero em linha sem base'
        p['formulaEvidence']='Reconstrução 2027 compatível com fórmulas inspecionadas no modelo histórico 2024'
    rows=[]
    for c in current['classes']:
        name=c['class']; label=source_name(name); doc=pdf.get(label)
        parts={p['component']:Decimal(p['monthlyCents'])/100 for p in current['expenseLedger'] if p['classId']==c['classId']}
        if sum(parts.values())!=dec(c['totalCostMonthly']): raise AssertionError('Razão atual diverge')
        reasons=[]
        if parts['pessoal-institucional']>0 and prior['institutionalBridge']['nominalBudget2027ReconciledCents'] is None:
            reasons.append('Composição nominal 2027/cobertura do saldo institucional não comprovada')
        if parts['gerais-liquidos-pcld']>0 and not prior['discountImportGate']['ready']:
            reasons.append('Cobertura das contas de descontos financeiros/condicionais versus receita não comprovada')
        flags=['CUST','REC','INST']
        if not doc:
            reasons.insert(0,'Sem linha correspondente no PDF; desdobramento não documentado financeiramente')
            flags=['SEM LINHA','INST']
        elif doc['students']==0:
            reasons.insert(0,'Linha original zerada; ativação e atribuição de custos ainda sem ponte documental')
            flags=['PDF ZERADO','INST']
        elif doc['freeTuition']:
            flags.append('GRAT')
        if name.startswith('G2 '): flags.append('G2 PRESERVADO')
        r=dict(name=name,currentCapacity=c['capacity'],documentCapacity=None,
            capacitySource='Cadastro homologado do projeto; PDF p. 8/9 sem capacidade preenchida',
            mappedSource=label,mappingType='SERIE_EM_PARA_SECAO_A' if name.endswith(' EM') else 'NOME_EXATO_NORMALIZADO',
            document=doc,current=c,currentComponents=parts,documentaryPE=int(doc['pe']) if doc else None,
            currentPE=c['breakEvenStudents'],peDifference=c['breakEvenStudents']-int(doc['pe']) if doc else None,
            costDifference=dec(c['totalCostMonthly'])-doc['totalCost'] if doc else None,
            revenueDifference=dec(c['capacityRevenueMonthly'])-doc['netRevenue'] if doc else None,
            costBridge=cost_bridge(doc,c,parts) if doc else None,revenueBridge=revenue_bridge(doc,c) if doc else None,
            unprovenReasons=reasons,possibleDuplicateRisk=bool(reasons),provenDuplicate=False,onlyRoundingVerified=False,
            causeFlags=flags,newDefinitivePE=None,
            g2ConclusionPreserved=dict(documentary=15,provisional=22,certifiedNewPE=None) if name.startswith('G2 ') else None)
        r['status']=classify(r)
        rows.append(r)
    matched=[r for r in rows if r['document'] is not None]
    missing=[r['name'] for r in rows if r['document'] is None]
    mapped={r['mappedSource'] for r in rows}
    excluded=[p for name,p in pdf.items() if name not in mapped]
    counts=Counter(r['status'] for r in rows)
    stats=dict(totalCurrentClasses=len(rows),currentCapacity=sum(r['currentCapacity'] for r in rows),
        categories={k:counts[k] for k in 'ABCDE'},numericPEMatches=sum(r['peDifference']==0 for r in matched),
        possibleDuplicateRisk=sum(r['possibleDuplicateRisk'] for r in rows),provenDuplicates=0,
        mappedClasses=len(matched),missingDocumentaryClasses=missing,
        mappedInactive=[r['name'] for r in matched if r['document']['students']==0],
        all41DocumentaryPESum=None,all41CurrentPESum=sum(r['currentPE'] for r in rows),all41ComparableDifference=None,
        mappedDocumentaryPESum=sum(r['documentaryPE'] for r in matched),mappedCurrentPESum=sum(r['currentPE'] for r in matched),
        mappedDifference=sum(r['peDifference'] for r in matched),
        sourceRows=len(pdf),sourceDisplayedPESum=sum(int(p['pe']) for p in pdf.values()),
        sourceDisplayedBelowCoveringInteger=sum(p['students']>0 and p['minimumCoveringInteger']>p['pe'] for p in pdf.values()),
        sourceQuotientSum=sum(p['ratio'] or 0 for p in pdf.values()),sourcePrintedTotalPE=1000,
        sourcePrintedTotalCost=Decimal('841486.96'),sourceSumDisplayedCosts=sum(p['totalCost'] for p in pdf.values()),
        allSourcePEsReproduced=True)
    if len(rows)!=41 or stats['all41CurrentPESum']!=831: raise AssertionError('Checkpoint alterado')
    return dict(scope='AUDITORIA_SOMENTE_LEITURA_SEM_ADOCAO_DE_METODOLOGIA',rows=rows,summary=stats,
        sourceRows=list(pdf.values()),excludedSourceRows=excluded,sourceHashes=source['sourceHashes'],
        existingSnapshotSha256=sha(PRIOR),readyToChangeDefinitivePE=False)


def cause(r):
    if not r['document']: return 'Sem linha no PDF; sem ponte de desdobramento'
    if not r['document']['students']: return 'PDF zerado; turma atual ativa; INST pendente'
    b=r['costBridge']['parts']
    precision=f"precisão {money(b['printedRowPrecision'])}; " if b['printedRowPrecision'] else ''
    return f"Pess {money(b['personnelReallocation'])}; gerais {money(b['generalGrossReallocation'])}; PCLD {money(b['neutralizePCLD'])}; "+precision+'/'.join(r['causeFlags'])


def report(d):
    s=d['summary']
    lines=['# Auditoria documental das 41 turmas — PE 2027', '',
        '**O cenário estrutural atual não reproduz integralmente o cenário documental. Nenhum PE foi alterado.** '
        'G2 A/B permanecem documentalmente em 15; 22 é somente cenário provisório não comprovado. '
        '“7&7 atual” nesta comparação é o checkpoint local da auditoria anterior (soma 831), '
        'não uma leitura de produção nem declaração de adoção definitiva desse cenário.', '',
        '## Fontes, universo e metodologia', '',
        'Fonte primária: `D:/Orçamento COLEGIO ADVENTISTA DE JUAZEIRO 2027 (1).pdf`: '
        'p. 4 (receitas, custos e PE), p. 8/9 (cadastro), p. 10/11 (benefícios, inadimplência e ticket), '
        'p. 2 (parâmetros) e p. 12–15 (contas). Folha/cadastro de agosto e decisões aprovadas permanecem como '
        'evidências históricas; não foram convertidos automaticamente em orçamento nominal 2027.', '',
        '**Capacidade:** a coluna “Cap. P/ Sala” está vazia em todas as linhas individuais do PDF; os zeros são subtotais. '
        'Não se interpretou a quantidade de alunos previstos como capacidade. A coluna CAPACIDADE abaixo usa as '
        '1.103 vagas do cadastro homologado; a capacidade documental PDF está registrada como desconhecida no JSON.', '',
        '**Correspondência:** 38 linhas atuais têm correspondência no PDF, incluindo o 8º C zerado. '
        '1º D, 2º D e 3º D não têm linha original. Os nomes atuais do EM foram associados explicitamente à seção A '
        'de cada série, sem anexar a eles gratuidades de outras seções. O 7º C do PDF não integra o cadastro atual. '
        'Não houve divisão automática dos valores das turmas C entre C/D.', '',
        '**Fórmula reconstruída:** `PE fracionário = alunos previstos × custo total ÷ receita líquida p. 4`, '
        'com exibição sem casas decimais. As 48 linhas do documento, inclusive inativas, foram reproduzidas. '
        'O modelo histórico 2024 `Orcamento Escolar - Juazeiro.xlsm`, `Resumo!O7:O54`, '
        'contém a mesma estrutura `IFERROR(B/H*M,0)` e formato `0`. Isso sustenta a reconstrução; '
        'a planilha geradora de 2027, com fórmulas e precisão interna, continua ausente. Em linha sem alunos/base, '
        'PE impresso zero significa ausência de cenário, não viabilidade gratuita.', '',
        '**PDF:** receita líquida = mensalidades + outras receitas − gratuidades/convênios − desconto comercial. '
        '**Atual:** ticket = mensalidade × 0,97 × 0,955, arredondado em centavos; receita = ticket × capacidade; '
        'PE = teto(custo/ticket). A receita da página 4 não deduz inadimplência, mas a PCLD está nas despesas oficiais. '
        'A página 10/11 apresenta separadamente a receita efetivamente recebida.', '',
        '## Tabela única — 41 turmas', '',
        'Valores mensais em R$. Receita documental é a receita líquida p. 4 para **alunos previstos**; receita 7&7 é '
        'a receita estrutural na **capacidade**. A memória por turma separa o efeito quantidade dos efeitos de ticket. '
        'Diferença de PE = atual − documental. ND = ausência de linha, nunca zero presumido. '
        'Na causa: Pess = mudança total da alocação de pessoal, incluindo institucional; gerais = mudança antes da PCLD; '
        'CUST/REC = bases de custo/receita diferentes; GRAT = benefícios documentais não usados individualmente no cenário estrutural; '
        'INST = composição institucional pendente. Esses códigos não são parcelas adicionais.', '',
        '| TURMA | CAPACIDADE 7&7 | PE DOCUMENTAL | PE 7&7 ATUAL | DIFERENÇA | CUSTO DOCUMENTAL | CUSTO 7&7 | RECEITA DOCUMENTAL | RECEITA 7&7 | CAUSA DA DIFERENÇA | STATUS |',
        '|---|---:|---:|---:|---:|---:|---:|---:|---:|---|---|']
    for r in d['rows']:
        p=r['document']; c=r['current']
        lines.append('| '+' | '.join([r['name'],str(r['currentCapacity']),str(r['documentaryPE']) if p else 'ND',
            str(r['currentPE']),str(r['peDifference']) if p else 'ND',money(p['totalCost']) if p else 'ND',
            money(c['totalCostMonthly']),money(p['netRevenue']) if p else 'ND',money(c['capacityRevenueMonthly']),
            cause(r),r['status']])+' |')
    lines += ['', '## Resultado da classificação', '',
        'A classificação avalia a reprodução integral do cálculo, não apenas a coincidência do inteiro final. '
        'A = reprodução exata; B = somente arredondamento; C = diferença integralmente comprovada; '
        'D = diferença causada por duplicidade comprovada; E = comprovação insuficiente.', '',
        f"- **A: {s['categories']['A']}; B: {s['categories']['B']}; C: {s['categories']['C']}; D: {s['categories']['D']}; E: {s['categories']['E']}.**",
        f"- Coincidência numérica de PE: **{s['numericPEMatches']} turma** (1º C: 21). Seus custos/receitas diferem; por isso não é reprodução exata A.",
        '- Todas as 41 recebem parcela do saldo institucional sem composição nominal 2027 suficiente; '
        'essa pendência impede classificar a diferença total como C, mesmo quando sua aritmética e parte das causas estão demonstradas.',
        f"- **Possível sobreposição a investigar: {s['possibleDuplicateRisk']} turmas**, pelo envelope institucional e/ou descontos financeiros. "
        'São riscos dentro de E, não 41 duplicidades comprovadas nem uma sexta categoria somável. D comprovada: zero.', '',
        '## Somas — sem transformar ausência em zero', '',
        '- **Soma documental das 41: indeterminada**, pois faltam as três linhas D.',
        f"- **Soma documental das 38 correspondências: {s['mappedDocumentaryPESum']}**, incluindo o zero explícito do 8º C.",
        f"- **Soma atual das 41: {s['all41CurrentPESum']}.** Nas mesmas 38 correspondências: **{s['mappedCurrentPESum']}**.",
        f"- **Diferença comparável nas 38: {s['mappedDifference']} alunos** ({s['mappedCurrentPESum']} − {s['mappedDocumentaryPESum']}). "
        'A diferença total das 41 não é determinável. Subtrair 831 de um subtotal documental de 38 turmas mistura universos.',
        f"- No PDF completo: soma dos PEs inteiros das 48 linhas = **{s['sourceDisplayedPESum']}**; "
        f"soma dos quocientes reconstruídos = **{s['sourceQuotientSum']:.6f}**, que se exibe como **1.000**, tal como o total impresso. "
        'São arredondamentos em momentos diferentes, não um aluno a transferir entre turmas.',
        '- O 7º C excluído do cadastro atual tem PE documental **34**. A passagem 1.001 − 34 = 967 '
        'explica o subtotal de correspondências. O PDF também tem 3º EM B sem alunos e com receita líquida negativa '
        'de R$ 9.474,89, proveniente de benefícios registrados; não foi absorvida automaticamente pelo 3º EM atual.', '',
        '## Auditoria dos componentes e representação no motor', '',
        '**Outras receitas:** contas 3182019 R$ 2.805,00 e 3195130 R$ 24.116,03 somam R$ 26.921,03/mês. '
        'O PDF as distribui por receita bruta; o ticket estrutural não as inclui. A soma das parcelas impressas é '
        'R$ 26.921,01, diferença de dois centavos por precisão. A projeção de receita na capacidade não representa esse financiamento institucional. '
        'Não foi incluído sem decisão sobre condições de uso/recorrência.', '',
        '**Gratuidades/convênios:** o PDF deduz R$ 262.352,03/mês, separados na p. 10/11 em bolsas parciais/integrais '
        'e convenção coletiva. O cenário estrutural atual aplica desconto uniforme de 3% e não esses benefícios por turma. '
        'Isso explica parte importante das quedas de PE em turmas com muitas gratuidades. Desconto comercial não equivale '
        'a todas as bolsas. Os valores por turma foram extraídos do PDF, sem abrir nem importar a planilha de descontos reais.', '',
        '**Desconto comercial:** ambos partem de 3%, mas no PDF a base varia com certos benefícios. '
        'Logo, R$ 34.646,32 do orçamento não equivale a 3% do total bruto sem ajustes. A memória calcula a diferença entre '
        'o desconto documental e o desconto estrutural para a mesma quantidade de alunos.', '',
        '**Inadimplência e PCLD:** o PDF registra R$ 42.117,80 de inadimplência p. 11 e a mesma provisão em despesa '
        '4124001. O atual aplica 4,5% no ticket e neutraliza R$ 42.117,80 do custo gerencial. Isso evita repetição global '
        'do mesmo conceito, porém os critérios por turma não são idênticos: inadimplência documental usa a receita líquida '
        'de benefícios, gerais oficiais seguem receita bruta e o estorno atual usa capacidade. Cada memória apresenta '
        'inadimplência documental, estorno atual e a parcela de PCLD estimada pelo direcionador original; '
        'essa última é inferência de rateio, não lançamento nominal comprovado.', '',
        '**Folha/apoio/encargos:** os blocos oficiais R$ 267.630,16 + R$ 119.961,77 = R$ 387.591,93 já contêm pessoal. '
        'O atual reconstrói docentes, estágio, auxiliares e coordenação e distribui o saldo institucional. '
        'O total global preservado não comprova inclusão exclusiva de cada pessoa em cada bloco. '
        'Sem a ponte nominal 2027, reclassificar e depois somar uma parcela à turma pode deslocar custo ou sobrepor '
        'alocação documental, mesmo sem duplicar o total da escola. Não se aplicou novo fator de encargos, DSR ou reajuste.', '',
        '**Auxiliares/coordenação/institucional:** as bases de agosto e os critérios atuais são rastreáveis, '
        'mas não demonstram a formação dos valores de julho/2026 que originaram o orçamento 2027. '
        'Lohana permanece somente coordenação; Marcelo não é projetado como professor. '
        'A falta de custo de Paula em G3 B continua pendente. Nenhuma lotação foi inventada.', '',
        '**Descontos como despesas:** 4126005 R$ 48.388,20 e 4126007 R$ 129.512,88 permanecem nos gerais, '
        'somando R$ 177.901,08. Sua cobertura perante 3149026, gratuidades e futuro desconto real ainda não está '
        'conciliada por evento. A exposição é identificada como possível sobreposição, sem exclusão por suposição.', '',
        '**Calendário:** o PDF apresenta totais anuais como 12 meses; o cenário atual mantém 11 mensalidades e '
        'matrícula de janeiro separada. A presente comparação é mensal. Nem a soma de PEs mensais nem a margem '
        'mensal multiplicada por 12 prova equilíbrio anual no calendário atual.', '',
        '**Arredondamento:** despesas por linha do PDF somam R$ 841.487,00 versus total impresso R$ 841.486,96. '
        'As diferenças de centavos foram registradas nas pontes, não distribuídas artificialmente. '
        'O PDF exibe PE sem casas; o atual usa teto. Essa diferença metodológica existe, mas nenhuma das 41 '
        'pode ser classificada B, pois também há diferenças materiais de base/custo e pendências.', '',
        f"Em **{s['sourceDisplayedBelowCoveringInteger']} linhas ativas do PDF**, o PE exibido é menor que o teto do quociente reconstruído. "
        'Logo, a exibição inteira do documento não garante sempre o menor número de alunos que cobre o custo. '
        'Isso é um diagnóstico da metodologia, não autorização para corrigir os PEs documentais.', '',
        '## Evidências/documentos ainda necessários', '',
        '1. **Planilha geradora do orçamento 2027 com fórmulas e memória nominal por funcionário, rubrica e centro/turma.** '
        'Uso: confirmar bases e precisão; conciliar folha, apoio, auxiliares, coordenação, encargos/isenções, institucional '
        'e distribuição da PCLD sem sobreposição. Deve incluir a previsão de Paula e as substituições já aprovadas.',
        '2. **Ponte financeira aprovada entre o cadastro do PDF e as 41 turmas:** abertura de 1º D, 2º D e 3º D, '
        'ativação de 8º C, destino de 7º C e tratamento do 3º EM B com benefícios sem alunos. '
        'Uso: tornar comparáveis os totais, sem dividir, absorver ou zerar custos por inferência.',
        '3. **Condições de repasse/uso das receitas 3182019 e 3195130.** '
        'Uso: decidir se e como direitos de propriedade/subvenções financiam o PE gerencial por turma.',
        '4. **Memória de cobertura dos benefícios das contas 4126005/4126007 em relação a 3149026 e bolsas/convênios.** '
        'Uso: impedir dupla contagem entre despesas e redução de receita antes de qualquer futura importação.', '',
        'Não é necessário rediscutir as capacidades já homologadas para reproduzir o PDF: ele simplesmente não as informa. '
        'A escolha entre PE orçamentário com benefícios e receitas institucionais ou PE estrutural por capacidade '
        'é posterior à auditoria; nenhuma metodologia foi adotada nesta execução.', '',
        '## Artefatos e validação', '',
        '`MEMORIA.md` contém os valores extraídos, fórmulas e pontes de cada turma; `auditoria.json` contém os dados '
        'e campos numéricos/pendentes, inclusive as linhas originais sem correspondência. '
        '`TESTES.md` e `integridade.json` registram validação e preservação das 338 verificações anteriores. '
        'Não houve mudança de PE, importação de descontos, deploy, push, acesso remoto ou alteração de produção.', '']
    return '\n'.join(lines)


def memory(d):
    lines=['# Memória documental por turma', '',
        'Valores mensais. Pontes aritméticas explicam divergências; não autorizam custos nem corrigem o documento. '
        'Capacidade PDF desconhecida em todas as linhas. Fórmula com base zero não certifica PE econômico zero.', '']
    for r in d['rows']:
        p=r['document']; c=r['current']; parts=r['currentComponents']
        lines += [f"## {r['name']}", '',f"Correspondência: {r['mappedSource']}; capacidade atual {r['currentCapacity']}; capacidade PDF ND."]
        if p:
            f=p['studentFinancials']
            lines += [f"PDF p. 4: alunos {p['students']}; mensalidades {money(p['grossRevenue'])}; outras receitas {money(p['otherRevenue'])}; "
                f"gratuidades/convênios {money(p['freeTuition'])}; desconto comercial {money(p['commercialDiscount'])}; receita líquida {money(p['netRevenue'])}.",
                f"Custos PDF: folha/encargos {money(p['payroll'])}; apoio {money(p['support'])}; gerais {money(p['generalExpenses'])}; "
                f"empréstimos {money(p['loans'])}; total {money(p['totalCost'])}; resultado impresso {money(p['result'])}.",
                f"PDF p. {f['sourcePage']}: mensalidade {money(f['tuition'])}; bolsas parciais {money(f['freePartial'])}; integrais {money(f['freeFull'])}; "
                f"convenção parcial {money(f['collectivePartial'])}; integral {money(f['collectiveFull'])}; inadimplência {money(f['delinquency'])}; "
                f"receita após benefícios {money(f['netRevenue'])}; recebida {money(f['receivedRevenue'])}; mensalidade líquida exibida {money(f['netTuition'])}.",
                f"PE PDF: {p['students']} × {money(p['totalCost'])} ÷ {money(p['netRevenue'])} = "
                f"{str(p['ratio']) if p['ratio'] is not None else 'sem base (tratamento zero da linha inativa)'}; "
                f"exibido/reproduzido {p['reconstructedDisplayedPE']}; impresso {int(p['pe'])}; teto contrafactual {p['minimumCoveringInteger']}."]
        else: lines += ['**Sem linha no PDF: receitas, despesas e PE documental desconhecidos. Não foi criado PE zero.**']
        lines += [f"Atual: mensalidade {money(c['grossTuition'])} × 0,97 × 0,955 → ticket {money(c['netTicket'])}; "
            f"receita na capacidade {money(c['capacityRevenueMonthly'])}; custo {money(c['totalCostMonthly'])}; "
            f"PE = teto({money(c['totalCostMonthly'])}/{money(c['netTicket'])}) = {c['breakEvenStudents']}.",
            'Parcelas atuais sem sobreposição na razão: '+ '; '.join(f'{k}: {money(v)}' for k,v in parts.items())+'.']
        if p:
            cb=r['costBridge']; rb=r['revenueBridge']
            lines += ['Ponte custo: '+money(cb['start'])+' + ['+'; '.join(f'{k} = {money(v)}' for k,v in cb['parts'].items())+'] = '+money(cb['end'])+'.',
                'Ponte receita: '+money(rb['start'])+' + ['+'; '.join(f'{k} = {v}' for k,v in rb['parts'].items())+'] = '+money(rb['end'])+'.',
                f"Controle de quantidade: receita atual aos mesmos {p['students']} alunos do PDF = {money(rb['currentAtDocumentaryStudents'])}; "
                f"efeito adicional da troca para capacidade = {money(rb['parts']['forecastToCapacity'])}.",
                f"PCLD: inadimplência documental {money(p['studentFinancials']['delinquency'])}; estorno atual {money(c['pcldNeutralizedMonthly'])}; "
                f"estimativa pelo direcionador original de gerais (42.117,80 × receita bruta / 1.175.229,48) = "
                f"{money(Decimal('42117.80')*p['grossRevenue']/Decimal('1175229.48'))}. Esta estimativa não é verba nominal."]
        lines += ['Status '+r['status']+': '+'; '.join(r['unprovenReasons'])+'. Novo PE definitivo: não emitido.', '']
    lines += ['## Linhas do PDF sem correspondência no cadastro atual', '']
    for p in d['excludedSourceRows']:
        lines.append(f"- {p['sourceName']}: alunos {p['students']}, receita líquida {money(p['netRevenue'])}, "
                     f"custo {money(p['totalCost'])}, PE {int(p['pe'])}; não redistribuído.")
    return '\n'.join(lines)+'\n'


def main():
    if '--extract' in sys.argv: extract()
    data=build()
    (OUT/'auditoria.json').write_text(json.dumps(data,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    (OUT/'RELATORIO.md').write_text(report(data),encoding='utf-8')
    (OUT/'MEMORIA.md').write_text(memory(data),encoding='utf-8')
    print(json.dumps(data['summary'],ensure_ascii=False,default=str))


if __name__=='__main__': main()
