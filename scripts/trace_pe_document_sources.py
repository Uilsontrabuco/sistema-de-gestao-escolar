"""Extração somente leitura de fontes locais já referidas no projeto CAJ."""
import hashlib
import json
from pathlib import Path
from pypdf import PdfReader
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/reconciliacao-pendencias-pe-2027/fontes'
NAMES=['Relação CAJ - Agosto 2026.pdf','Planilha CAJ.xlsx',
       'Relação Estagiárias OF.xlsx','Relação Estagiárias.xlsx',
       'Orçamento COLEGIO ADVENTISTA DE JUAZEIRO 2027 (1).pdf',
       'Orçamento COLEGIO ADVENTISTA DE JUAZEIRO 2027 (2).pdf',
       'Orçamento COLEGIO ADVENTISTA DE JUAZEIRO 2027 (3).pdf',
       'Orçamento COLEGIO ADVENTISTA DE JUAZEIRO 2027.pdf']


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    manifest=[]
    seen={}
    for name in NAMES:
        path=Path('D:/')/name
        if not path.exists():
            manifest.append(dict(path=str(path),status='AUSENTE'));continue
        sha=hashlib.sha256(path.read_bytes()).hexdigest()
        info=dict(path=str(path),sha256=sha)
        if sha in seen:
            info['identicalTo']=seen[sha];manifest.append(info);continue
        seen[sha]=name
        if path.suffix=='.pdf':
            reader=PdfReader(path)
            info['pages']=len(reader.pages)
            text='\n\n'.join(f'PÁGINA {i+1}\n{p.extract_text() or ""}' for i,p in enumerate(reader.pages))
            (OUT/(name+'.txt')).write_text(text,encoding='utf-8')
            if name == 'Relação CAJ - Agosto 2026.pdf':
                layout='\n\n'.join(f'PAGINA {i+1}\n{p.extract_text(extraction_mode="layout") or ""}' for i,p in enumerate(reader.pages))
                (OUT/'agosto-layout.txt').write_text(layout,encoding='utf-8')
        else:
            book=load_workbook(path,read_only=True,data_only=True)
            sheets=[]
            for sheet in book:
                rows=[dict(row=i,cells={str(j):v for j,v in enumerate(row,1) if v is not None})
                      for i,row in enumerate(sheet.iter_rows(values_only=True),1) if any(v is not None for v in row)]
                sheets.append(dict(sheet=sheet.title,rows=rows))
            (OUT/(name+'.json')).write_text(json.dumps(sheets,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
            info['sheets']=[dict(name=s['sheet'],nonemptyRows=len(s['rows'])) for s in sheets]
            book.close()
        if hashlib.sha256(path.read_bytes()).hexdigest()!=sha:raise AssertionError('Fonte mudou')
        manifest.append(info)
    (OUT/'manifesto.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
