"""Integração HTTP com banco e contas temporários, nunca na base escolar."""
import base64,copy,http.cookiejar,io,json,sys,tempfile,threading,unittest,urllib.request,urllib.error
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from server import serve,Handler
from services import financial_import_workbook
from types import SimpleNamespace

class HttpTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.server=serve(Path(self.tmp.name)/'http.sqlite3',port=0);self.admin=self.server.store.create_user('Fixture HTTP','http@fixture.invalid','Test-only-password-2026',is_admin=True);self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start();self.base='http://127.0.0.1:'+str(self.server.server_port);self.client=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()));self.csrf='';_,body=self.request('/api/login',{'email':'http@fixture.invalid','password':'Test-only-password-2026'});self.csrf=json.loads(body)['csrf']
    def tearDown(self):self.server.shutdown();self.server.server_close();self.thread.join();self.tmp.cleanup()
    def request(self,path,payload=None,headers=None):
        req=urllib.request.Request(self.base+path,data=json.dumps(payload).encode() if payload is not None else None,headers={'Content-Type':'application/json','X-CSRF-Token':self.csrf,**(headers or {})})
        try:
            with self.client.open(req,timeout=10) as response:return response.status,response.read()
        except urllib.error.HTTPError as error:return error.code,error.read()
    def test_auth_csrf_and_private_files(self):
        self.assertEqual(self.request('/data/caj.sqlite3')[0],404);self.assertEqual(self.request('/server.py')[0],404);self.assertEqual(self.request('/api/state',{'changes':{},'version':0,'module':'budget'},headers={'X-CSRF-Token':'wrong'})[0],403);self.assertEqual(self.request('/api/users',{},headers={'Origin':'https://untrusted.invalid'})[0],403)
    def test_teaching_cost_read_only_and_permission(self):
        before=self.server.store.state()
        code,body=self.request('/api/teaching-load/costs')
        self.assertEqual(code,200,body)
        data=json.loads(body)
        self.assertEqual(len(data['classes']),41)
        self.assertEqual(data['status'],'FINALIZADO')
        self.assertEqual(data['validated_cost_cents'],2567735)
        self.assertEqual(data['conflicted_cost_cents']+data['unassigned_cost_cents'],0)
        self.assertEqual(len([p for p in data['operational_allocations'] if p['bucket']!='validated']),0)
        self.assertEqual(data['accounting']['additional_expense'],0)
        self.assertEqual(self.server.store.state(),before)
        self.server.store.create_user('Sem acesso','no-finance@fixture.invalid','Test-only-password-2026',{'classes':['view']},actor=self.admin)
        self.client=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
        self.request('/api/login',{'email':'no-finance@fixture.invalid','password':'Test-only-password-2026'})
        self.assertEqual(self.request('/api/teaching-load/costs')[0],401)
    def test_integrated_financial_and_pe_endpoints_read_only(self):
        before=self.server.store.state()
        for path in ('/api/financial/teaching-integration?year=2027','/api/break-even/integrated?year=2027'):
            status,body=self.request(path);self.assertEqual(status,200)
            data=json.loads(body);self.assertEqual(data['teachingCostWeekly'],25677.35)
            self.assertEqual(data['teachingCostMonthly'],115548.16);self.assertEqual(data['monthlyFactor'],4.5);self.assertIsNone(data['breakEvenStudents'])
            self.assertEqual(data['classBreakEvenCounts']['undetermined'],0);self.assertEqual(len(data['classes']),41)
            g2=next(row for row in data['classes'] if row['name']=='G2 A')
            self.assertEqual(g2['teachingCostWeekly'],532.80);self.assertEqual(g2['teachingCostMonthly'],2397.60)
            self.assertEqual(g2['breakEvenStudents'],3);self.assertEqual(g2['physicalMarginStudents'],15)
        self.assertEqual(self.server.store.state(),before)
        self.assertEqual(self.request('/api/financial/teaching-integration?year=2028')[0],400)
    def test_sse_windows_disconnect_is_normal(self):
        handler=object.__new__(Handler);handler.path='/api/events';handler.current=lambda:self.admin;handler.token=lambda:'fixture';handler.send_response=lambda *a:None;handler.send_header=lambda *a:None;handler.end_headers=lambda:None
        handler.server=SimpleNamespace(store=SimpleNamespace(session=lambda token:(self.admin,'csrf'),overdue=lambda:None,state=lambda:({},1)))
        class AbortedStream:
            def write(self,data):raise ConnectionAbortedError('Windows client disconnected')
        handler.wfile=AbortedStream();self.assertIsNone(handler.do_GET())
    def test_import_preview_confirm_replay_and_export(self):
        payload={'name':'fixture.csv','kind':'budget','sourceConfirmed':True,'content':base64.b64encode(b'Categoria;Orcado\nFixture;100\n').decode()};code,body=self.request('/api/import/preview',payload);self.assertEqual(code,200,body);preview=json.loads(body);self.assertEqual(self.server.store.state()[0]['budget'],[])
        code,body=self.request('/api/import/confirm',{'token':preview['token'],'confirmation':'CONFIRMAR IMPORTAÇÃO'});self.assertEqual(code,200,body);self.assertEqual(self.server.store.state()[0]['budget'][0]['budget'],100);self.assertEqual(self.request('/api/import/confirm',{'token':preview['token'],'confirmation':'CONFIRMAR IMPORTAÇÃO'})[0],400)
        code,pdf=self.request('/api/export',{'kind':'budget','title':'Teste HTTP','format':'pdf'});self.assertEqual(code,200,pdf);self.assertTrue(pdf.startswith(b'%PDF'))
        code,xlsx=self.request('/api/export',{'kind':'budget','format':'xlsx'});self.assertEqual(code,200,xlsx);self.assertTrue(xlsx.startswith(b'PK'))
    def test_financial_discount_template_download_is_read_only(self):
        before=self.server.store.state();code,xlsx=self.request('/api/financial-classifications/template',{});self.assertEqual(code,200,xlsx);self.assertTrue(xlsx.startswith(b'PK'));self.assertEqual(self.server.store.state(),before)
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(xlsx));self.assertEqual(book.sheetnames,['Nominal','Consolidada por turma','Instruções']);book['Nominal'].append(['ALUNO MODELO FICTÍCIO','G2 A','Grupo 2','Sem desconto',0,'Teste de reimportação']);buffer=io.BytesIO();book.save(buffer);book.close();code,body=self.request('/api/financial-classifications/preview',{'name':'modelo-preenchido.xlsx','year':2027,'content':base64.b64encode(buffer.getvalue()).decode()});self.assertEqual(code,200,body);self.assertEqual(json.loads(body)['automatic'],1);self.assertEqual(self.server.store.state(),before)
    def test_financial_discount_button_to_http_preview_uses_multisheet_reader_without_writes(self):
        before=self.server.store.state();code,interface=self.request('/professional.js');self.assertEqual(code,200);self.assertIn(b'Baixar modelo de planilha',interface);self.assertIn(b"api('financial-classifications/preview'",interface)
        fixture=financial_import_workbook(homologation=True);code,body=self.request('/api/financial-classifications/preview',{'name':'HOMOLOGACAO-FICTICIA-NAO-CONFIRMAR.xlsx','year':2027,'content':base64.b64encode(fixture).decode()});self.assertEqual(code,200,body)
        self.assertNotIn('inadimplência',body.decode().lower());preview=json.loads(body);self.assertEqual(preview['sheet'],'Nominal, Consolidada por turma');self.assertEqual(len(preview['classSummary']),41);self.assertEqual(self.server.store.state(),before)
        g2=next(row for row in preview['classSummary'] if row['className']=='G2 A');g3=next(row for row in preview['classSummary'] if row['className']=='G3 A');self.assertTrue(any(row['percent']==16 and row['quantity']==1 for row in g2['variables']));self.assertEqual(g3['fixed']['markingLives45'],1);self.assertEqual(g3['fixed']['philanthropic50'],1);self.assertEqual(g2['status'],'blocked');self.assertGreaterEqual(preview['pending'],3)
    def test_classes_button_to_real_pdf_preview_has_41_rows_without_writes(self):
        pdf=Path(r'D:\Turmas CAJ 2027 (1).pdf')
        if not pdf.exists():self.skipTest('PDF oficial não disponível neste ambiente')
        before=self.server.store.state();code,interface=self.request('/professional.js');self.assertEqual(code,200);self.assertIn('Importar atualização de turmas'.encode(),interface);self.assertIn(b"api('classes-update/preview'",interface)
        code,body=self.request('/api/classes-update/preview',{'name':pdf.name,'content':base64.b64encode(pdf.read_bytes()).decode()});self.assertEqual(code,200,body);preview=json.loads(body);self.assertEqual(len(preview['rows']),41);self.assertEqual(preview['newTotals'],{'capacity':1103,'students':752,'new':34,'renewed':718,'vacancies':351});self.assertTrue(preview['canConfirm']);self.assertEqual(self.server.store.state(),before)
        code,body=self.request('/api/classes-update/confirm',{'token':preview['token'],'confirmation':'CONFIRMAR ATUALIZAÇÃO DE TURMAS'});self.assertEqual(code,200,body);updated=json.loads(body)['state'];self.assertEqual(sum(x['students'] for x in updated['classes']),752);self.assertEqual(len(updated['classes']),41);self.assertEqual(len([x for x in updated['imports'] if x.get('kind')=='class_update']),1)
        self.assertEqual(self.request('/api/classes-update/confirm',{'token':preview['token'],'confirmation':'CONFIRMAR ATUALIZAÇÃO DE TURMAS'})[0],400)
        code,body=self.request('/api/classes-update/preview',{'name':pdf.name,'content':base64.b64encode(pdf.read_bytes()).decode()});self.assertEqual(code,200,body);self.assertFalse(json.loads(body)['canConfirm']);self.assertTrue(json.loads(body)['duplicateFile'])
    def test_two_users_share_state_and_revocation(self):
        user=self.server.store.create_user('Lívia','livia@fixture.invalid','Test-only-password-2026',{'benefits':['view','create','edit']},actor=self.admin);other=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()));self.client,adminclient=other,self.client;_,body=self.request('/api/login',{'email':user['email'],'password':'Test-only-password-2026'});self.csrf,admincsrf=json.loads(body)['csrf'],self.csrf;_,body=self.request('/api/state');state=json.loads(body);code,body=self.request('/api/state',{'module':'benefits','action':'create','version':state['version'],'changes':{'benefits':[{'id':'b','name':'Fixture','quantity':1,'type':'Fixture','active':True}]}});self.assertEqual(code,200,body)
        self.client=adminclient;self.csrf=admincsrf;_,body=self.request('/api/state');self.assertEqual(json.loads(body)['state']['benefits'][0]['name'],'Fixture');self.request('/api/users',{'action':'edit','user':dict(user,active=False)});self.client=other;self.assertEqual(self.request('/api/state')[0],401)
    def test_livia_quantity_change_is_shared_and_audited(self):
        user=self.server.store.create_user('Lívia','livia@fixture.invalid','Test-only-password-2026',{'enrollments':['view','create','edit'],'classes':['view']},actor=self.admin);adminclient,admincsrf=self.client,self.csrf;self.client=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()));_,body=self.request('/api/login',{'email':user['email'],'password':'Test-only-password-2026'});self.csrf=json.loads(body)['csrf'];_,body=self.request('/api/state');snapshot=json.loads(body);classes=copy.deepcopy(snapshot['state']['classes']);g2=next(item for item in classes if item['name']=='G2 A');g2['opening']['new']+=1
        code,body=self.request('/api/state',{'module':'enrollments','action':'edit','version':snapshot['version'],'changes':{'classes':classes}});self.assertEqual(code,200,body)
        self.client,self.csrf=adminclient,admincsrf;_,body=self.request('/api/state');shared=json.loads(body)['state'];g2=next(item for item in shared['classes'] if item['name']=='G2 A');self.assertEqual(len(shared['classes']),41);self.assertEqual(g2['opening']['new'],10);self.assertEqual(g2['students'],10);self.assertEqual(sum(item['students'] for item in shared['classes']),625);self.assertEqual(shared['enrollments']['new'],28);self.assertEqual(shared['enrollments']['re'],597);self.assertEqual(sum(max(0,item['capacity']-item['students']) for item in shared['classes']),500);self.assertTrue(any(item['user']=='Lívia' and item['field']=='G2 A · Alunos novos' and item['before']==9 and item['after']==10 for item in shared['audit']))
    def test_revenue_2027_preview_and_confirmation_are_versioned(self):
        campaign={'years':{'2027':{'year':2027,'calendar':{'enrollmentMonth':1,'tuitionMonths':[2,3,4,5,6,7,8,9,10,11,12]},'campaigns':[{'id':'jan','name':'Janeiro fixture','startDate':'2027-01-01','endDate':'2027-01-31','discountPercent':10.0,'appliesEnrollment':True,'exceptions':[],'classExceptions':[]}],'imports':[],'records':[]}}}
        _,body=self.request('/api/state');snapshot=json.loads(body);code,body=self.request('/api/state',{'module':'financial','action':'edit','version':snapshot['version'],'changes':{'revenuePlanning':campaign}});self.assertEqual(code,200,body)
        content=b'ID aluno;ID lancamento;Turma;Tipo;Data;Valor bruto;Status\nALUNO-HTTP;REC-HTTP-1;G2 A;Matricula;10/01/2027;1000;Realizada\n'
        code,body=self.request('/api/revenue-2027/preview',{'name':'receitas.csv','year':2027,'content':base64.b64encode(content).decode()});self.assertEqual(code,200,body);preview=json.loads(body);self.assertEqual(preview['ready'],1);self.assertEqual(preview['rows'][0]['net'],900)
        code,body=self.request('/api/revenue-2027/confirm',{'token':preview['token'],'decisions':[{'id':preview['rows'][0]['id'],'action':'accept','reason':''}],'confirmation':'APROVAR RECEITAS CONSOLIDADAS'});self.assertEqual(code,200,body);state=json.loads(body)['state'];records=state['revenuePlanning']['years']['2027']['records'];self.assertEqual(len(records),1);self.assertEqual(records[0]['eventType'],'enrollment');self.assertEqual(state['revenuePlanning']['years']['2027']['imports'][0]['version'],1);self.assertEqual(len(state['classes']),41)

if __name__=='__main__':unittest.main(verbosity=2)
