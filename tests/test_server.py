"""Todas as identidades e lançamentos abaixo são fixtures em bancos temporários."""
import copy, io, json, sys, tempfile, unittest, zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from server import Store,blank,validate,recalculate,upgrade_enrollment_classes,LEGACY_CORE_CLASSES
from services import build_preview,apply_preview,classify,report_rows,xlsx_report,pdf_report,_budget_class_rows,financial_classification_preview,apply_financial_classification_preview,financial_import_workbook,revenue_preview,apply_revenue_preview,revenue_year,teaching_load_preview,normalize_teaching_class_code,extract_teaching_occurrence,teaching_load_positional_pilot,teaching_load_geometric_preview,build_teaching_weekly_cost_audit,build_teaching_weekly_cost_with_shared_rateio_audit

class ServerTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.store=Store(Path(self.tmp.name)/'test.sqlite3');self.admin=self.store.create_user('Administrador fixture','admin@fixture.invalid','Test-only-password-2026',is_admin=True)
    def tearDown(self):self.tmp.cleanup()
    def user(self,name='Lívia',access=None):return self.store.create_user(name,name+'@fixture.invalid','Test-only-password-2026',access or {'enrollments':['view','create','edit'],'benefits':['view','create','edit'],'requests':['view','create','edit']},actor=self.admin)
    def patch(self,changes,user=None,module='classes',confirmation=None):return self.store.patch(user or self.admin,changes,self.store.state()[1],module,confirmation=confirmation)
    def room(self):return {'id':'c','name':'Turma fixture','capacity':10,'opening':{'new':0,'re':0},'unclassified':0,'students':0}
    def request(self,user,uid='r'):return {'id':uid,'title':'Solicitação fixture','status':'requested','creatorId':user['id'],'creatorName':user['name'],'requestedBy':user['name'],'recipientId':self.admin['id'],'recipientName':self.admin['name'],'createdAt':'2026-09-10T10:00:00','history':[],'comments':[],'due':'2026-01-01'}
    def test_report_2027_classes_and_password_hash(self):
        state,_=self.store.state();self.assertEqual(len(state['classes']),41);self.assertEqual([item['name'] for item in state['classes'][:2]],['G2 A','G2 B']);self.assertEqual(state['classes'][-1]['name'],'3º EM')
        self.assertEqual(sum(item['capacity'] for item in state['classes']),1103);self.assertEqual(state['enrollments']['new'],27);self.assertEqual(state['enrollments']['re'],597)
        self.assertEqual(sum(item['students'] for item in state['classes']),624);self.assertEqual(sum(max(0,item['capacity']-item['students']) for item in state['classes']),501)
        self.assertEqual([item['name'] for item in state['classes'] if item['students']>item['capacity']],['G4 A','G5 A','4º B','7º A'])
        with self.store.db() as db:self.assertNotIn('Test-only',db.execute('SELECT password FROM users').fetchone()[0])
    def test_only_empty_generic_classes_are_upgraded(self):
        state={'classes':[{'id':name,'name':name,'capacity':0,'students':0,'template':True,'opening':{'new':0,'re':0},'unclassified':0} for name in LEGACY_CORE_CLASSES],'enrollments':{'new':0,'re':0,'history':[],'unallocated':{'new':0,'re':0}}}
        self.assertTrue(upgrade_enrollment_classes(state));self.assertEqual(len(state['classes']),41);self.assertEqual(state['enrollments']['new'],27)
        state['classes'][0]['students']=1;state['enrollments']['new']=1
        self.assertFalse(upgrade_enrollment_classes(state));self.assertEqual(len(state['classes']),41)
    def test_login_and_revocation(self):
        user=self.user();token,csrf,_=self.store.login(user['email'],'Test-only-password-2026','local');self.assertEqual(self.store.session(token)[0]['id'],user['id']);self.store.user_action(self.admin,dict(user,active=False),'edit');self.assertIsNone(self.store.session(token)[0])
    def test_delete_user_preserves_audit(self):
        user=self.user();self.store.user_action(self.admin,{'id':user['id']},'delete');self.assertTrue(next(u for u in self.store.users() if u['id']==user['id'])['deleted']);self.assertGreater(len(self.store.snapshot(self.admin)['state']['audit']),1)
    def test_last_admin_protected(self):
        with self.assertRaises(ValueError):self.store.user_action(self.admin,{'id':self.admin['id']},'delete')
    def test_permissions_server_side(self):
        user=self.user()
        with self.assertRaises(PermissionError):self.patch({'classes':[self.room()]},user)
        with self.assertRaises(PermissionError):self.store.user_action(user,{'name':'X','email':'x@fixture.invalid'},'create')
    def test_only_admin_changes_capacity_and_quantities(self):
        self.patch({'classes':[self.room()]});user=self.user();classes=copy.deepcopy(self.store.state()[0]['classes']);classes[0]['capacity']=12;classes[0]['opening']={'new':2,'re':3}
        with self.assertRaises(PermissionError):self.patch({'classes':classes},user,'enrollments')
        self.patch({'classes':classes},self.admin,'enrollments');state,_=self.store.state();self.assertEqual(state['classes'][0]['capacity'],12);self.assertEqual(state['classes'][0]['students'],5)
        fields=[item['field'] for item in self.store.snapshot(self.admin)['state']['audit']];self.assertIn('c.capacity',fields);self.assertIn('c.opening',fields)
    def test_livia_reconciles_quantities_without_turma_administration_and_is_audited(self):
        livia=self.user('Lívia',{'enrollments':['view','create','edit'],'classes':['view'],'benefits':['view','create','edit'],'requests':['view','create','edit']});before,_=self.store.state();classes=copy.deepcopy(before['classes']);g2=next(item for item in classes if item['name']=='G2 A');g2['opening']['new']+=1;g2['unclassified']+=2;recalculate({'classes':classes,'enrollments':copy.deepcopy(before['enrollments'])})
        self.patch({'classes':classes},livia,'enrollments');after,_=self.store.state();changed=next(item for item in after['classes'] if item['name']=='G2 A');self.assertFalse(livia['isAdmin']);self.assertEqual(livia['access']['classes'],['view']);self.assertEqual(changed['capacity'],18);self.assertEqual(changed['opening']['new'],10);self.assertEqual(changed['unclassified'],2);self.assertEqual(changed['students'],12);self.assertEqual(len(after['classes']),41)
        audit=self.store.snapshot(self.admin)['state']['audit'];details=[item for item in audit if item['userId']==livia['id']];self.assertTrue(any(item['field']=='G2 A · Alunos novos' and item['before']==9 and item['after']==10 for item in details));self.assertTrue(any(item['field']=='G2 A · Matriculados ainda não classificados' and item['before']==0 and item['after']==2 for item in details))
        forbidden=copy.deepcopy(after['classes']);next(item for item in forbidden if item['name']=='G2 A')['capacity']=19
        with self.assertRaises(PermissionError):self.patch({'classes':forbidden},livia,'enrollments')
        with self.assertRaises(PermissionError):self.patch({'classes':after['classes']+[self.room()]},livia,'classes')
    def test_financial_years_keep_2027_data_and_audit_classification(self):
        before,_=self.store.state();self.assertEqual(len(before['classes']),41);year=copy.deepcopy(before['academicYears']);current=next(x for x in year if x['year']==2027);self.assertEqual({x['id']:x['tuition'] for x in current['parameters']},{'early':930.69,'fundamental1':964.67,'fundamental2':1239.81,'secondary12':1425.59,'secondary3':1461.05});g2=next(x for x in before['classes'] if x['name']=='G2 A');current['classifications'][g2['id']]={'fixed':{'noDiscount':1,'philanthropic100':1,'philanthropic50':1,'staffChild100':0,'staffChild80':0,'workerChild100':0,'markingLives45':1},'variables':[{'id':'v16','percent':16,'quantity':1}]}
        self.patch({'academicYears':year},module='financial');after,_=self.store.state();self.assertEqual(len(after['classes']),41);self.assertEqual(next(x for x in after['classes'] if x['id']==g2['id'])['students'],9);audit=self.store.snapshot(self.admin)['state']['audit'];self.assertTrue(any(x['field']=='Ano 2027 · G2 A · Bolsa filantrópica · 100%' and x['before']==0 and x['after']==1 for x in audit));self.assertTrue(any(x['field']=='Ano 2027 · G2 A · Outros descontos variáveis · 16%' for x in audit))
        years=copy.deepcopy(after['academicYears']);years.append({'id':'2028','year':2028,'parameters':[{'id':key,'label':label,'tuition':None} for key,label in [('early','Educação Infantil'),('fundamental1','Fundamental I'),('fundamental2','Fundamental II'),('secondary12','1º e 2º ano do Ensino Médio'),('secondary3','3º ano do Ensino Médio')]],'classifications':{}});self.patch({'academicYears':years,'activeAcademicYear':2028},module='financial');final,_=self.store.state();self.assertEqual(len(final['classes']),41);self.assertIn(2027,[x['year'] for x in final['academicYears']]);self.assertIn(2028,[x['year'] for x in final['academicYears']])
        with self.assertRaises(PermissionError):self.patch({'academicYears':final['academicYears']},self.user(),module='financial')
    def test_financial_report_preserves_unclassified_and_uses_active_year(self):
        state,_=self.store.state();g2=next(x for x in state['classes'] if x['name']=='G2 A');year=next(x for x in state['academicYears'] if x['year']==2027)
        year['classifications'][g2['id']]={'fixed':{'noDiscount':1,'philanthropic100':1,'philanthropic50':0,'staffChild100':0,'staffChild80':0,'workerChild100':0,'markingLives45':0},'variables':[]}
        self.patch({'academicYears':state['academicYears']},module='financial');columns,rows=report_rows(self.store.state()[0],'financial');row=next(x for x in rows if x[1]=='G2 A')
        self.assertEqual(columns[0],'Ano');self.assertEqual(row[0],2027);self.assertEqual(row[3],9);self.assertEqual(row[4],930.69);self.assertEqual(row[13],7);self.assertEqual(row[16],930.69)
    def test_variable_discount_saves_when_classification_is_below_enrollment_total(self):
        before,_=self.store.state();g2=next(x for x in before['classes'] if x['name']=='G2 A');enrollments=copy.deepcopy(before['enrollments']);enrollments['history'].append({'id':'qa-g2-plus-one','classId':g2['id'],'className':'G2 A','type':'new','quantity':1,'date':'2026-09-11','user':'Fixture'})
        self.patch({'enrollments':enrollments},module='enrollments');state,_=self.store.state();years=copy.deepcopy(state['academicYears']);year=next(x for x in years if x['year']==2027);year['classifications'][g2['id']]={'fixed':{'noDiscount':3,'philanthropic100':1,'philanthropic50':1,'staffChild100':1,'staffChild80':1,'workerChild100':0,'markingLives45':1},'variables':[{'id':'variable-16','percent':16,'quantity':1}]}
        self.patch({'academicYears':years},module='financial');saved,_=self.store.state();room=next(x for x in saved['classes'] if x['id']==g2['id']);self.assertEqual(len(saved['classes']),41);self.assertEqual(room['students'],10);classification=next(x for x in saved['academicYears'] if x['year']==2027)['classifications'][g2['id']];self.assertEqual(sum(classification['fixed'].values())+classification['variables'][0]['quantity'],9);self.assertEqual(classification['variables'][0]['percent'],16)
        _,rows=report_rows(saved,'financial');row=next(x for x in rows if x[1]=='G2 A');self.assertEqual(row[13],1);self.assertEqual(round(row[14],2),8376.21);self.assertEqual(round(row[15],2),3639.00);self.assertEqual(round(row[16],2),4737.21);self.assertEqual(round(row[17],2),526.36)
    def test_break_even_layer_is_separate_audited_and_requires_financial_permission(self):
        state,_=self.store.state();self.assertEqual(state['breakEven']['plans'],[]);plan={'id':'official-2027','year':2027,'source':{'sourceName':'Orçamento oficial 2027'},'officialTotals':{'totalExpenses':10097843.57,'metaStudents':1065},'costLines':[{'id':'pending-total','label':'Total despesas oficial','amount':10097843.57,'classification':'pending','reconciliation':'pending'}],'mappings':[{'id':'m-g2','operationalClassId':state['classes'][0]['id'],'status':'pending'}],'rateioRules':[]}
        self.patch({'breakEven':{'plans':[plan]}},module='financial');saved,_=self.store.state();self.assertEqual(len(saved['classes']),41);self.assertEqual(saved['classes'][0]['students'],9);self.assertEqual(saved['breakEven']['plans'][0]['officialTotals']['totalExpenses'],10097843.57);self.assertTrue(any(item['module']=='financial' and item['field']=='breakEven' for item in self.store.snapshot(self.admin)['state']['audit']))
        user=self.user('Financeiro sem acesso',{'financial':['view']})
        with self.assertRaises(PermissionError):self.patch({'breakEven':{'plans':[]}},user,'financial')
    def test_official_2027_personnel_composition_is_auditable_and_not_a_new_cost(self):
        state,_=self.store.state();plan={'id':'official-personnel-2027','year':2027,'source':{'sourceName':'Orçamento COLEGIO ADVENTISTA DE JUAZEIRO 2027 (1).pdf'},'officialTotals':{'totalExpenses':10097843.57,'totalExpensesMonthly':841486.96},'costLines':[],'mappings':[],'rateioRules':[]}
        self.patch({'breakEven':{'plans':[plan]}},module='financial');saved,_=self.store.state();stored=saved['breakEven']['plans'][0];audit=stored['personnelCostAudit']
        self.assertEqual(len(saved['classes']),41);self.assertEqual(stored['officialTotals']['totalExpensesMonthly'],841486.96);self.assertEqual(stored['costLines'],[]);self.assertTrue(audit['includedInOfficialTotalExpenses']);self.assertEqual(audit['allocationStatus'],'not_allocated_to_teachers_or_classes');self.assertAlmostEqual(audit['totalMonthly'],387591.93);self.assertAlmostEqual(audit['payrollMonthly']+audit['personalChargesAndBenefitsMonthly']+audit['nonEmployeeServicesMonthly'],audit['totalMonthly'])
    def test_budget_versions_preserve_operational_classes_and_audit(self):
        state,_=self.store.state();plans=[{'id':'budget-2027-v1','year':2027,'version':1,'status':'approved','source':{'sourceName':'oficial.pdf'},'officialTotals':{'metaStudents':1065,'totalExpenses':10097843.57},'plannedClasses':[],'tuitionParameters':[],'delinquency':{'officialPercent':4.5,'scenarioPercent':None},'installments':11,'enrollmentRule':{'status':'pending_validation'},'costCenters':[],'costLines':[],'rateioRules':[],'mappings':[],'extractions':[{'id':'x1','original':1065,'used':1065,'action':'accept','origin':'página 2'}],'scenarios':[]},{'id':'budget-2028-v1','year':2028,'version':1,'status':'analysis','source':{'sourceName':'revisao.xlsx'},'officialTotals':{},'plannedClasses':[],'tuitionParameters':[],'delinquency':{'officialPercent':None,'scenarioPercent':None},'installments':11,'enrollmentRule':{'status':'pending_validation'},'costCenters':[],'costLines':[],'rateioRules':[],'mappings':[],'extractions':[],'scenarios':[]}]
        self.patch({'breakEven':{'plans':plans}},module='financial');saved,_=self.store.state();self.assertEqual(len(saved['classes']),41);self.assertEqual(len(saved['breakEven']['plans']),2);self.assertEqual(saved['breakEven']['plans'][0]['officialTotals']['metaStudents'],1065);self.assertTrue(any(x['field']=='breakEven' for x in self.store.snapshot(self.admin)['state']['audit']))
    def test_official_budget_rows_are_separate_from_operational_classes(self):
        line='Grupo 2 A 16 14.890,97 341,11 15.232,08 - 446,73 14.785,35 6.271,93 1.520,00 5.751,16 - 13.543,09 1.242,25 15'
        pending='3º E. Médio B 0 - - - 9.496,80 (21,92) (9.474,89) - - - - - (9.474,89) 0'
        rows=_budget_class_rows([line+'\n'+pending]);self.assertEqual(len(rows),2)
        g2=rows[0];self.assertEqual(g2['students'],16);self.assertEqual(g2['breakEvenStudents'],15);self.assertEqual(g2['netRevenue'],14785.35);self.assertEqual(g2['calculatedResult'],1242.26);self.assertEqual(g2['status'],'recognized')
        self.assertEqual(rows[1]['status'],'pending_review')
        before,_=self.store.state();plan={'id':'source-2027','year':2027,'officialTotals':{},'officialBudget':{'status':'source_imported','classRows':rows,'totals':{'students':1065}},'managerialPlanning':{'status':'draft','installments':11,'scenarios':[]}}
        self.patch({'breakEven':{'plans':[plan]}},module='financial');after,_=self.store.state();self.assertEqual(len(after['classes']),41);self.assertEqual(after['classes'],before['classes']);self.assertEqual(len(after['breakEven']['plans'][0]['officialBudget']['classRows']),2)
    def test_imported_official_source_is_immutable_but_financial_classifications_remain_editable(self):
        state,_=self.store.state();source={'status':'source_imported','classRows':[{'id':'official-row','className':'Grupo 2 A','students':16}],'totals':{'students':1065}}
        plan={'id':'official-source','year':2027,'officialTotals':{},'officialBudget':source,'managerialPlanning':{'status':'draft','installments':11,'scenarios':[]}}
        self.patch({'breakEven':{'plans':[plan]}},module='financial');changed=copy.deepcopy(self.store.state()[0]['breakEven']);changed['plans'][0]['officialBudget']['totals']['students']=999
        with self.assertRaises(ValueError):self.patch({'breakEven':changed},module='financial')
        years=copy.deepcopy(self.store.state()[0]['academicYears']);g2=next(x for x in self.store.state()[0]['classes'] if x['name']=='G2 A');years[0]['classifications'][g2['id']]={'fixed':{'noDiscount':8,'philanthropic100':1,'philanthropic50':0,'staffChild100':0,'staffChild80':0,'workerChild100':0,'markingLives45':0},'variables':[]}
        self.patch({'academicYears':years},module='financial');saved,_=self.store.state();self.assertEqual(saved['breakEven']['plans'][0]['officialBudget'],source);self.assertEqual(len(saved['classes']),41)
    def test_direct_cost_conciliation_requires_evidence_and_administrator(self):
        state,_=self.store.state();g2=next(item for item in state['classes'] if item['name']=='G2 A');source={'status':'source_imported','classRows':[{'id':'official-g2','className':'Grupo 2 A','status':'recognized','totalExpenses':13543.09,'origin':'página fixture'}],'totals':{'totalExpensesMonthly':13543.09}}
        mapping={'id':'map-g2','operationalClassId':g2['id'],'operationalClassName':g2['name'],'budgetClassName':'Grupo 2 A','sourceRowId':'official-g2','costMonthly':13543.09,'status':'mapped','costLinkType':'direct_source','costEvidence':{'status':'verified','method':'equivalência única','origin':'página fixture'}}
        plan={'id':'cost-source','year':2027,'officialTotals':{},'officialBudget':source,'managerialPlanning':{'status':'draft','installments':11,'scenarios':[]},'mappings':[mapping]}
        self.patch({'breakEven':{'plans':[plan]}},module='financial');saved,_=self.store.state();self.assertEqual(len(saved['classes']),41);self.assertEqual(saved['breakEven']['plans'][0]['mappings'][0]['costMonthly'],13543.09);self.assertEqual(saved['breakEven']['plans'][0].get('rateioRules',[]),[])
        finance_user=self.user('Financeiro conciliação',{'financial':['view','edit']});changed=copy.deepcopy(saved['breakEven']);changed['plans'][0]['mappings'][0]['costMonthly']=1
        with self.assertRaises(PermissionError):self.patch({'breakEven':changed},finance_user,'financial')
    def test_financial_discount_import_preview_apply_and_recalculation(self):
        state,_=self.store.state();content='Turma;Benefício;Percentual;Quantidade\nG2 A;Sem desconto;0;3\nG2 A;Bolsa filantrópica;100;1\nG2 A;Desconto variável;16;1\n'.encode()
        before=copy.deepcopy(state);preview=financial_classification_preview('descontos.csv',content,state,2027)
        self.assertEqual(state,before);self.assertEqual(preview['automatic'],3);self.assertEqual(preview['pending'],0)
        decisions=[{'id':item['id'],'action':'accept','classId':item['classId'],'category':item['category'],'percent':item['percent'],'quantity':item['quantity'],'reason':''} for item in preview['rows']]
        updated=apply_financial_classification_preview(state,preview,decisions,self.admin);g2=next(x for x in updated['classes'] if x['name']=='G2 A');classification=next(x for x in updated['academicYears'] if x['year']==2027)['classifications'][g2['id']]
        self.assertEqual(classification['fixed']['noDiscount'],3);self.assertEqual(classification['fixed']['philanthropic100'],1);self.assertEqual(classification['variables'][0]['percent'],16);self.assertEqual(classification['variables'][0]['quantity'],1);self.assertEqual(updated['imports'][-1]['kind'],'financial_classifications')
        self.patch({'academicYears':updated['academicYears'],'imports':updated['imports']},module='financial');saved,_=self.store.state();_,rows=report_rows(saved,'financial');row=next(x for x in rows if x[1]=='G2 A');self.assertEqual(row[3],9);self.assertEqual(row[13],4);self.assertEqual(round(row[14],2),4653.45);self.assertEqual(round(row[15],2),1079.60);self.assertEqual(round(row[16],2),3573.85);self.assertEqual(round(row[17],2),714.77);self.assertEqual(len(saved['classes']),41)
    def test_financial_discount_import_flags_unknown_student_percent_and_duplicates(self):
        state,_=self.store.state();content='Turma;Benefício;Percentual;Quantidade;Aluno\nG2 A;Desconto variável;16;1;\nG2 A;Desconto variável;16;1;\nTurma inexistente;Sem desconto;0;1;\nG2 A;Sem desconto;120;1;\nG2 A;Sem desconto;0;1;Aluno fixture\n'.encode()
        preview=financial_classification_preview('pendencias.csv',content,state,2027);self.assertEqual(preview['duplicates'],1);self.assertEqual(preview['pending'],1);self.assertEqual(len(preview['invalid']),1)
        individual=next(item for item in preview['rows'] if item.get('student'))
        self.assertEqual(individual['status'],'ready');self.assertEqual(individual['quantity'],1)
        applied=apply_financial_classification_preview(state,preview,[{'id':individual['id'],'action':'accept'}],self.admin);g2=next(x for x in applied['classes'] if x['name']=='G2 A');self.assertEqual(next(x for x in applied['academicYears'] if x['year']==2027)['classifications'][g2['id']]['fixed']['noDiscount'],1)
    def test_financial_discount_preview_has_all_classes_and_blocks_class_overflow(self):
        state,_=self.store.state();content='Aluno;Turma;Benefício;Percentual\nAluno A;G2 A;Bolsa filantrópica;100\nAluno B;G2 A;Desconto;16\nAluno sem turma;;Sem desconto;0\n'.encode()
        preview=financial_classification_preview('nominal.csv',content,state,2027);self.assertEqual(len(preview['classSummary']),41);self.assertEqual(preview['unrecognized'],1);self.assertTrue(preview['canConfirm'])
        g2=next(x for x in preview['classSummary'] if x['className']=='G2 A');self.assertEqual(g2['fixed']['philanthropic100'],1);self.assertEqual(g2['variables'],[{'percent':16.0,'quantity':1}]);self.assertEqual(g2['pending'],0)
        overflow='Turma;Benefício;Percentual;Quantidade\nG2 A;Sem desconto;0;999\n'.encode();blocked=financial_classification_preview('excesso.csv',overflow,state,2027);self.assertFalse(blocked['canConfirm']);self.assertEqual(next(x for x in blocked['classSummary'] if x['className']=='G2 A')['status'],'blocked')
        decision={'id':blocked['rows'][0]['id'],'action':'accept'}
        with self.assertRaisesRegex(ValueError,'excede os matriculados'):apply_financial_classification_preview(state,blocked,[decision],self.admin)
    def test_financial_discount_template_and_isolated_homologation_fixture(self):
        from openpyxl import load_workbook
        model=financial_import_workbook();book=load_workbook(io.BytesIO(model),data_only=True);self.assertEqual(book.sheetnames,['Nominal','Consolidada por turma','Instruções']);self.assertEqual([cell.value for cell in book['Nominal'][1]],['Aluno','Turma','Série/Ano','Categoria/Benefício','Percentual','Observação']);self.assertIn('Projeto Marcando Vidas — 45%',[row[0].value for row in book['Instruções']]);book.close()
        state,_=self.store.state();before=copy.deepcopy(state);fixture=financial_import_workbook(homologation=True);preview=financial_classification_preview('HOMOLOGACAO-FICTICIA-NAO-CONFIRMAR.xlsx',fixture,state,2027)
        self.assertEqual(state,before);self.assertEqual(len(preview['classSummary']),41);self.assertFalse(preview['canConfirm']);self.assertTrue(any(row['classificationKey']=='variable' and row['percent']==16 for row in preview['rows']));self.assertTrue(any(row['className']=='TURMA FICTÍCIA INEXISTENTE' and row['status']=='pending_review' for row in preview['rows']));self.assertTrue(any(row['classId'] is None and row['source'].get('className','')=='' and row['status']=='pending_review' for row in preview['rows']));self.assertTrue(any(row['category']=='Categoria ambígua' and row['status']=='pending_review' for row in preview['rows']))
        g2=next(row for row in preview['classSummary'] if row['className']=='G2 A');self.assertEqual(g2['status'],'blocked');self.assertGreater(g2['classified'],g2['enrolled']);self.assertEqual(state,before)
    def test_financial_discount_xlsx_ignores_instructions_and_never_uses_delinquency_error(self):
        from openpyxl import Workbook
        def workbook(sheets):
            book=Workbook();book.remove(book.active)
            for title,rows in sheets:
                sheet=book.create_sheet(title)
                for row in rows:sheet.append(row)
            buffer=io.BytesIO();book.save(buffer);book.close();return buffer.getvalue()
        state,_=self.store.state();valid=[['Turma','Categoria/Benefício','Percentual','Quantidade'],['G2 A','Sem desconto',0,1]]
        content=workbook([('Dados',valid),('Instruções',[['Texto auxiliar sem dados']])]);preview=financial_classification_preview('com-instrucoes.xlsx',content,state,2027);self.assertEqual(preview['sheet'],'Dados');self.assertEqual(preview['automatic'],1)
        ambiguous=workbook([('Janeiro',valid),('Fevereiro',valid),('Instruções',[['Ignore esta aba']])]);selection=financial_classification_preview('ambiguo.xlsx',ambiguous,state,2027);self.assertTrue(selection['needsSheetSelection']);self.assertEqual(selection['sheetCandidates'],['Janeiro','Fevereiro'])
        selected=financial_classification_preview('ambiguo.xlsx',ambiguous,state,2027,'Fevereiro');self.assertEqual(selected['sheet'],'Fevereiro');self.assertEqual(selected['automatic'],1)
        invalid=workbook([('Instruções',[['Leia antes']]),('Apoio',[['Sem colunas financeiras']])])
        with self.assertRaises(ValueError) as context:financial_classification_preview('invalido.xlsx',invalid,state,2027)
        self.assertNotIn('inadimplência',str(context.exception).lower());self.assertIn('descontos e benefícios',str(context.exception).lower())
    def test_financial_discount_duplicate_invalid_cancel_and_recalculation_chain(self):
        from financial_integration import integration_snapshot
        state,_=self.store.state();year=next(row for row in state['academicYears'] if row['year']==2027)
        for room in state['classes']:year['classifications'][room['id']]={'fixed':{'noDiscount':room['students'],'philanthropic100':0,'philanthropic50':0,'staffChild100':0,'staffChild80':0,'workerChild100':0,'markingLives45':0},'variables':[]}
        source=[];mappings=[]
        for room in state['classes']:
            source.append({'id':'source-'+room['id'],'className':room['name'],'status':'recognized'});mappings.append({'operationalClassId':room['id'],'budgetClassName':room['name'],'sourceRowId':'source-'+room['id'],'status':'mapped','costMonthly':5000,'costEvidence':{'status':'verified'}})
        state['breakEven']={'plans':[{'id':'fixture-pe','year':2027,'officialBudget':{'classRows':source,'totals':{'totalExpensesMonthly':205000}},'mappings':mappings,'delinquency':{'scenarioPercent':0}}]};ledger={'target_year':2027,'conflicted_cost_cents':0,'unassigned_cost_cents':0,'validated_cost_cents':4100,'classes':[{'class_id':room['id'],'weekly_cost':1} for room in state['classes']]}
        before=integration_snapshot(state,ledger,2027);g2=next(room for room in state['classes'] if room['name']=='G2 A');content=f'Turma;Categoria/Benefício;Percentual;Quantidade\nG2 A;Sem desconto;0;{g2["students"]-1}\nG2 A;Desconto variável;16;1\n'.encode();preview=financial_classification_preview('recalculo.csv',content,state,2027);self.assertEqual(len(preview['classSummary']),41)
        self.assertEqual(state['imports'],[])
        with self.assertRaises(ValueError):financial_classification_preview('invalido.csv',b'cabecalho desconhecido\nvalor',state,2027)
        decisions=[{'id':row['id'],'action':'accept'} for row in preview['rows']];updated=apply_financial_classification_preview(state,preview,decisions,self.admin);after=integration_snapshot(updated,ledger,2027);before_g2=next(row for row in before['classes'] if row['name']=='G2 A');after_g2=next(row for row in after['classes'] if row['name']=='G2 A')
        self.assertLess(after_g2['netRevenueMonthly'],before_g2['netRevenueMonthly']);self.assertLess(after_g2['netTicketMonthly'],before_g2['netTicketMonthly']);self.assertGreaterEqual(after_g2['breakEvenStudents'],before_g2['breakEvenStudents']);self.assertLess(after['recurringNetRevenueMonthly'],before['recurringNetRevenueMonthly']);self.assertGreaterEqual(after['breakEvenStudents'],before['breakEvenStudents'])
        with self.assertRaisesRegex(ValueError,'já foi confirmado'):financial_classification_preview('mesmo-arquivo.csv',content,updated,2027)
    def test_financial_discount_import_reads_xlsx_pdf_and_requires_administrator(self):
        state,_=self.store.state();xlsx=xlsx_report(['Turma','Benefício','Percentual','Quantidade'],[['G2 A','Sem desconto',0,2]])
        self.assertEqual(financial_classification_preview('descontos.xlsx',xlsx,state,2027)['automatic'],1)
        from reportlab.pdfgen import canvas
        buffer=io.BytesIO();pdf=canvas.Canvas(buffer);pdf.drawString(30,800,'Turma;Benefício;Percentual;Quantidade');pdf.drawString(30,780,'G2 A;Sem desconto;0;2');pdf.save()
        preview=financial_classification_preview('descontos.pdf',buffer.getvalue(),state,2027);self.assertEqual(preview['automatic'],1)
        decision={'id':preview['rows'][0]['id'],'action':'accept','classId':preview['rows'][0]['classId'],'category':'Sem desconto','percent':0,'quantity':2,'reason':''}
        finance_user=self.user('Financeiro importador',{'financial':['view','edit']})
        with self.assertRaises(PermissionError):apply_financial_classification_preview(state,preview,[decision],finance_user)
    def test_revenue_2027_keeps_enrollment_separate_and_versions_imports(self):
        state,_=self.store.state();g2=next(x for x in state['classes'] if x['name']=='G2 A');config=revenue_year(state,2027);config['campaigns']=[{'id':'jan','name':'Janeiro','startDate':'2027-01-01','endDate':'2027-01-31','discountPercent':10.0,'appliesEnrollment':True,'exceptions':[],'classExceptions':[]}]
        content=('ID aluno;ID lançamento;Turma;Tipo;Data;Valor bruto;Desconto;Status;Baixa financeira\n'
                 'ALUNO-001;FIN-001;G2 A;Matrícula;10/01/2027;1000;0;Realizada;900\n'
                 'ALUNO-001;FIN-002;G2 A;Mensalidade;10/02/2027;930,69;0;A receber;\n').encode()
        preview=revenue_preview('receitas.csv',content,state,2027);self.assertEqual(preview['ready'],2);self.assertEqual(preview['pending'],0);self.assertEqual(round(preview['rows'][0]['discount'],2),100);self.assertEqual(round(preview['rows'][0]['net'],2),900);self.assertEqual(preview['rows'][0]['financialDivergence'],0)
        decisions=[{'id':row['id'],'action':'accept','reason':''} for row in preview['rows']];updated=apply_revenue_preview(state,preview,decisions,self.admin);records=revenue_year(updated,2027)['records'];self.assertEqual(len(records),2);self.assertEqual(records[0]['eventType'],'enrollment');self.assertEqual(records[1]['eventType'],'tuition');self.assertEqual(revenue_year(updated,2027)['imports'][0]['version'],1);self.assertEqual(len(updated['classes']),41);validate(updated,state)
    def test_revenue_2027_rejects_name_only_duplicate_and_invalid_month(self):
        state,_=self.store.state();content=('Nome;ID lançamento;Turma;Tipo;Data;Valor bruto;Status\n'
                 'Aluno sem id;X-1;G2 A;Matrícula;10/02/2027;100;Realizada\n'
                 'ALUNO-2;X-2;G2 A;Mensalidade;10/01/2027;100;Realizada\n').encode()
        preview=revenue_preview('invalidas.csv',content,state,2027);self.assertEqual(len(preview['invalid']),2)
        duplicate=('ID aluno;ID lançamento;Turma;Tipo;Data;Valor bruto;Status\nALUNO-2;DUP;G2 A;Matrícula;10/01/2027;100;Realizada\nALUNO-2;DUP-2;G2 A;Matrícula;10/01/2027;100;Realizada\n').encode()
        self.assertEqual(revenue_preview('duplicadas.csv',duplicate,state,2027)['duplicates'],1)
    def test_revenue_campaign_exception_preserves_individual_amount(self):
        state,_=self.store.state();config=revenue_year(state,2027);config['campaigns']=[{'id':'jan','name':'Campanha fixture','startDate':'2027-01-01','endDate':'2027-01-31','discountPercent':25.0,'appliesEnrollment':True,'exceptions':['ALUNO-ISENTO'],'classExceptions':[]}]
        content=('ID aluno;ID lançamento;Turma;Tipo;Data;Valor bruto;Status\nALUNO-COM-DESC;A;G2 A;Matrícula;05/01/2027;1000;Realizada\nALUNO-ISENTO;B;G2 A;Matrícula;05/01/2027;1000;Realizada\n').encode()
        preview=revenue_preview('campanha.csv',content,state,2027);self.assertEqual([row['discount'] for row in preview['rows']],[250.0,0.0]);self.assertEqual([row['net'] for row in preview['rows']],[750.0,1000.0])
    def test_revenue_transfers_do_not_add_and_refunds_preserve_history(self):
        state,_=self.store.state();content=('ID aluno;ID lançamento;Turma;Tipo;Data;Valor bruto;Status;Estorno de\n'
                 'ALUNO-3;ORIGINAL;G2 A;Matrícula;10/01/2027;100;Realizada;\n'
                 'ALUNO-3;MOVE-1;G2 A;Transferência;11/02/2027;100;Realizada;\n'
                 'ALUNO-3;REF-1;G2 A;Estorno;12/02/2027;100;Realizada;ORIGINAL\n').encode()
        preview=revenue_preview('eventos.csv',content,state,2027);self.assertEqual(preview['ready'],3);updated=apply_revenue_preview(state,preview,[{'id':x['id'],'action':'accept','reason':''} for x in preview['rows']],self.admin);records=revenue_year(updated,2027)['records'];self.assertEqual(len(records),3);self.assertEqual(next(x for x in records if x['externalId']=='MOVE-1')['direction'],0);self.assertEqual(sum(x['net']*x['direction'] for x in records),0)
    def test_teaching_load_preview_is_readonly_and_maps_exact_codes(self):
        from reportlab.pdfgen import canvas
        b=io.BytesIO();c=canvas.Canvas(b);c.drawString(20,800,'Professor: Docente Fixture (99) Período letivo: 2026-AL');c.drawString(20,780,'EFUND03MB - MATEM | EFUND03MB - CIENC');c.drawString(20,760,'07:20 - 45');c.save();state,_=self.store.state();before=copy.deepcopy(state);p=teaching_load_preview('carga_horaria.pdf',b.getvalue(),state)
        self.assertEqual(state,before);self.assertEqual(p['status'],'PREVIEW');self.assertEqual(p['source_year'],2026);self.assertEqual(p['target_year'],2027);self.assertEqual(len(state['classes']),41);self.assertEqual(p['records'][0]['professor_id'],'99');self.assertEqual(p['records'][0]['duration_minutes'],45);self.assertTrue(p['records'][0]['shared_class']);self.assertEqual(p['records'][0]['reconciliation_status'],'CONCILIADO')
        state['teachingLoad']={'versions':[{'source_hash':p['source_hash']}]}
        with self.assertRaises(ValueError):teaching_load_preview('carga_horaria.pdf',b.getvalue(),state)
    def test_teaching_code_normalization_is_conservative(self):
        rooms=blank()['classes'];self.assertEqual(normalize_teaching_class_code('EFUND03MB',rooms)['status'],'CONCILIADO')
        for code in ['EINFA02MA','EMERE01MA','EMICN02MA','EMICH03MA','DESCONHECIDO']:
            self.assertEqual(normalize_teaching_class_code(code,rooms)['status'],'PENDENTE')
    def test_teaching_occurrences_preserve_real_duration_and_shared_interval(self):
        for minutes,end in [(40,'11:25'),(45,'08:05'),(50,'10:45')]:
            start={'11:25': '10:45','08:05':'07:20','10:45':'09:55'}[end];row=extract_teaching_occurrence('Docente','7','Segunda',f'EFUND03MB - MATEM',f'{start} - {minutes}');self.assertEqual(row['duration_minutes'],minutes);self.assertEqual(row['end_time'],end);self.assertEqual(row['day_of_week'],'Segunda');self.assertEqual(row['discipline'],'MATEM')
        shared=extract_teaching_occurrence('Docente','7','Terça','EFUND03MB - MATEM | EFUND03MB - CIENC','07:20 - 45');self.assertTrue(shared['shared_class']);self.assertEqual(len(shared['source_class_codes']),2);self.assertEqual(shared['duration_minutes'],45)
        pending=extract_teaching_occurrence(None,None,None,'texto sem código','');self.assertEqual(pending['status'],'PENDENTE');self.assertIsNone(pending['duration_minutes'])
    def test_positional_teaching_pilot_uses_grid_not_legend_or_empty_cells(self):
        from reportlab.pdfgen import canvas
        b=io.BytesIO();c=canvas.Canvas(b,pagesize=(595,842));c.drawString(18,760,'Professor: Docente Piloto (77)');c.drawString(18,744,'Turno: Manhã')
        days=['Segunda','Terça','Quarta','Quinta','Sexta'];lefts=[50,151,252,353,454]
        for day,left in zip(days,lefts):c.drawString(left+30,726,day);c.rect(left,630,100,95);c.line(left,642,left+100,642);c.setFont('Helvetica',8);c.drawString(left+8,690,'EFUND03MB - MATEM' if day=='Segunda' else ('EFUND03MB - CIENC' if day=='Terça' else '--'));c.drawString(left+30,634,'--' if day=='Terça' else '07:20 - 45');c.setFont('Helvetica',12)
        c.drawString(18,610,'Legenda: EFUND03MB - MATEM não é aula');c.save()
        pilot=teaching_load_positional_pilot(b.getvalue(),blank()['classes']);self.assertEqual(len(pilot['occurrences']),2);row=pilot['occurrences'][0];self.assertEqual(row['day_of_week'],'Segunda');self.assertEqual(row['start_time'],'07:20');self.assertEqual(row['end_time'],'08:05');self.assertEqual(row['duration_minutes'],45);self.assertEqual(row['status_temporal'],'CONFIRMADO');self.assertEqual(row['status_normalization'],'CONCILIADO');pending=pilot['occurrences'][1];self.assertEqual(pending['day_of_week'],'Terça');self.assertIsNone(pending['start_time']);self.assertEqual(pending['status_temporal'],'PENDENTE');self.assertTrue(pilot['legend_ignored']);self.assertEqual(pilot['columns'][0]['x0'],50.0)
    def test_positional_teaching_pilot_keeps_documentally_absent_days_absent(self):
        from reportlab.pdfgen import canvas
        b=io.BytesIO();c=canvas.Canvas(b,pagesize=(595,842));days=['Terça','Quinta','Sexta'];lefts=[50,151,252]
        for day,left in zip(days,lefts):c.drawString(left+30,726,day);c.rect(left,630,100,95);c.line(left,642,left+100,642);c.setFont('Helvetica',8);c.drawString(left+8,690,'EFUND03MB - MATEM' if day=='Quinta' else '--');c.drawString(left+30,634,'07:20 - 45');c.setFont('Helvetica',12)
        c.save();pilot=teaching_load_positional_pilot(b.getvalue(),blank()['classes']);self.assertEqual([x['day'] for x in pilot['columns']],days);self.assertEqual(len(pilot['occurrences']),1);self.assertEqual(pilot['occurrences'][0]['day_of_week'],'Quinta')
    def test_positional_teaching_pilot_preserves_shared_cell_as_one_interval(self):
        from reportlab.pdfgen import canvas
        b=io.BytesIO();c=canvas.Canvas(b,pagesize=(595,842));c.drawString(80,726,'Segunda');c.rect(50,630,100,95);c.line(50,642,150,642);c.setFont('Helvetica',8);c.drawString(58,690,'EFUND03MB - MATEM |');c.drawString(58,680,'EFUND03MB - CIENC');c.drawString(80,634,'07:20 - 45');c.save()
        pilot=teaching_load_positional_pilot(b.getvalue(),blank()['classes']);self.assertEqual(len(pilot['occurrences']),1);row=pilot['occurrences'][0];self.assertTrue(row['shared_class']);self.assertEqual(row['source_class_codes'],['EFUND03MB','EFUND03MB']);self.assertEqual(row['duration_minutes'],45);self.assertEqual(row['status_normalization'],'PENDENTE')
    def test_geometric_preview_is_readonly_and_keeps_legend_only_page_out_of_records(self):
        from reportlab.pdfgen import canvas
        b=io.BytesIO();c=canvas.Canvas(b,pagesize=(595,842));c.drawString(18,760,'Professor: Docente Completo (88)');c.drawString(80,726,'Segunda');c.rect(50,630,100,95);c.line(50,642,150,642);c.setFont('Helvetica',8);c.drawString(58,690,'EFUND03MB - MATEM');c.drawString(80,634,'07:20 - 45');c.showPage();c.drawString(18,760,'Professor: Docente Completo (88)');c.drawString(18,730,'Legenda: (EFUND03MB - MATEM) material documental');c.save()
        before=copy.deepcopy(blank());preview=teaching_load_geometric_preview(b.getvalue(),before['classes']);self.assertEqual(before,blank());self.assertEqual(preview['pages_processed'],2);self.assertEqual(len(preview['records']),1);self.assertEqual(preview['summary']['occurrences'],1);self.assertEqual(preview['inconsistencies']['pages_without_weekly_grid'],[2]);self.assertEqual(len(blank()['classes']),41)
    def test_teaching_weekly_cost_audit_uses_reajusted_rates_without_monthly_or_shared_rateio(self):
        rooms=blank()['classes'];preview={'records':[
            {'professor':'Docente A','professor_id':'1','turma_operacional_7e7':'1º A','source_class_codes':['EFUND01MA'],'disciplina_componente':'MATEM','duracao_minutos':45,'status_temporal':'CONFIRMADO','status_vinculo_turma':'CONFIRMADO','aula_compartilhada':False},
            {'professor':'Docente A','professor_id':'1','turma_operacional_7e7':'6º A','source_class_codes':['EFUND06MA'],'disciplina_componente':'CIENC','duracao_minutos':50,'status_temporal':'CONFIRMADO','status_vinculo_turma':'CONFIRMADO','aula_compartilhada':False},
            {'professor':'Docente A','professor_id':'1','turma_operacional_7e7':None,'source_class_codes':['EFUND01MA','EFUND01MB'],'disciplina_componente':'MATEM | CIENC','duracao_minutos':40,'status_temporal':'CONFIRMADO','status_vinculo_turma':'PENDENTE','aula_compartilhada':True},
            {'professor':'Docente B','professor_id':'2','turma_operacional_7e7':None,'source_class_codes':['EMERE01MA'],'disciplina_componente':'ELETIVA','duracao_minutos':45,'status_temporal':'CONFIRMADO','status_vinculo_turma':'PENDENTE','aula_compartilhada':False},
        ]}
        before_preview=copy.deepcopy(preview);before_rooms=copy.deepcopy(rooms);audit=build_teaching_weekly_cost_audit(preview,rooms);one=next(item for item in audit['classes'] if item['class_name']=='1º A');six=next(item for item in audit['classes'] if item['class_name']=='6º A')
        self.assertEqual(preview,before_preview);self.assertEqual(rooms,before_rooms);self.assertEqual(len(audit['classes']),41);self.assertEqual(audit['summary']['included_unshared_confirmed_occurrences'],2);self.assertAlmostEqual(audit['summary']['reference_weekly_cost'],42.31);self.assertIsNone(audit['summary']['monthly_cost']);self.assertEqual(one['weekly_minutes'],45);self.assertAlmostEqual(one['reference_weekly_cost'],16.20);self.assertTrue(one['coverage_status'].startswith('PARCIAL'));self.assertEqual(six['coverage_status'],'CONCILIADO');self.assertAlmostEqual(six['reference_weekly_cost'],26.11);self.assertEqual(len(audit['included_occurrences']),2);self.assertEqual(len(audit['pending_occurrences']),2);self.assertEqual(audit['professors'][0]['durations'][40],1);self.assertIn('rateio financeiro',audit['pending_occurrences'][0]['reason'])
    def test_shared_lesson_equal_rateio_closes_to_original_without_duplicate_teacher_cost(self):
        rooms=blank()['classes'];preview={'records':[
            {'professor':'Docente','professor_id':'1','turma_operacional_7e7':None,'source_class_codes':['EFUND01MA','EFUND01MB'],'disciplina_componente':'ARTE | ARTE','duracao_minutos':45,'status_temporal':'CONFIRMADO','status_vinculo_turma':'PENDENTE','aula_compartilhada':True,'pagina':1},
            {'professor':'Docente EM','professor_id':'2','turma_operacional_7e7':None,'source_class_codes':['EMERE01MA','EMERE02MA','EMERE03MA'],'disciplina_componente':'CULTG','duracao_minutos':45,'status_temporal':'CONFIRMADO','status_vinculo_turma':'PENDENTE','aula_compartilhada':True,'pagina':2},
        ]};before=copy.deepcopy(preview);audit=build_teaching_weekly_cost_with_shared_rateio_audit(preview,rooms);rateio=audit['shared_lesson_rateio']
        self.assertEqual(preview,before);self.assertEqual(len(rooms),41);self.assertEqual(rateio['criterio'],'RATEIO_IGUALITARIO');self.assertEqual(rateio['origem'],'DECISAO_ADMINISTRATIVA');self.assertEqual(rateio['finalidade'],'CUSTO_POR_TURMA_PE');self.assertEqual(rateio['occurrences'],2);self.assertAlmostEqual(rateio['original_cost_total'],54.70);self.assertAlmostEqual(rateio['allocated_to_operational_classes'],16.20);self.assertEqual(rateio['allocated_pending_code'],38.50);self.assertEqual(rateio['closure_difference'],0);self.assertTrue(rateio['teacher_cost_counted_once'])
        two,three=rateio['allocations'];self.assertEqual([x['amount'] for x in two['allocations']],[8.1,8.1]);self.assertEqual(sum(x['amount'] for x in three['allocations']),38.5);self.assertEqual(sorted(x['amount'] for x in three['allocations']),[12.83,12.83,12.84]);self.assertEqual(sum(x['weekly_minutes'] for x in audit['professors']),90)
    def test_only_admin_can_update_enrollment_goal(self):
        user=self.user()
        with self.assertRaises(PermissionError):self.patch({'meta':700},user,'enrollments')
        self.patch({'meta':700},self.admin,'enrollments')
        self.assertEqual(self.store.state()[0]['meta'],700)
    def test_enrollment_math_and_conflict(self):
        self.patch({'classes':[self.room()]});state,version=self.store.state();en=copy.deepcopy(state['enrollments']);en['history']=[{'id':'h','classId':'c','type':'new','quantity':3,'date':'2026-09-10'}];self.patch({'enrollments':en},self.user(),module='enrollments');state,_=self.store.state();self.assertEqual(state['classes'][0]['students'],3);self.assertEqual(state['enrollments']['new'],3)
        with self.assertRaises(RuntimeError):self.store.patch(self.admin,{'classes':[]},version,'classes')
    def test_new_enrollment_keeps_report_class_rows_separate_from_history(self):
        before,_=self.store.state();g2=next(item for item in before['classes'] if item['name']=='G2 A');class_ids={item['id'] for item in before['classes']};before_total=sum(item['students'] for item in before['classes']);before_vacancies=sum(max(0,item['capacity']-item['students']) for item in before['classes']);enrollments=copy.deepcopy(before['enrollments']);enrollments['history'].append({'id':'isolated-g2-a-new','classId':g2['id'],'className':g2['name'],'type':'new','quantity':1,'date':'2026-09-11'})
        self.patch({'enrollments':enrollments},module='enrollments')
        after,_=self.store.state();changed=next(item for item in after['classes'] if item['id']==g2['id'])
        self.assertEqual(len(after['classes']),41);self.assertEqual({item['id'] for item in after['classes']},class_ids);self.assertEqual(changed['students'],g2['students']+1);self.assertEqual(after['enrollments']['new'],before['enrollments']['new']+1);self.assertEqual(sum(item['students'] for item in after['classes']),before_total+1);self.assertEqual(sum(max(0,item['capacity']-item['students']) for item in after['classes']),before_vacancies-1);self.assertEqual(len(after['enrollments']['history']),len(before['enrollments']['history'])+1)
    def test_benefit_audit_not_client_authored(self):
        user=self.user();self.patch({'benefits':[{'id':'b','name':'Fixture','type':'Fixture','quantity':1,'active':True,'history':[{'user':'spoof'}]}]},user,'benefits');s,_=self.store.state();self.assertEqual(s['benefits'][0]['history'][0]['user'],'Lívia')
        with self.assertRaises(PermissionError):self.patch({'benefits':[]},user,'benefits')
    def test_request_flow_reason_comments_and_outbox(self):
        user=self.user();self.patch({'requests':[self.request(user)]},user,'requests');state,_=self.store.state();r=state['requests'][0];r['status']='waiting';self.patch({'requests':[r]},user,'requests');r=self.store.state()[0]['requests'][0];r['status']='rejected'
        with self.assertRaises(ValueError):self.patch({'requests':[r]},module='requests')
        r['rejectionReason']='Motivo fixture';self.patch({'requests':[r]},module='requests');out=self.store.snapshot(self.admin)['state']['notifications'];self.assertTrue(out);self.assertTrue(all(x['attemptedAt'] is None and x['status']=='awaiting_configuration' for x in out))
    def test_user_cannot_approve(self):
        user=self.user();self.patch({'requests':[self.request(user)]},user,'requests');r=self.store.state()[0]['requests'][0];r['status']='waiting';self.patch({'requests':[r]},user,'requests');r=self.store.state()[0]['requests'][0];r['status']='approved'
        with self.assertRaises(PermissionError):self.patch({'requests':[r]},user,'requests')
    def test_private_requests_are_not_deleted_by_other_snapshot(self):
        first=self.user();second=self.user('Outro');self.patch({'requests':[self.request(first)]},first,'requests');self.patch({'requests':[self.request(second,'r2')]},second,'requests');self.assertEqual(len(self.store.state()[0]['requests']),2);self.assertEqual(len(self.store.snapshot(first)['state']['requests']),1)
    def test_bulk_requires_strong_confirmation(self):
        self.patch({'requests':[self.request(self.admin,'a'),self.request(self.admin,'b')]},module='requests')
        with self.assertRaises(PermissionError):self.patch({'requests':[]},module='requests')
        self.patch({'requests':[]},module='requests',confirmation='APAGAR TODAS AS SOLICITAÇÕES');self.assertEqual(self.store.state()[0]['requests'],[])
    def test_snapshot_no_finance_for_livia(self):
        self.patch({'budget':[{'id':'b','category':'Fixture','budget':10,'actual':0}]},module='budget');self.assertEqual(self.store.snapshot(self.user())['state']['budget'],[])
    def test_csv_preview_is_readonly_and_duplicates(self):
        s=blank();s['budget']=[{'id':'b','category':'Energia fixture','code':'01','budget':100,'actual':0}];csv=b'Data;Codigo;Descricao;Valor\n10/09/2026;01;Fixture;12,50\n10/09/2026;01;Fixture;12,50\n';before=copy.deepcopy(s);p=build_preview('fixture.csv',csv,'expenses',s,True);self.assertEqual(s,before);self.assertEqual(p['automatic'],1);self.assertEqual(p['duplicates'],1);updated=apply_preview(s,p,self.admin);self.assertEqual(updated['budget'][0]['actual'],12.5)
    def test_ambiguous_classification_waits(self):
        s=blank();p=build_preview('fixture.csv',b'Data;Descricao;Valor\n10/09/2026;Item fixture;20\n','expenses',s,True);self.assertEqual(p['pending'],1);updated=apply_preview(s,p,self.admin);self.assertEqual(len(updated['reviewQueue']),1);self.assertEqual(updated['budget'],[])
    def test_other_year_and_unconfirmed_source_rejected(self):
        content=b'Data;Descricao;Valor\n10/09/2025;Fixture;20\n'
        with self.assertRaises(ValueError):build_preview('fixture.csv',content,'expenses',blank(),True)
        with self.assertRaises(ValueError):build_preview('fixture.csv',content,'expenses',blank(),False)
    def test_delinquency_month_and_percentage(self):
        content='Mês;Financeira (%);Dívida;Alunos;Responsáveis\n09/2026;2,5;100;2;1\n'.encode();p=build_preview('fixture.csv',content,'financial',blank());s=apply_preview(blank(),p,self.admin);self.assertEqual(s['delinquency']['financialPercent'],2.5);self.assertEqual(s['delinquency']['monthly'][0]['month'],'2026-09')
    def test_report_xlsx_is_valid_and_strings_not_formulas(self):
        result=xlsx_report(['Valor','Descrição'],[[2.5,'=HYPERLINK("fixture")']]);from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(result));self.assertEqual(book.active['A2'].value,2.5);self.assertEqual(book.active['B2'].data_type,'s');book.close()
    def test_only_2026_sheet_is_imported(self):
        content=xlsx_report(['Categoria','Orçado'],[['Fixture',25]])
        source=zipfile.ZipFile(io.BytesIO(content));buffer=io.BytesIO()
        with zipfile.ZipFile(buffer,'w') as z:
            for item in source.infolist():
                data=source.read(item.filename)
                if item.filename=='xl/workbook.xml':data=data.replace('Relatório CAJ'.encode(),'DESPESAS 2026'.encode())
                z.writestr(item,data)
        source.close();p=build_preview('fixture.xlsx',buffer.getvalue(),'budget',blank());self.assertEqual(p['sheet'],'DESPESAS 2026');self.assertEqual(p['rows'][0]['budget'],25)
        with self.assertRaises(ValueError):build_preview('other.xlsx',content,'budget',blank())
    def test_pdf_text_import(self):
        from reportlab.pdfgen import canvas
        buffer=io.BytesIO();c=canvas.Canvas(buffer);c.drawString(30,800,'Categoria;Orcado');c.drawString(30,780,'Fixture;50');c.save();p=build_preview('fixture.pdf',buffer.getvalue(),'budget',blank(),True);self.assertEqual(p['rows'][0]['budget'],50)
    def test_enrollment_delete_and_audit(self):
        self.patch({'classes':[self.room()]});en=self.store.state()[0]['enrollments'];en['history']=[{'id':'h','classId':'c','type':'re','quantity':2,'date':'2026-09-10'}];self.patch({'enrollments':en},module='enrollments');en=self.store.state()[0]['enrollments'];en['history']=[];self.patch({'enrollments':en},module='enrollments');s=self.store.state()[0];self.assertEqual(s['classes'][0]['students'],0);self.assertEqual(s['enrollments']['re'],0);self.assertTrue(self.store.snapshot(self.admin)['state']['audit'])
    def test_new_month_does_not_reuse_old_percentage(self):
        s=blank();s['delinquency']['monthly']=[{'month':'2026-08','financialPercent':9}];p=build_preview('fixture.csv','Mês;Contábil (%)\n09/2026;2,5\n'.encode(),'accounting',s);updated=apply_preview(s,p,self.admin);self.assertIsNone(updated['delinquency']['financialPercent']);self.assertEqual(updated['delinquency']['accountingPercent'],2.5)
    def test_pdf_has_repeated_letterhead_and_pages(self):
        from pypdf import PdfReader
        result=pdf_report('Validação isolada — sem dados reais',['Indicador','Valor'],[['Linha de teste '+str(i),i] for i in range(100)]);reader=PdfReader(io.BytesIO(result));self.assertGreater(len(reader.pages),1)
        for page in reader.pages:
            text=page.extract_text();self.assertIn('Colégio Adventista de Juazeiro',text);self.assertIn('07.114.699/0047-42',text);self.assertIn('Página',text)

if __name__=='__main__':unittest.main(verbosity=2)
